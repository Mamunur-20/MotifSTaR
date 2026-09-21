#!/usr/bin/env python3
"""MotifSTaR: allele-level STR motif reconstruction from REViewer SVGs.

The program reconstructs one nucleotide sequence per full repeat-spanning read,
normalises that sequence to the biological reporting strand, derives complete
motif compositions, and then infers allele-level structures by read support.

It intentionally does *not* replace the ExpansionHunter genotype.  REViewer
panel labels (or an override table) define the expected genotype branch, while
reported allele sizes are reconstructed from the complete supported motif path.
"""

from __future__ import annotations

import argparse
import colorsys
import csv
import gzip
import hashlib
import html
import itertools
import json
import logging
import math
import os
import re
import shutil
import statistics
import sys
import traceback
import zipfile
import xml.etree.ElementTree as ElementTree
from collections import Counter, defaultdict
from concurrent.futures import ProcessPoolExecutor
from dataclasses import asdict, dataclass, field, replace
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Iterator, Mapping, Sequence
import xml.etree.ElementTree as ET


VERSION = "1.4.7"
PRIMARY_REPEAT_ONLY_LOCI = frozenset(
    {"HTT", "ATXN7", "FXN", "ATXN8OS", "PRNP"}
)
SEQUENCE_COMPOSITION_FIGURE_MAX_ROWS = 40
INDEL_SENTINEL = "NA(indel detected)"
INDEL_REVIEW_FLAG = "indel_detected_manual_review_required"
INDEL_REVIEW_REASON = "Indel detected manual review required"
TARGET_INDEL_REASONS = frozenset(
    {"target_deletion", "target_insertion_bases_unavailable"}
)
DNA_BASES = frozenset("ACGTN")
IUPAC_BASES = frozenset("ACGTRYSWKMBDHVN")
IUPAC_MATCH = {
    "A": frozenset("A"), "C": frozenset("C"), "G": frozenset("G"),
    "T": frozenset("T"), "R": frozenset("AG"), "Y": frozenset("CT"),
    "S": frozenset("GC"), "W": frozenset("AT"), "K": frozenset("GT"),
    "M": frozenset("AC"), "B": frozenset("CGT"), "D": frozenset("AGT"),
    "H": frozenset("ACT"), "V": frozenset("ACG"), "N": frozenset("ACGTN"),
}
COMPLEMENT = str.maketrans(
    "ACGTRYSWKMBDHVNacgtryswkmbdhvn",
    "TGCAYRSWMKVHDBNtgcayrswmkvhdbn",
)

MAIN_COLUMNS = [
    "sampleID",
    "allele1",
    "allele2",
    "allele1_size",
    "allele2_size",
    "read_depth_allele1",
    "read_depth_allele2",
]

QC_COLUMNS = [
    "batch",
    "sampleID",
    "locus",
    "parent_locus",
    "component_index",
    "svg_component_index",
    "locus_structure",
    "component_count",
    "allele1_components",
    "allele2_components",
    "source_svg",
    "EH_genotype",
    "genotype_source",
    "genotype_class",
    "EH_repeat_unit",
    "reported_repeat_unit",
    "gene",
    "gene_strand",
    "reporting_orientation",
    "catalog_status",
    "annotation_status",
    "n_panels",
    "n_full_spanning_reads_raw",
    "n_total_full_spanning_reads",
    "n_callable_reads",
    "n_cross_panel_duplicates_removed",
    "svg_read_rows_allele1",
    "svg_read_rows_allele2",
    "full_spanning_read_depth_allele1",
    "full_spanning_read_depth_allele2",
    "callable_full_spanning_read_depth_allele1",
    "callable_full_spanning_read_depth_allele2",
    "consensus_eligible_read_depth_allele1",
    "consensus_eligible_read_depth_allele2",
    "supporting_read_depth_allele1",
    "supporting_read_depth_allele2",
    "interruption_threshold_allele1",
    "interruption_threshold_allele2",
    "consensus_interruption_positions_allele1",
    "consensus_interruption_positions_allele2",
    "evidence_scope_allele1",
    "evidence_scope_allele2",
    "n_exact_motif_composition_clusters",
    "n_motif_composition_clusters",
    "n_strong_motif_composition_clusters",
    "same_size_candidate_cluster_count",
    "same_size_selected_cluster_count",
    "same_size_selected_cluster_coverage",
    "same_size_assigned_read_count",
    "same_size_unassigned_read_count",
    "same_size_ambiguous_read_count",
    "allele1_support_fraction",
    "allele2_support_fraction",
    "minor_cluster_count",
    "pure_or_interrupted_allele1",
    "pure_or_interrupted_allele2",
    "allele1_motif_classification",
    "allele2_motif_classification",
    "qc_status",
    "qc_flags",
    "notes",
]

READ_COLUMNS = [
    "batch",
    "sampleID",
    "locus",
    "parent_locus",
    "component_index",
    "svg_component_index",
    "panel",
    "panel_declared_size",
    "read_row_y",
    "read_signature",
    "full_spanning",
    "technical_sequence",
    "reported_sequence",
    "motif_composition",
    "n_repeat_units",
    "callable",
    "exclusion_reason",
]


class ToolError(RuntimeError):
    """Expected, user-actionable error."""


def reverse_complement(sequence: str) -> str:
    """Return the reverse complement of a DNA/IUPAC sequence."""
    return sequence.translate(COMPLEMENT)[::-1]


def natural_key(value: Any) -> tuple[Any, ...]:
    return tuple(
        int(part) if part.isdigit() else part.casefold()
        for part in re.split(r"(\d+)", str(value))
    )


def normalise_key(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", "", str(value or "").upper())


def normalise_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def sanitise_filename(value: str) -> str:
    clean = re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("._")
    return clean or "unnamed"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def as_list(value: Any) -> list[Any]:
    if value is None:
        return []
    return list(value) if isinstance(value, list) else [value]


def safe_float(value: Any) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def safe_int(value: Any) -> int | None:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def format_fraction(numerator: int | None, denominator: int | None) -> str:
    if numerator is None or not denominator:
        return "NA"
    return f"{numerator / denominator:.4f}"


def semicolon_join(values: Iterable[str]) -> str:
    return ";".join(sorted({v for v in values if v}))


def rle_motifs(motifs: Sequence[str]) -> str:
    """Collapse motifs using the required ``(MOTIF)n`` notation."""
    if not motifs:
        return "NA"
    output: list[str] = []
    current = motifs[0]
    count = 1
    for motif in motifs[1:]:
        if motif == current:
            count += 1
        else:
            output.append(f"({current}){count}")
            current, count = motif, 1
    output.append(f"({current}){count}")
    return "".join(output)


def parse_rle_motifs(composition: str) -> list[str]:
    if not composition or composition == "NA":
        return []
    parts = re.findall(r"\(([A-Za-z]+)\)(\d+)", composition)
    if not parts or "".join(f"({m}){n}" for m, n in parts) != composition:
        return []
    motifs: list[str] = []
    for motif, count in parts:
        motifs.extend([motif.upper()] * int(count))
    return motifs


def motif_distance(left: Sequence[str], right: Sequence[str]) -> int:
    if len(left) != len(right):
        return max(len(left), len(right))
    return sum(a != b for a, b in zip(left, right))


def classify_structure(motifs: Sequence[str], expected: str) -> str:
    if not motifs:
        return "unresolved"
    return "pure" if all(motif_matches(m, expected) for m in motifs) else "interrupted"


def motif_matches(observed: str, pattern: str) -> bool:
    observed = observed.upper()
    pattern = pattern.upper()
    return len(observed) == len(pattern) and all(
        base in IUPAC_MATCH.get(code, frozenset(code))
        for base, code in zip(observed, pattern)
    )


def split_motifs(sequence: str, motif_length: int) -> tuple[str, ...] | None:
    if motif_length < 1 or not sequence or len(sequence) % motif_length:
        return None
    return tuple(
        sequence[i : i + motif_length]
        for i in range(0, len(sequence), motif_length)
    )


@dataclass(frozen=True)
class CatalogComponent:
    index: int
    variant_id: str
    reference_region: str
    motif: str
    variant_type: str = "Repeat"
    is_primary: bool = False
    quantifier: str = ""


@dataclass
class CatalogLocus:
    locus_id: str
    gene: str
    locus_structure: str
    repeat_unit: str
    main_reference_region: str
    components: list[CatalogComponent]
    pathogenic_motifs: list[str] = field(default_factory=list)
    benign_motifs: list[str] = field(default_factory=list)
    normal_max: int | None = None
    pathogenic_min: int | None = None
    raw: dict[str, Any] = field(default_factory=dict)

    @property
    def primary_index(self) -> int:
        for component in self.components:
            if component.is_primary:
                return component.index
        return 0


@dataclass(frozen=True)
class Annotation:
    key: str
    strand: str
    orientation: str
    reported_motif: str = ""
    gene: str = ""
    raw: Mapping[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class HeaderLabel:
    x: float
    y: float
    units: int


@dataclass(frozen=True)
class Rect:
    x: float
    y: float
    width: float
    height: float
    fill: str

    @property
    def end(self) -> float:
        return self.x + self.width


@dataclass(frozen=True)
class Glyph:
    x: float
    y: float
    text: str


@dataclass(frozen=True)
class LineSegment:
    x1: float
    y1: float
    x2: float
    y2: float


@dataclass
class RepeatBlock:
    index: int
    start: float
    end: float
    unit_width: float
    declared_units: int
    fill: str
    pitch: float
    motif_length: int
    scaffold_positions: list[float]
    scaffold_sequence: str


@dataclass
class SvgRead:
    panel: int
    row_y: float
    signature: str
    block_sequences: list[str]
    block_indel_reasons: list[str]
    full_spanning: bool
    exclusion_reason: str = ""


@dataclass
class SvgPanel:
    index: int
    reference_y: float
    blocks: list[RepeatBlock]
    reads: list[SvgRead]
    excluded_read_counts: dict[str, int]


@dataclass
class ParsedSvg:
    panels: list[SvgPanel]
    warnings: list[str]


@dataclass(frozen=True)
class ComponentMapping:
    """Link a catalogue component to its rendered SVG block.

    ReViewer omits a zero-copy optional component, so catalogue and SVG indices
    cannot be assumed to be identical.  ``catalog_index`` is ``None`` only when
    no catalogue model exists for the SVG locus.
    """

    catalog_index: int | None
    svg_index: int
    score: int = 0


@dataclass
class ReadEvidence:
    panel: int
    declared_size: int
    row_y: float
    signature: str
    technical_sequence: str
    reported_sequence: str
    motifs: tuple[str, ...] | None
    exclusion_reason: str = ""

    @property
    def callable(self) -> bool:
        return self.motifs is not None and not self.exclusion_reason


@dataclass
class Cluster:
    representative: tuple[str, ...]
    count: int
    first_order: int
    exact_paths: list[tuple[str, ...]]

    @property
    def composition(self) -> str:
        return rle_motifs(self.representative)


@dataclass
class RunConfig:
    min_spanning_reads: int = 3
    min_cluster_reads: int = 3
    max_near_motif_distance: int = 0
    near_cluster_parent_ratio: float = 3.0
    strong_cluster_fraction: float = 0.0
    heterozygous_interruption_min_fraction: float = 0.50
    same_size_min_cluster_reads: int = 2
    same_size_min_cluster_fraction: float = 0.20
    same_size_min_top_two_coverage: float = 0.90
    same_size_within_cluster_consensus_min_fraction: float = 0.50
    same_size_max_reported_clusters: int = 2
    same_size_max_assignment_distance: int = 2
    homozygous_deduplication: str = "none"
    component_mode: str = "combined"
    unresolved_strand_policy: str = "skip"
    allow_duplicate_sample_locus: bool = False


@dataclass
class WorkerContext:
    svg_root: str
    catalog_loci: list[CatalogLocus]
    annotation_rows: list[Annotation]
    genotype_overrides: dict[str, list[int]]
    file_overrides: dict[str, dict[str, str]]
    selected_loci: list[str]
    config: RunConfig


_WORKER_CONTEXT: WorkerContext | None = None
_WORKER_CATALOG_INDEX: dict[str, CatalogLocus] = {}
_WORKER_ANNOTATION_INDEX: dict[str, Annotation] = {}


def _column_index(cell_reference: str) -> int:
    letters = re.match(r"[A-Za-z]+", cell_reference or "")
    if not letters:
        return 0
    index = 0
    for char in letters.group(0).upper():
        index = index * 26 + ord(char) - 64
    return index - 1


def _read_xlsx_rows_stdlib(path: Path, sheet_name: str | None = None) -> list[list[Any]]:
    """Small dependency-free XLSX reader used if openpyxl is unavailable."""
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    pkg_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"

    with zipfile.ZipFile(path) as archive:
        shared: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in root.findall(f"{{{main_ns}}}si"):
                shared.append("".join(node.text or "" for node in item.iter(f"{{{main_ns}}}t")))

        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {
            rel.attrib["Id"]: rel.attrib["Target"]
            for rel in relationships.findall(f"{{{pkg_rel_ns}}}Relationship")
        }
        sheets = workbook.find(f"{{{main_ns}}}sheets")
        if sheets is None or not list(sheets):
            return []
        chosen = None
        for sheet in sheets:
            if sheet_name is None or sheet.attrib.get("name") == sheet_name:
                chosen = sheet
                break
        if chosen is None:
            raise ToolError(f"Worksheet {sheet_name!r} not found in {path}")
        relationship_id = chosen.attrib.get(f"{{{rel_ns}}}id")
        target = targets.get(relationship_id or "")
        if not target:
            raise ToolError(f"Cannot resolve worksheet XML for {path}")
        target = target.replace("\\", "/").lstrip("/")
        if not target.startswith("xl/"):
            target = "xl/" + target
        sheet_root = ET.fromstring(archive.read(target))

        rows: list[list[Any]] = []
        for row in sheet_root.findall(f".//{{{main_ns}}}sheetData/{{{main_ns}}}row"):
            values: dict[int, Any] = {}
            for cell in row.findall(f"{{{main_ns}}}c"):
                col = _column_index(cell.attrib.get("r", ""))
                cell_type = cell.attrib.get("t", "")
                value_node = cell.find(f"{{{main_ns}}}v")
                inline = cell.find(f"{{{main_ns}}}is")
                raw = value_node.text if value_node is not None else None
                value: Any = None
                if cell_type == "s" and raw is not None:
                    value = shared[int(raw)]
                elif cell_type == "inlineStr" and inline is not None:
                    value = "".join(
                        node.text or "" for node in inline.iter(f"{{{main_ns}}}t")
                    )
                elif cell_type in {"str", "e"}:
                    value = raw or ""
                elif cell_type == "b":
                    value = raw == "1"
                elif raw is not None:
                    try:
                        number = float(raw)
                        value = int(number) if number.is_integer() else number
                    except ValueError:
                        value = raw
                values[col] = value
            if values:
                width = max(values) + 1
                rows.append([values.get(i) for i in range(width)])
        return rows


def read_xlsx_rows(path: Path, sheet_name: str | None = None) -> list[list[Any]]:
    try:
        from openpyxl import load_workbook  # type: ignore

        workbook = load_workbook(path, read_only=True, data_only=True)
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ToolError(f"Worksheet {sheet_name!r} not found in {path}")
            sheet = workbook[sheet_name]
        else:
            sheet = workbook[workbook.sheetnames[0]]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    except ImportError:
        return _read_xlsx_rows_stdlib(path, sheet_name)


def read_delimited_rows(path: Path) -> list[list[Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(8192)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
        except csv.Error:
            dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
        return [list(row) for row in csv.reader(handle, dialect)]


def read_table_records(path: Path, sheet_name: str | None = None) -> list[dict[str, Any]]:
    rows = (
        read_xlsx_rows(path, sheet_name)
        if path.suffix.lower() in {".xlsx", ".xlsm"}
        else read_delimited_rows(path)
    )
    rows = [list(row) for row in rows if any(value not in {None, ""} for value in row)]
    if not rows:
        return []
    headers = [normalise_header(value) for value in rows[0]]
    records: list[dict[str, Any]] = []
    for row in rows[1:]:
        record: dict[str, Any] = {}
        for index, header in enumerate(headers):
            if header and index < len(row):
                record[header] = row[index]
        if any(value not in {None, ""} for value in record.values()):
            records.append(record)
    return records


def first_value(record: Mapping[str, Any], aliases: Sequence[str]) -> Any:
    for alias in aliases:
        value = record.get(alias)
        if value not in {None, ""}:
            return value
    return None


def normalise_strand(value: Any) -> str:
    text = str(value or "").strip().lower().replace(" ", "_")
    if text in {"+", "+1", "1", "forward", "plus", "positive", "sense"}:
        return "+"
    if text in {"-", "-1", "reverse", "minus", "negative", "antisense"}:
        return "-"
    return ""


def normalise_orientation(value: Any, strand: str) -> str:
    text = str(value or "").strip().lower().replace("-", "_").replace(" ", "_")
    if text in {"reference", "technical", "eh", "forward", "same"}:
        return "reference"
    if text in {"reverse", "reverse_complement", "revcomp", "rc"}:
        return "reverse_complement"
    if text in {"coding", "gene", "transcript", "biological"}:
        if strand not in {"+", "-"}:
            return ""
        return "reverse_complement" if strand == "-" else "reference"
    if text:
        return ""
    return "reverse_complement" if strand == "-" else "reference" if strand == "+" else ""


def load_annotations(path: Path, sheet_name: str | None = None) -> list[Annotation]:
    records = read_table_records(path, sheet_name)
    if not records:
        raise ToolError(f"No annotation rows found in {path}")
    output: list[Annotation] = []
    for record in records:
        key = first_value(
            record,
            ["locusid", "locus_id", "locus", "gene", "gene_symbol", "symbol"],
        )
        strand = normalise_strand(
            first_value(record, ["strand", "gene_strand", "genestrand", "orientation"])
        )
        if not key:
            continue
        explicit_orientation = first_value(
            record, ["reporting_orientation", "biological_orientation"]
        )
        orientation = normalise_orientation(explicit_orientation, strand)
        reported_motif = str(
            first_value(
                record,
                [
                    "reported_repeat_unit",
                    "reported_motif",
                    "biological_repeat_unit",
                    "coding_repeat_unit",
                ],
            )
            or ""
        ).strip().upper()
        output.append(
            Annotation(
                key=str(key).strip(),
                strand=strand,
                orientation=orientation,
                reported_motif=reported_motif,
                gene=str(first_value(record, ["gene", "gene_symbol", "symbol"]) or key),
                raw=dict(record),
            )
        )
    if not output:
        raise ToolError(
            f"Annotation table {path} has no usable locus/gene and strand rows"
        )
    return output


def _extract_repeat_terms(locus_structure: str) -> list[tuple[str, str]]:
    return [
        (match.group(1).upper(), match.group(2))
        for match in re.finditer(r"\(([A-Za-z]+)\)(\*|\+|\?)", locus_structure or "")
    ]


def _extract_repeat_groups(locus_structure: str) -> list[str]:
    """Return repeated motifs while retaining the legacy helper API."""
    return [motif for motif, _ in _extract_repeat_terms(locus_structure)]


def _locus_structure_tokens(locus_structure: str) -> list[tuple[str, Any]]:
    """Return ordered repeat-component indices and intervening fixed sequences."""
    tokens: list[tuple[str, Any]] = []
    cursor = 0
    component_index = 0
    for match in re.finditer(r"\(([A-Za-z]+)\)(\*|\+|\?)", locus_structure or ""):
        literal = re.sub(r"[^A-Za-z]", "", locus_structure[cursor : match.start()]).upper()
        if literal:
            tokens.append(("literal", literal))
        tokens.append(("repeat", component_index))
        component_index += 1
        cursor = match.end()
    literal = re.sub(r"[^A-Za-z]", "", (locus_structure or "")[cursor:]).upper()
    if literal:
        tokens.append(("literal", literal))
    return tokens


def load_catalog(path: Path) -> list[CatalogLocus]:
    with path.open("r", encoding="utf-8-sig") as handle:
        payload = json.load(handle)
    if isinstance(payload, dict):
        for key in ("Loci", "loci", "VariantCatalog", "variant_catalog"):
            if isinstance(payload.get(key), list):
                payload = payload[key]
                break
        else:
            payload = [payload]
    if not isinstance(payload, list):
        raise ToolError(f"Unsupported catalogue JSON structure in {path}")

    loci: list[CatalogLocus] = []
    for raw_value in payload:
        if not isinstance(raw_value, dict):
            continue
        raw = dict(raw_value)
        locus_id = str(raw.get("LocusId") or raw.get("LocusID") or "").strip()
        if not locus_id:
            continue
        structure = str(raw.get("LocusStructure") or "")
        repeat_unit = str(raw.get("RepeatUnit") or "").upper()
        repeat_terms = _extract_repeat_terms(structure)
        groups = [motif for motif, _ in repeat_terms]
        quantifiers = [quantifier for _, quantifier in repeat_terms]
        regions = [str(value) for value in as_list(raw.get("ReferenceRegion"))]
        variant_ids = [str(value) for value in as_list(raw.get("VariantId"))]
        variant_types = [str(value) for value in as_list(raw.get("VariantType"))]
        n_components = max(len(groups), len(regions), len(variant_ids), 1)
        if not groups:
            groups = [repeat_unit]
            quantifiers = [""]
        while len(groups) < n_components:
            groups.append(repeat_unit or groups[-1])
            quantifiers.append("")
        if not regions:
            regions = [str(raw.get("MainReferenceRegion") or "")]
        while len(regions) < n_components:
            regions.append("")
        if not variant_ids:
            variant_ids = [locus_id]
        while len(variant_ids) < n_components:
            variant_ids.append(f"{locus_id}_component{len(variant_ids) + 1}")
        if not variant_types:
            variant_types = ["Repeat"] * n_components
        while len(variant_types) < n_components:
            variant_types.append("Repeat")

        main_region = str(raw.get("MainReferenceRegion") or "")
        primary_index: int | None = None
        if main_region and main_region in regions:
            primary_index = regions.index(main_region)
        if primary_index is None:
            for index, variant_id in enumerate(variant_ids):
                if normalise_key(variant_id) == normalise_key(locus_id):
                    primary_index = index
                    break
        if primary_index is None and repeat_unit:
            for index, motif in enumerate(groups):
                if motif == repeat_unit:
                    primary_index = index
                    break
        if primary_index is None:
            primary_index = 0

        components: list[CatalogComponent] = []
        for index in range(n_components):
            motif = repeat_unit if index == primary_index and repeat_unit else groups[index]
            components.append(
                CatalogComponent(
                    index=index,
                    variant_id=variant_ids[index],
                    reference_region=regions[index],
                    motif=motif.upper(),
                    variant_type=variant_types[index],
                    is_primary=index == primary_index,
                    quantifier=quantifiers[index],
                )
            )

        diseases = [value for value in as_list(raw.get("Diseases")) if isinstance(value, dict)]
        normal_values = [safe_int(item.get("NormalMax")) for item in diseases]
        pathogenic_values = [safe_int(item.get("PathogenicMin")) for item in diseases]
        normal_values = [value for value in normal_values if value is not None]
        pathogenic_values = [value for value in pathogenic_values if value is not None]
        loci.append(
            CatalogLocus(
                locus_id=locus_id,
                gene=str(raw.get("Gene") or "").strip(),
                locus_structure=structure,
                repeat_unit=repeat_unit,
                main_reference_region=main_region,
                components=components,
                pathogenic_motifs=[str(v).upper() for v in as_list(raw.get("PathogenicMotifs"))],
                benign_motifs=[str(v).upper() for v in as_list(raw.get("BenignMotifs"))],
                normal_max=max(normal_values) if normal_values else None,
                pathogenic_min=min(pathogenic_values) if pathogenic_values else None,
                raw=raw,
            )
        )
    if not loci:
        raise ToolError(f"No loci found in catalogue {path}")
    return loci


def build_catalog_index(loci: Sequence[CatalogLocus]) -> dict[str, CatalogLocus]:
    index: dict[str, CatalogLocus] = {}
    for locus in loci:
        for value in [locus.locus_id, locus.gene]:
            key = normalise_key(value)
            if key:
                index.setdefault(key, locus)
        for component in locus.components:
            key = normalise_key(component.variant_id)
            if key:
                index.setdefault(key, locus)
    return index


def resolve_catalog_locus(
    svg_locus: str, index: Mapping[str, CatalogLocus]
) -> CatalogLocus | None:
    candidates = [svg_locus]
    if svg_locus.upper().startswith("PRE-"):
        candidates.append(svg_locus[4:])
    candidates.append(re.sub(r"_\d+$", "", svg_locus))
    if normalise_key(svg_locus) == "ATXN8OS":
        candidates.append("ATXN8")
    for candidate in candidates:
        locus = index.get(normalise_key(candidate))
        if locus is not None:
            return locus
    return None


def build_annotation_index(rows: Sequence[Annotation]) -> dict[str, Annotation]:
    index: dict[str, Annotation] = {}
    for row in rows:
        index[normalise_key(row.key)] = row
        if row.gene:
            index.setdefault(normalise_key(row.gene), row)
    return index


def resolve_annotation(
    svg_locus: str,
    catalog: CatalogLocus | None,
    index: Mapping[str, Annotation],
) -> Annotation | None:
    candidates: list[str] = [svg_locus]
    if svg_locus.upper().startswith("PRE-"):
        candidates.append(svg_locus[4:])
    candidates.append(re.sub(r"_\d+$", "", svg_locus))
    if catalog:
        candidates.extend([catalog.locus_id, catalog.gene])
    if normalise_key(svg_locus) == "ATXN8OS":
        candidates.append("ATXN8")
    for candidate in candidates:
        annotation = index.get(normalise_key(candidate))
        if annotation is not None:
            return annotation
    return None


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _parse_svg_number(value: Any, default: float = 0.0) -> float:
    if value is None:
        return default
    match = re.search(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)", str(value))
    return float(match.group(0)) if match else default


def _style_value(element: ET.Element, key: str) -> str:
    if key in element.attrib:
        return str(element.attrib[key])
    style = str(element.attrib.get("style", ""))
    for part in style.split(";"):
        if ":" in part:
            name, value = part.split(":", 1)
            if name.strip() == key:
                return value.strip()
    return ""


def _normalise_colour(value: str) -> str:
    return re.sub(r"\s+", "", str(value or "")).casefold()


def _round_coord(value: float) -> float:
    return round(value, 3)


def _deduplicate_headers(headers: Sequence[HeaderLabel]) -> list[HeaderLabel]:
    seen: set[tuple[float, float, int]] = set()
    output: list[HeaderLabel] = []
    for header in sorted(headers, key=lambda h: (h.y, h.x, h.units)):
        key = (_round_coord(header.x), _round_coord(header.y), header.units)
        if key not in seen:
            seen.add(key)
            output.append(header)
    return output


def _mode_float(values: Sequence[float]) -> float:
    if not values:
        return 0.0
    counts = Counter(round(value, 2) for value in values)
    return float(counts.most_common(1)[0][0])


def _infer_pitch(positions: Sequence[float], fallback: float = 10.0) -> float:
    ordered = sorted(set(round(value, 4) for value in positions))
    differences = [
        round(right - left, 4)
        for left, right in zip(ordered, ordered[1:])
        if right - left > 0.1
    ]
    if not differences:
        return fallback
    counts = Counter(differences)
    most_common = counts.most_common()
    highest = most_common[0][1]
    candidates = [difference for difference, count in most_common if count == highest]
    pitch = min(candidates)
    return pitch if pitch > 0 else fallback


def _group_reference_rectangles(rectangles: Sequence[Rect]) -> list[list[Rect]]:
    groups: list[list[Rect]] = []
    tolerance = 0.2
    for rectangle in sorted(rectangles, key=lambda r: (r.x, r.width, r.fill)):
        if not groups:
            groups.append([rectangle])
            continue
        previous = groups[-1][-1]
        contiguous = abs(rectangle.x - previous.end) <= tolerance
        same_fill = rectangle.fill.casefold() == previous.fill.casefold()
        same_width = abs(rectangle.width - previous.width) <= tolerance
        if contiguous and same_fill and same_width:
            groups[-1].append(rectangle)
        else:
            groups.append([rectangle])
    return groups


def _merge_intervals(rectangles: Sequence[Rect], tolerance: float = 0.2) -> list[tuple[float, float]]:
    intervals: list[list[float]] = []
    for rectangle in sorted(rectangles, key=lambda r: (r.x, r.end)):
        if not intervals or rectangle.x > intervals[-1][1] + tolerance:
            intervals.append([rectangle.x, rectangle.end])
        else:
            intervals[-1][1] = max(intervals[-1][1], rectangle.end)
    return [(left, right) for left, right in intervals]


def _coverage_reason(
    rectangles: Sequence[Rect],
    start: float,
    end: float,
    pitch: float,
    flank_fills: set[str] | None = None,
) -> tuple[bool, str]:
    intervals = _merge_intervals(rectangles, tolerance=max(0.2, pitch * 0.05))
    if not intervals:
        return False, "no_aligned_sequence"
    left_needed = start - pitch * 0.95
    right_needed = end + pitch * 0.95
    continuous = any(left <= left_needed and right >= right_needed for left, right in intervals)
    if continuous:
        if flank_fills:
            has_left_flank = any(
                _normalise_colour(rectangle.fill) in flank_fills
                and rectangle.x < start
                and rectangle.end <= start + pitch * 0.2
                for rectangle in rectangles
            )
            has_right_flank = any(
                _normalise_colour(rectangle.fill) in flank_fills
                and rectangle.end > end
                and rectangle.x >= end - pitch * 0.2
                for rectangle in rectangles
            )
            if not has_left_flank:
                return False, "missing_left_blue_flank"
            if not has_right_flank:
                return False, "missing_right_blue_flank"
        return True, ""
    left_covered = any(left <= left_needed and right >= start for left, right in intervals)
    right_covered = any(left <= end and right >= right_needed for left, right in intervals)
    repeat_covered = any(left <= start and right >= end for left, right in intervals)
    if not left_covered:
        return False, "missing_left_flank"
    if not right_covered:
        return False, "missing_right_flank"
    if not repeat_covered:
        return False, "gap_over_repeat_locus"
    return False, "not_full_spanning"


def _nearest_glyph_map(
    glyphs: Sequence[Glyph], target_y: float, height: float
) -> dict[float, str]:
    candidates = [
        glyph
        for glyph in glyphs
        if target_y - 0.5 <= glyph.y <= target_y + height + 0.5
        and len(glyph.text) <= 1
    ]
    if not candidates:
        return {}
    y_counts = Counter(round(glyph.y, 3) for glyph in candidates)
    expected_y = target_y + height / 2.0
    chosen_y = min(
        y_counts,
        key=lambda value: (-y_counts[value], abs(value - expected_y)),
    )
    return {
        _round_coord(glyph.x): glyph.text
        for glyph in candidates
        if abs(glyph.y - chosen_y) <= 0.01
    }


def _lookup_glyph(mapping: Mapping[float, str], x: float, tolerance: float) -> str:
    key = _round_coord(x)
    if key in mapping:
        return mapping[key]
    if not mapping:
        return ""
    nearest = min(mapping, key=lambda value: abs(value - x))
    return mapping[nearest] if abs(nearest - x) <= tolerance else ""


def _read_svg_xml(path: Path) -> ET.Element:
    if path.name.lower().endswith(".svg.gz"):
        with gzip.open(path, "rb") as handle:
            return ET.fromstring(handle.read())
    return ET.parse(path).getroot()


def parse_reviewer_svg(path: Path) -> ParsedSvg:
    """Parse ReViewer geometry without relying on XML attribute order or fixed pixels."""
    root = _read_svg_xml(path)
    rectangles: list[Rect] = []
    glyphs: list[Glyph] = []
    line_segments: list[LineSegment] = []
    headers: list[HeaderLabel] = []
    warnings: list[str] = []

    flank_fills = {"#8da0cb"}
    for gradient in root.iter():
        if _local_name(gradient.tag) != "linearGradient":
            continue
        gradient_id = str(gradient.attrib.get("id", "")).casefold()
        if "blue" not in gradient_id:
            continue
        for stop in gradient:
            if _local_name(stop.tag) == "stop":
                colour = _normalise_colour(_style_value(stop, "stop-color"))
                if colour:
                    flank_fills.add(colour)

    for element in root.iter():
        name = _local_name(element.tag)
        if name == "rect":
            width = _parse_svg_number(element.attrib.get("width"))
            height = _parse_svg_number(element.attrib.get("height"))
            if width > 0 and height > 0:
                rectangles.append(
                    Rect(
                        x=_parse_svg_number(element.attrib.get("x")),
                        y=_parse_svg_number(element.attrib.get("y")),
                        width=width,
                        height=height,
                        fill=_style_value(element, "fill"),
                    )
                )
        elif name == "text":
            content = "".join(element.itertext())
            x = _parse_svg_number(element.attrib.get("x"))
            y = _parse_svg_number(element.attrib.get("y"))
            glyphs.append(Glyph(x=x, y=y, text=content))
            header_match = re.fullmatch(r"\s*(\d+)\s+units?\s*", content, re.IGNORECASE)
            if header_match:
                headers.append(HeaderLabel(x=x, y=y, units=int(header_match.group(1))))
        elif name == "line" and not (
            element.attrib.get("marker-start") or element.attrib.get("marker-end")
        ):
            line_segments.append(
                LineSegment(
                    x1=_parse_svg_number(element.attrib.get("x1")),
                    y1=_parse_svg_number(element.attrib.get("y1")),
                    x2=_parse_svg_number(element.attrib.get("x2")),
                    y2=_parse_svg_number(element.attrib.get("y2")),
                )
            )

    headers = _deduplicate_headers(headers)
    if not headers:
        raise ToolError("no ReViewer 'N units' panel labels were found")

    gradient_rectangles = [
        rectangle
        for rectangle in rectangles
        if "url(" in rectangle.fill.casefold()
    ]
    if not gradient_rectangles:
        raise ToolError("no ReViewer reference repeat rectangles were found")
    reference_rows: dict[float, list[Rect]] = defaultdict(list)
    for rectangle in gradient_rectangles:
        reference_rows[_round_coord(rectangle.y)].append(rectangle)

    solid_rectangles = [
        rectangle
        for rectangle in rectangles
        if "url(" not in rectangle.fill.casefold()
    ]
    modal_read_height = _mode_float([rectangle.height for rectangle in solid_rectangles])
    if not modal_read_height:
        modal_read_height = 10.0
    read_rows: dict[float, list[Rect]] = defaultdict(list)
    for rectangle in solid_rectangles:
        if abs(rectangle.height - modal_read_height) <= max(0.2, modal_read_height * 0.05):
            read_rows[_round_coord(rectangle.y)].append(rectangle)

    reference_y_values = sorted(reference_rows)
    panels: list[SvgPanel] = []
    for panel_index, reference_y in enumerate(reference_y_values, start=1):
        ref_rects = reference_rows[reference_y]
        reference_height = statistics.median(rectangle.height for rectangle in ref_rects)
        eligible_headers = [
            header
            for header in headers
            if 0 < reference_y - header.y <= max(40.0, reference_height * 2.5)
            and not any(
                0 < other_y - header.y < reference_y - header.y
                for other_y in reference_y_values
            )
        ]
        candidate_groups = _group_reference_rectangles(ref_rects)
        chosen_pairs: list[tuple[HeaderLabel, list[Rect]]] = []
        available = list(candidate_groups)
        for header in sorted(eligible_headers, key=lambda value: value.x):
            if not available:
                break
            group = min(
                available,
                key=lambda values: abs(
                    (values[0].x + values[-1].end) / 2.0 - header.x
                ),
            )
            midpoint = (group[0].x + group[-1].end) / 2.0
            tolerance = max(group[0].width, 20.0)
            if abs(midpoint - header.x) <= tolerance:
                chosen_pairs.append((header, group))
                available.remove(group)
        chosen_pairs.sort(key=lambda pair: pair[1][0].x)
        if not chosen_pairs:
            raise ToolError(f"reference row y={reference_y:g} could not be matched to unit labels")

        reference_glyph_map = _nearest_glyph_map(glyphs, reference_y, reference_height)
        all_reference_positions = sorted(
            x
            for x, text in reference_glyph_map.items()
            if str(text).strip().upper() in IUPAC_BASES
        )
        pitch = _infer_pitch(all_reference_positions)
        blocks: list[RepeatBlock] = []
        for block_index, (header, group) in enumerate(chosen_pairs):
            start = group[0].x
            end = group[-1].end
            unit_width = statistics.median(rectangle.width for rectangle in group)
            block_positions = [
                x for x in all_reference_positions if start <= x < end
            ]
            block_pitch = _infer_pitch(block_positions, pitch)
            motif_length = max(1, int(round(unit_width / block_pitch)))
            sequence_chars: list[str] = []
            for position in block_positions:
                char = str(reference_glyph_map.get(position, "")).strip().upper()
                sequence_chars.append(char if char in IUPAC_BASES else "N")
            scaffold = "".join(sequence_chars)
            expected_length = header.units * motif_length
            if len(scaffold) != expected_length:
                warnings.append(
                    f"panel_{panel_index}_component_{block_index + 1}_scaffold_length_"
                    f"{len(scaffold)}_expected_{expected_length}"
                )
                if not scaffold:
                    scaffold = "N" * expected_length
                elif len(scaffold) > expected_length:
                    scaffold = scaffold[:expected_length]
                    block_positions = block_positions[:expected_length]
            blocks.append(
                RepeatBlock(
                    index=block_index,
                    start=start,
                    end=end,
                    unit_width=unit_width,
                    declared_units=header.units,
                    fill=group[0].fill,
                    pitch=block_pitch,
                    motif_length=motif_length,
                    scaffold_positions=block_positions,
                    scaffold_sequence=scaffold,
                )
            )

        locus_start = min(block.start for block in blocks)
        locus_end = max(block.end for block in blocks)
        next_reference_y = (
            reference_y_values[panel_index]
            if panel_index < len(reference_y_values)
            else math.inf
        )
        reads: list[SvgRead] = []
        excluded = Counter()
        for row_y in sorted(read_rows):
            if row_y < reference_y + reference_height - 0.1 or row_y >= next_reference_y:
                continue
            row_rects = read_rows[row_y]
            row_pitch = statistics.median(block.pitch for block in blocks)
            full_spanning, reason = _coverage_reason(
                row_rects, locus_start, locus_end, row_pitch, flank_fills
            )
            row_glyph_map = _nearest_glyph_map(glyphs, row_y, modal_read_height)
            block_sequences: list[str] = []
            block_indel_reasons: list[str] = []
            row_centre = row_y + modal_read_height / 2.0
            for block in blocks:
                sequence_chars = list(block.scaffold_sequence)
                for index, position in enumerate(block.scaffold_positions[: len(sequence_chars)]):
                    observed = _lookup_glyph(
                        row_glyph_map, position, tolerance=max(0.1, block.pitch * 0.1)
                    ).strip().upper()
                    if observed in IUPAC_BASES:
                        sequence_chars[index] = observed
                block_sequences.append("".join(sequence_chars))
                has_deletion = any(
                    abs(line.y1 - line.y2) <= 0.2
                    and abs(line.y1 - row_centre) <= 0.3
                    and max(line.x1, line.x2) > block.start
                    and min(line.x1, line.x2) < block.end
                    for line in line_segments
                )
                # Deliberately include a vertical marker at either repeat/flank
                # boundary.  It is not a clean repeat-to-blue-flank transition,
                # and the SVG does not expose the inserted bases needed to call
                # the complete repeat sequence.
                has_insertion = any(
                    abs(line.x1 - line.x2) <= 0.2
                    and min(line.y1, line.y2) <= row_y + 0.3
                    and max(line.y1, line.y2) >= row_y + modal_read_height - 0.3
                    and block.start - 0.2 <= line.x1 <= block.end + 0.2
                    for line in line_segments
                )
                reasons = []
                if has_deletion:
                    reasons.append("target_deletion")
                if has_insertion:
                    reasons.append("target_insertion_bases_unavailable")
                block_indel_reasons.append(semicolon_join(reasons))

            intervals = _merge_intervals(row_rects, tolerance=max(0.2, row_pitch * 0.05))
            explicit = sorted(
                (
                    round((x - locus_start) / row_pitch, 3),
                    text.strip().upper(),
                )
                for x, text in row_glyph_map.items()
                if text.strip().upper() in IUPAC_BASES
            )
            signature_payload = {
                "intervals": [
                    [round((left - locus_start) / row_pitch, 3), round((right - locus_start) / row_pitch, 3)]
                    for left, right in intervals
                ],
                "changes": explicit,
                "blocks": block_sequences,
            }
            signature = hashlib.sha1(
                json.dumps(signature_payload, sort_keys=True, separators=(",", ":")).encode("utf-8")
            ).hexdigest()[:16]
            if not full_spanning:
                if any("target_deletion" in value for value in block_indel_reasons):
                    reason = "target_deletion"
                excluded[reason] += 1
            reads.append(
                SvgRead(
                    panel=panel_index,
                    row_y=row_y,
                    signature=signature,
                    block_sequences=block_sequences,
                    block_indel_reasons=block_indel_reasons,
                    full_spanning=full_spanning,
                    exclusion_reason=reason,
                )
            )

        panels.append(
            SvgPanel(
                index=panel_index,
                reference_y=reference_y,
                blocks=blocks,
                reads=reads,
                excluded_read_counts=dict(excluded),
            )
        )

    return ParsedSvg(panels=panels, warnings=warnings)


def parse_svg_identity(
    svg_path: Path, svg_root: Path, override: Mapping[str, str] | None = None
) -> tuple[str, str, str]:
    """Return ``(batch, sample_id, locus)`` using metadata then robust filename rules."""
    if override:
        sample = str(first_value(override, ["sampleid", "sample_id", "sample"]) or "").strip()
        locus = str(first_value(override, ["locus", "locusid", "locus_id", "gene"]) or "").strip()
        batch = str(first_value(override, ["batch", "batch_id", "cohort"]) or "").strip()
        if sample and locus:
            return batch or svg_path.parent.parent.name, sample, locus

    name = svg_path.name
    if name.lower().endswith(".gz"):
        name = name[:-3]
    if name.lower().endswith(".svg"):
        name = name[:-4]
    match = re.search(r"(?:\.|_)reviewer\.([^.]+)$", name, re.IGNORECASE)
    if not match:
        raise ToolError(
            "cannot confidently parse locus from filename; supply --file-map"
        )
    locus = match.group(1)
    prefix = name[: match.start()]
    prefix = re.sub(
        rf"[._](?:{re.escape(locus)})[._]eh_realigned$",
        "",
        prefix,
        flags=re.IGNORECASE,
    )
    prefix = re.sub(
        rf"_EH[^_]*_{re.escape(locus)}$",
        "",
        prefix,
        flags=re.IGNORECASE,
    )
    sample_from_filename = prefix.strip("._-")
    parent = svg_path.parent.name
    sample = parent if parent and parent != svg_root.name else sample_from_filename
    if not sample:
        raise ToolError(
            "cannot confidently parse sample ID from filename/path; supply --file-map"
        )
    try:
        relative = svg_path.relative_to(svg_root)
        parts = relative.parts
        batch = parts[-3] if len(parts) >= 3 else svg_root.name
    except ValueError:
        batch = svg_path.parent.parent.name
    return batch, sample, locus


def _genotype_key(sample: str, locus: str) -> str:
    return f"{normalise_key(sample)}|{normalise_key(locus)}"


def parse_genotype(value: Any) -> list[int]:
    text = str(value or "").strip()
    if not text or text.upper() in {"NA", "N/A", ".", "NONE"}:
        return []
    tokens = re.findall(r"\d+", text)
    return [int(token) for token in tokens[:2]]


def load_genotype_overrides(path: Path | None) -> dict[str, list[int]]:
    if path is None:
        return {}
    records = read_table_records(path)
    output: dict[str, list[int]] = {}
    for record in records:
        sample = first_value(record, ["sampleid", "sample_id", "sample"])
        locus = first_value(record, ["locus", "locusid", "locus_id", "gene"])
        genotype = parse_genotype(
            first_value(record, ["genotype", "eh_genotype", "repeat_genotype"])
        )
        if not genotype:
            left = safe_int(first_value(record, ["allele1_size", "allele_1_size"]))
            right = safe_int(first_value(record, ["allele2_size", "allele_2_size"]))
            genotype = [value for value in [left, right] if value is not None]
        if sample and locus and genotype:
            output[_genotype_key(str(sample), str(locus))] = genotype
    return output


def load_file_overrides(path: Path | None, svg_root: Path) -> dict[str, dict[str, str]]:
    if path is None:
        return {}
    records = read_table_records(path)
    output: dict[str, dict[str, str]] = {}
    for record in records:
        raw_path = first_value(record, ["svg_path", "path", "file", "filename"])
        if not raw_path:
            continue
        candidate = Path(str(raw_path))
        if not candidate.is_absolute():
            candidate = svg_root / candidate
        try:
            key = str(candidate.resolve())
        except OSError:
            key = str(candidate.absolute())
        output[key] = {str(k): str(v) for k, v in record.items() if v is not None}
    return output


def orient_sequence(sequence: str, orientation: str) -> str:
    if orientation == "reference":
        return sequence.upper()
    if orientation == "reverse_complement":
        return reverse_complement(sequence.upper())
    raise ToolError(f"unsupported reporting orientation {orientation!r}")


def transform_motif(motif: str, orientation: str) -> str:
    return orient_sequence(motif.upper(), orientation)


def _path_contains_protected_difference(
    source: Sequence[str], target: Sequence[str], protected: set[str]
) -> bool:
    return any(
        left != right and (left in protected or right in protected)
        for left, right in zip(source, target)
    )


def cluster_motif_paths(
    evidence: Sequence[ReadEvidence],
    config: RunConfig,
    protected_motifs: set[str] | None = None,
) -> tuple[list[Cluster], int]:
    """Cluster exact paths; optional abundance-aware near merging is opt-in."""
    callable_evidence = [item for item in evidence if item.callable and item.motifs]
    counter: Counter[tuple[str, ...]] = Counter(item.motifs for item in callable_evidence if item.motifs)
    first_seen: dict[tuple[str, ...], int] = {}
    for index, item in enumerate(callable_evidence):
        if item.motifs is not None:
            first_seen.setdefault(item.motifs, index)
    exact_count = len(counter)
    if not counter:
        return [], 0

    assignments = {path: path for path in counter}
    protected = protected_motifs or set()
    if config.max_near_motif_distance > 0:
        ordered = sorted(counter, key=lambda path: (-counter[path], first_seen[path], path))
        for source in reversed(ordered):
            candidates: list[tuple[int, int, tuple[str, ...]]] = []
            for target in ordered:
                if target == source or counter[target] <= counter[source]:
                    continue
                if counter[target] < counter[source] * config.near_cluster_parent_ratio:
                    continue
                distance = motif_distance(source, target)
                if distance > config.max_near_motif_distance:
                    continue
                if _path_contains_protected_difference(source, target, protected):
                    continue
                candidates.append((distance, -counter[target], target))
            if candidates:
                candidates.sort()
                best = candidates[0]
                tied = [item for item in candidates if item[:2] == best[:2]]
                if len(tied) == 1:
                    assignments[source] = best[2]

        for path in list(assignments):
            target = assignments[path]
            seen = {path}
            while assignments.get(target, target) != target and target not in seen:
                seen.add(target)
                target = assignments[target]
            assignments[path] = target

    grouped: dict[tuple[str, ...], list[tuple[str, ...]]] = defaultdict(list)
    for path, target in assignments.items():
        grouped[target].append(path)
    clusters = [
        Cluster(
            representative=target,
            count=sum(counter[path] for path in paths),
            first_order=min(first_seen[path] for path in paths),
            exact_paths=sorted(paths),
        )
        for target, paths in grouped.items()
    ]
    clusters.sort(key=lambda cluster: (-cluster.count, cluster.first_order, cluster.composition))
    return clusters, exact_count


def deduplicate_cross_panel_exact(
    evidence: Sequence[ReadEvidence],
) -> tuple[list[ReadEvidence], int]:
    """Keep max per-panel multiplicity for an identical full-row signature."""
    by_signature: dict[str, dict[int, list[ReadEvidence]]] = defaultdict(lambda: defaultdict(list))
    for item in evidence:
        by_signature[item.signature][item.panel].append(item)
    output: list[ReadEvidence] = []
    removed = 0
    for signature in sorted(by_signature):
        panel_groups = by_signature[signature]
        best_panel = min(
            panel_groups,
            key=lambda panel: (-len(panel_groups[panel]), panel),
        )
        kept = panel_groups[best_panel]
        output.extend(kept)
        removed += sum(len(items) for items in panel_groups.values()) - len(kept)
    output.sort(key=lambda item: (item.panel, item.row_y, item.signature))
    return output, removed


def _strong_clusters(
    clusters: Sequence[Cluster], total: int, config: RunConfig
) -> list[Cluster]:
    return [
        cluster
        for cluster in clusters
        if cluster.count >= config.min_cluster_reads
        or (
            config.strong_cluster_fraction > 0
            and total
            and cluster.count / total >= config.strong_cluster_fraction
        )
    ]


def _motif_classification(motifs: Sequence[str], benign: set[str], pathogenic: set[str]) -> str:
    observed = set(motifs)
    labels: list[str] = []
    if observed & pathogenic:
        labels.append("pathogenic_motif_present")
    if observed & benign:
        labels.append("benign_motif_present")
    return semicolon_join(labels) or "none_catalogued"


def _empty_main(sample: str) -> dict[str, Any]:
    return {
        "sampleID": sample,
        "allele1": "NA",
        "allele2": "NA",
        "allele1_size": "NA",
        "allele2_size": "NA",
        "read_depth_allele1": "NA",
        "read_depth_allele2": "NA",
        # Extended fields are written to all_loci.calls.csv.  Per-locus files
        # deliberately retain the seven-column workflow contract in MAIN_COLUMNS.
        "supporting_read_depth_allele1": "NA",
        "supporting_read_depth_allele2": "NA",
        "svg_read_rows_allele1": "NA",
        "svg_read_rows_allele2": "NA",
        "full_spanning_read_depth_allele1": "NA",
        "full_spanning_read_depth_allele2": "NA",
        "callable_full_spanning_read_depth_allele1": "NA",
        "callable_full_spanning_read_depth_allele2": "NA",
        "consensus_eligible_read_depth_allele1": "NA",
        "consensus_eligible_read_depth_allele2": "NA",
        "allele1_support_fraction": "NA",
        "allele2_support_fraction": "NA",
        "interruption_threshold_allele1": "NA",
        "interruption_threshold_allele2": "NA",
        "consensus_interruption_positions_allele1": "NA",
        "consensus_interruption_positions_allele2": "NA",
        "evidence_scope_allele1": "NA",
        "evidence_scope_allele2": "NA",
    }


def _iupac_compatible(left: str, right: str) -> bool:
    """Return whether two equal-length IUPAC strings can describe one sequence."""
    if len(left) != len(right):
        return False
    return all(
        bool(
            IUPAC_MATCH.get(a.upper(), frozenset(a.upper()))
            & IUPAC_MATCH.get(b.upper(), frozenset(b.upper()))
        )
        for a, b in zip(left, right)
    )


def _component_block_score(component: CatalogComponent, block: RepeatBlock) -> int | None:
    """Score a catalogue-to-SVG match, rejecting incompatible motif lengths.

    The reference scaffold can contain a biological interruption, so matching
    considers every complete unit instead of requiring only the first unit to
    equal the catalogue motif.
    """
    motif = component.motif.upper()
    if not motif or len(motif) != block.motif_length:
        return None
    units = split_motifs(block.scaffold_sequence.upper(), block.motif_length) or ()
    if not units:
        return None
    compatible = sum(_iupac_compatible(unit, motif) for unit in units)
    exact = sum(unit == motif for unit in units)
    if compatible == 0:
        return None
    # Compatibility dominates exactness; the latter provides deterministic
    # separation when an IUPAC catalogue motif matches several candidates.
    return compatible * 100 + exact * 10


def _align_catalog_components(
    catalog: CatalogLocus, blocks: Sequence[RepeatBlock]
) -> list[ComponentMapping]:
    """Align rendered blocks to ordered catalogue components.

    Catalogue components may be absent from an SVG when a ``*`` component has
    zero copies.  Exhaustive ordered alignment is tiny here (typically <=3
    components) and avoids the unsafe assumption that both index spaces match.
    """
    if not blocks:
        return []
    if len(blocks) > len(catalog.components):
        raise ToolError(
            f"SVG renders {len(blocks)} components but catalogue defines "
            f"{len(catalog.components)}"
        )

    candidates: list[tuple[int, tuple[int, ...]]] = []
    for catalog_indices in itertools.combinations(
        range(len(catalog.components)), len(blocks)
    ):
        retained = set(catalog_indices)
        omitted = [
            component
            for index, component in enumerate(catalog.components)
            if index not in retained
        ]
        if any(component.quantifier not in {"*", "?"} for component in omitted):
            continue
        scores: list[int] = []
        for svg_index, catalog_index in enumerate(catalog_indices):
            score = _component_block_score(
                catalog.components[catalog_index], blocks[svg_index]
            )
            if score is None:
                break
            scores.append(score)
        else:
            candidates.append((sum(scores), catalog_indices))
    if not candidates:
        rendered = ", ".join(
            f"{block.motif_length}bp:{block.scaffold_sequence[:block.motif_length]}"
            for block in blocks
        )
        expected = ", ".join(component.motif for component in catalog.components)
        raise ToolError(
            "catalogue components cannot be aligned to rendered SVG blocks "
            f"(catalogue={expected}; SVG={rendered})"
        )

    candidates.sort(key=lambda value: (-value[0], value[1]))
    best_score, best_indices = candidates[0]
    if len(candidates) > 1 and candidates[1][0] == best_score:
        raise ToolError(
            "catalogue-to-SVG component alignment is ambiguous; provide a "
            "single-component catalogue or inspect this locus manually"
        )
    return [
        ComponentMapping(catalog_index=catalog_index, svg_index=svg_index, score=score or 0)
        for svg_index, catalog_index in enumerate(best_indices)
        for score in [
            _component_block_score(catalog.components[catalog_index], blocks[svg_index])
        ]
    ]


def _select_component_mappings(
    catalog: CatalogLocus | None, parsed: ParsedSvg, mode: str
) -> list[ComponentMapping]:
    if not parsed.panels or not parsed.panels[0].blocks:
        return []
    reference_count = len(parsed.panels[0].blocks)
    if any(len(panel.blocks) != reference_count for panel in parsed.panels):
        raise ToolError("SVG panels render different numbers of repeat components")
    if catalog is None:
        mappings = [
            ComponentMapping(catalog_index=None, svg_index=index)
            for index in range(reference_count)
        ]
    else:
        mappings = _align_catalog_components(catalog, parsed.panels[0].blocks)
        expected_indices = [mapping.catalog_index for mapping in mappings]
        for panel in parsed.panels[1:]:
            panel_indices = [
                mapping.catalog_index
                for mapping in _align_catalog_components(catalog, panel.blocks)
            ]
            if panel_indices != expected_indices:
                raise ToolError(
                    "catalogue-to-SVG component mapping differs between panels"
                )
    if mode in {"all", "combined"}:
        return mappings
    if catalog is None:
        return mappings[:1]
    return [
        mapping for mapping in mappings if mapping.catalog_index == catalog.primary_index
    ]


def _reported_expected_motif(
    component: CatalogComponent | None,
    inferred_technical_motif: str,
    annotation: Annotation,
) -> tuple[str, str]:
    technical = (component.motif if component and component.motif else inferred_technical_motif).upper()
    reported = annotation.reported_motif or transform_motif(technical, annotation.orientation)
    return technical, reported.upper()


def _make_read_evidence(
    panels: Sequence[SvgPanel],
    svg_component_index: int,
    orientation: str,
    motif_length: int,
) -> tuple[list[ReadEvidence], list[dict[str, Any]]]:
    evidence: list[ReadEvidence] = []
    audit: list[dict[str, Any]] = []
    for panel in panels:
        if svg_component_index >= len(panel.blocks):
            continue
        block = panel.blocks[svg_component_index]
        for read in panel.reads:
            if not read.full_spanning:
                audit.append(
                    {
                        "panel": panel.index,
                        "panel_declared_size": block.declared_units,
                        "read_row_y": f"{read.row_y:g}",
                        "read_signature": read.signature,
                        "full_spanning": "no",
                        "technical_sequence": "NA",
                        "reported_sequence": "NA",
                        "motif_composition": "NA",
                        "n_repeat_units": "NA",
                        "callable": "no",
                        "exclusion_reason": read.exclusion_reason or "not_full_spanning",
                    }
                )
                continue
            technical = read.block_sequences[svg_component_index]
            reported = orient_sequence(technical, orientation)
            reason = read.block_indel_reasons[svg_component_index]
            motifs = None if reason else split_motifs(reported, motif_length)
            if not reason and any(base not in "ACGT" for base in reported):
                reason = "ambiguous_iupac_base_in_repeat"
                motifs = None
            elif not reason and motifs is None:
                reason = "sequence_not_divisible_by_motif_length"
            item = ReadEvidence(
                panel=panel.index,
                declared_size=block.declared_units,
                row_y=read.row_y,
                signature=read.signature,
                technical_sequence=technical,
                reported_sequence=reported,
                motifs=motifs,
                exclusion_reason=reason,
            )
            evidence.append(item)
            audit.append(
                {
                    "panel": panel.index,
                    "panel_declared_size": block.declared_units,
                    "read_row_y": f"{read.row_y:g}",
                    "read_signature": read.signature,
                    "full_spanning": "yes",
                    "technical_sequence": technical,
                    "reported_sequence": reported,
                    "motif_composition": rle_motifs(motifs or []),
                    "n_repeat_units": len(motifs) if motifs is not None else "NA",
                    "callable": "yes" if item.callable else "no",
                    "exclusion_reason": reason,
                }
            )
    return evidence, audit


def _infer_genotype(
    panels: Sequence[SvgPanel],
    svg_component_index: int,
    override: Sequence[int] | None,
) -> tuple[list[int], str, str, list[str]]:
    flags: list[str] = []
    panel_sizes = [
        panel.blocks[svg_component_index].declared_units
        for panel in panels
        if svg_component_index < len(panel.blocks)
    ]
    if override:
        genotype = list(override[:2])
        source = "override_table"
        if panel_sizes and sorted(genotype) != sorted(panel_sizes[: len(genotype)]):
            flags.append("genotype_override_svg_label_mismatch")
    else:
        genotype = panel_sizes[:2]
        source = "reviewer_svg_labels"
        flags.append("genotype_inferred_from_svg_labels")
    if not genotype:
        return genotype, source, "unknown", flags
    if len(genotype) == 1:
        return genotype, source, "hemizygous", flags
    if genotype[0] == genotype[1]:
        return genotype, source, "homozygous-size", flags
    return genotype, source, "heterozygous-size", flags


def _cluster_panel_sets(
    evidence_sets: Sequence[Sequence[ReadEvidence]],
    config: RunConfig,
    protected: set[str],
) -> tuple[list[list[Cluster]], int, int]:
    outputs: list[list[Cluster]] = []
    exact_total = 0
    strong_total = 0
    for evidence in evidence_sets:
        clusters, exact_count = cluster_motif_paths(evidence, config, protected)
        outputs.append(clusters)
        exact_total += exact_count
        strong_total += len(_strong_clusters(clusters, sum(c.count for c in clusters), config))
    return outputs, exact_total, strong_total


def _interruption_consensus(
    evidence: Sequence[ReadEvidence],
    expected_size: int,
    expected_motif: str,
    threshold: float,
) -> tuple[tuple[str, ...] | None, list[ReadEvidence], list[str], list[str]]:
    """Build a nucleotide-position consensus using a strict > threshold.

    Only callable full-spanning reads with the expected allele length enter the
    denominator.  At each nucleotide column, the most frequent specific
    non-reference base replaces a concrete reference base only when its own
    fraction is strictly greater than ``threshold``.  Ambiguous catalogue bases
    such as N are resolved to the dominant observed A/C/G/T base using the same
    threshold, so reported alleles contain observed nucleotides rather than
    catalogue wildcards. A tied or insufficiently dominant ambiguous base is not
    guessed and is returned for QC review.
    """
    eligible = [
        item
        for item in evidence
        if item.callable and item.motifs is not None and len(item.motifs) == expected_size
    ]
    if not eligible or expected_size < 1:
        return None, eligible, [], []
    reference_sequence = expected_motif * expected_size
    consensus_bases = list(reference_sequence)
    events: list[str] = []
    ties: list[str] = []
    denominator = len(eligible)
    motif_length = len(expected_motif)
    observed_sequences = ["".join(item.motifs or ()) for item in eligible]
    for base_index, reference_base in enumerate(reference_sequence):
        ambiguous_reference = reference_base not in "ACGT"
        if ambiguous_reference:
            counts = Counter(
                sequence[base_index]
                for sequence in observed_sequences
                if sequence[base_index] in "ACGT"
            )
        else:
            counts = Counter(
                sequence[base_index]
                for sequence in observed_sequences
                if sequence[base_index]
                not in IUPAC_MATCH.get(reference_base, frozenset(reference_base))
            )
        if not counts:
            if ambiguous_reference:
                unit_position = base_index // motif_length + 1
                within_unit = base_index % motif_length + 1
                ties.append(
                    f"unit{unit_position}.base{within_unit}:{reference_base}>"
                    "unresolved:no_concrete_base"
                )
            continue
        best_count = max(counts.values())
        winners = sorted(motif for motif, count in counts.items() if count == best_count)
        fraction = best_count / denominator
        unit_position = base_index // motif_length + 1
        within_unit = base_index % motif_length + 1
        if fraction > threshold and len(winners) == 1:
            base = winners[0]
            consensus_bases[base_index] = base
            events.append(
                f"unit{unit_position}.base{within_unit}:{reference_base}>{base}:"
                f"{best_count}/{denominator}={fraction:.4f}"
            )
        elif len(winners) > 1 and fraction > threshold:
            ties.append(
                f"unit{unit_position}.base{within_unit}:{reference_base}>"
                f"{'/'.join(winners)}:{best_count}/{denominator}={fraction:.4f}"
            )
        elif ambiguous_reference:
            ties.append(
                f"unit{unit_position}.base{within_unit}:{reference_base}>"
                f"unresolved_best_{'/'.join(winners)}:"
                f"{best_count}/{denominator}={fraction:.4f}"
            )
    consensus = split_motifs("".join(consensus_bases), motif_length)
    return consensus, eligible, events, ties


def _path_support(evidence: Sequence[ReadEvidence], path: Sequence[str]) -> int:
    return sum(
        item.motifs is not None
        and len(item.motifs) == len(path)
        and all(motif_matches(observed, called) for observed, called in zip(item.motifs, path))
        for item in evidence
    )


def _assign_reads_to_same_size_seeds(
    evidence: Sequence[ReadEvidence],
    seeds: Sequence[Cluster],
    max_distance: int,
) -> tuple[list[list[ReadEvidence]], list[ReadEvidence], list[ReadEvidence]]:
    """Assign each read to a unique nearest same-size allele seed.

    Reads farther than ``max_distance`` motif units from every seed are left
    unassigned.  Equal-distance ties are kept as ambiguous instead of being
    guessed.  These categories remain visible through the same-size QC fields.
    """
    groups: list[list[ReadEvidence]] = [[] for _ in seeds]
    unassigned: list[ReadEvidence] = []
    ambiguous: list[ReadEvidence] = []
    for item in evidence:
        if item.motifs is None or not seeds:
            unassigned.append(item)
            continue
        distances = [motif_distance(item.motifs, seed.representative) for seed in seeds]
        minimum = min(distances)
        if minimum > max_distance:
            unassigned.append(item)
            continue
        winners = [index for index, distance in enumerate(distances) if distance == minimum]
        if len(winners) != 1:
            ambiguous.append(item)
            continue
        groups[winners[0]].append(item)
    return groups, unassigned, ambiguous


def _interruption_burden(path: Sequence[str], expected_motif: str) -> int:
    """Count repeat units that do not match the catalogue reporting motif."""
    return sum(not motif_matches(motif, expected_motif) for motif in path)


def _assign_same_size_calls_to_panels(
    calls: Sequence[tuple[tuple[str, ...], list[ReadEvidence], list[str], list[str]]],
    evidence: Sequence[ReadEvidence],
    panel_ids: Sequence[int],
    config: RunConfig,
) -> dict[int, int]:
    """Map panel IDs to candidate indices only when panel evidence is decisive.

    Pooling can discover structures but cannot establish parental phase.  For
    SVG reporting order, require a uniquely best exact-matching candidate with
    at least the usual minimum support and a strict majority of the panel's
    eligible reads.  Mixed/tied panels do not acquire an arbitrary label.
    """
    assignments: dict[int, int] = {}
    for panel in panel_ids:
        panel_reads = [item for item in evidence if item.panel == panel]
        if not panel_reads or not calls:
            continue
        counts = [_path_support(panel_reads, call[0]) for call in calls]
        best = max(counts)
        winners = [index for index, count in enumerate(counts) if count == best]
        if (
            len(winners) == 1
            and best >= max(config.min_spanning_reads, config.same_size_min_cluster_reads)
            and best / len(panel_reads) > 0.50
        ):
            assignments[panel] = winners[0]
    return assignments


def _infer_alleles(
    sample: str,
    genotype_class: str,
    genotype: Sequence[int],
    evidence: Sequence[ReadEvidence],
    panel_declared_sizes: Mapping[int, int],
    expected_motif: str,
    benign: set[str],
    pathogenic: set[str],
    config: RunConfig,
    panel_total_rows: Mapping[int, int] | None = None,
    preserve_panel_order: bool = True,
) -> tuple[dict[str, Any], dict[str, Any], list[str], list[str]]:
    main = _empty_main(sample)
    flags: list[str] = []
    notes: list[str] = []
    protected = benign | pathogenic
    callable_evidence = [item for item in evidence if item.callable]
    raw_full_count = len(evidence)
    duplicates_removed = 0
    exact_count = 0
    cluster_count = 0
    strong_count = 0
    minor_count = 0
    same_size_audit: dict[str, Any] = {
        "same_size_candidate_cluster_count": "NA",
        "same_size_selected_cluster_count": "NA",
        "same_size_selected_cluster_coverage": "NA",
        "same_size_assigned_read_count": "NA",
        "same_size_unassigned_read_count": "NA",
        "same_size_ambiguous_read_count": "NA",
    }
    panel_total_rows = dict(panel_total_rows or {})
    full_by_panel = Counter(item.panel for item in evidence)
    callable_by_panel = Counter(item.panel for item in callable_evidence)
    for panel in panel_declared_sizes:
        panel_total_rows.setdefault(panel, full_by_panel.get(panel, 0))

    allele_paths: dict[int, tuple[str, ...] | None] = {1: None, 2: None}
    allele_support: dict[int, int | None] = {1: None, 2: None}
    allele_denominator: dict[int, int] = {1: 0, 2: 0}
    allele_events: dict[int, list[str]] = {1: [], 2: []}
    allele_thresholds: dict[int, float | None] = {1: None, 2: None}
    allele_metrics: dict[int, dict[str, Any]] = {
        1: {"scope": "NA", "svg_rows": "NA", "full": "NA", "callable": "NA", "eligible": "NA"},
        2: {"scope": "NA", "svg_rows": "NA", "full": "NA", "callable": "NA", "eligible": "NA"},
    }

    def set_panel_metrics(number: int, panel: int, eligible: int) -> None:
        allele_metrics[number] = {
            "scope": f"panel_{panel}_heterozygous_consensus",
            "svg_rows": panel_total_rows.get(panel, 0),
            "full": full_by_panel.get(panel, 0),
            "callable": callable_by_panel.get(panel, 0),
            "eligible": eligible,
        }

    def set_pooled_metrics(number: int, scope: str, eligible: int) -> None:
        allele_metrics[number] = {
            "scope": scope,
            "svg_rows": sum(panel_total_rows.values()),
            "full": raw_full_count - duplicates_removed,
            "callable": len(callable_evidence) - duplicates_removed,
            "eligible": eligible,
        }

    def register_call(
        number: int,
        path: tuple[str, ...] | None,
        eligible: Sequence[ReadEvidence],
        events: list[str],
        ties: list[str],
        threshold: float | None,
    ) -> None:
        allele_denominator[number] = len(eligible)
        allele_thresholds[number] = threshold
        if len(eligible) < config.min_spanning_reads:
            flags.append(f"allele{number}_insufficient_consensus_read_support")
            allele_paths[number] = None
            allele_events[number] = []
            allele_support[number] = None
            return
        allele_paths[number] = path
        allele_events[number] = events
        if path is not None:
            allele_support[number] = _path_support(eligible, path)
        if ties:
            flags.append(f"allele{number}_interruption_consensus_tie")
            notes.append(f"allele{number} tied interruption candidates: " + ",".join(ties))

    if raw_full_count < config.min_spanning_reads:
        flags.append("insufficient_full_spanning_read_support")
    if len(callable_evidence) < raw_full_count:
        flags.append("uncallable_full_spanning_reads_present")

    if genotype_class == "heterozygous-size":
        ordered_panels = sorted(
            panel_declared_sizes,
            key=(
                (lambda panel: panel)
                if preserve_panel_order
                else (lambda panel: (panel_declared_sizes.get(panel, math.inf), panel))
            ),
        )
        if len(ordered_panels) < 2:
            flags.append("heterozygous_panel_assignment_unresolved")
        panel_sets = [
            [item for item in callable_evidence if item.panel == panel]
            for panel in ordered_panels[:2]
        ]
        clusters_by_panel, exact_count, strong_count = _cluster_panel_sets(
            panel_sets, config, protected
        )
        cluster_count = sum(len(clusters) for clusters in clusters_by_panel)
        minor_count = sum(max(0, len(clusters) - 1) for clusters in clusters_by_panel)
        for clusters, values in zip(clusters_by_panel, panel_sets):
            if len(_strong_clusters(clusters, len(values), config)) > 1:
                flags.append("multiple_strong_clusters_within_heterozygous_panel")
        for number, panel in enumerate(ordered_panels[:2], start=1):
            size = panel_declared_sizes.get(panel, 0)
            panel_evidence = [item for item in callable_evidence if item.panel == panel]
            path, eligible, events, ties = _interruption_consensus(
                panel_evidence,
                size,
                expected_motif,
                config.heterozygous_interruption_min_fraction,
            )
            set_panel_metrics(number, panel, len(eligible))
            register_call(
                number, path, eligible, events, ties,
                config.heterozygous_interruption_min_fraction,
            )

    elif genotype_class == "homozygous-size":
        if len(panel_declared_sizes) < 2:
            flags.append("homozygous_expected_two_panels_missing")
        pooled = callable_evidence
        if config.homozygous_deduplication == "cross_panel_exact":
            pooled, duplicates_removed = deduplicate_cross_panel_exact(pooled)
            if duplicates_removed:
                flags.append("cross_panel_exact_rows_deduplicated")
        expected_size = genotype[0] if genotype else next(iter(panel_declared_sizes.values()), 0)
        eligible_pool = [
            item
            for item in pooled
            if item.callable and item.motifs is not None and len(item.motifs) == expected_size
        ]
        # Same-size allele discovery starts with exact paths.  Nearby minor paths
        # are assigned only after independently supported seeds have been chosen,
        # which prevents sequencing noise from becoming a third allele while
        # preserving two real same-length sequence alleles such as ATXN2.
        exact_config = replace(config, max_near_motif_distance=0)
        clusters, exact_count = cluster_motif_paths(
            eligible_pool, exact_config, protected
        )
        cluster_count = len(clusters)
        total_eligible = len(eligible_pool)
        candidates = [
            cluster
            for cluster in clusters
            if cluster.count >= config.same_size_min_cluster_reads
            and total_eligible > 0
            and cluster.count / total_eligible >= config.same_size_min_cluster_fraction
        ]
        strong_count = len(candidates)
        selected = candidates[: config.same_size_max_reported_clusters]
        groups, unassigned, ambiguous = _assign_reads_to_same_size_seeds(
            eligible_pool,
            selected,
            config.same_size_max_assignment_distance,
        )
        assigned_count = sum(len(group) for group in groups)
        coverage = assigned_count / total_eligible if total_eligible else 0.0
        minor_count = max(0, cluster_count - len(selected))
        same_size_audit.update(
            {
                "same_size_candidate_cluster_count": len(candidates),
                "same_size_selected_cluster_count": len(selected),
                "same_size_selected_cluster_coverage": (
                    f"{coverage:.4f}" if total_eligible else "NA"
                ),
                "same_size_assigned_read_count": assigned_count,
                "same_size_unassigned_read_count": len(unassigned),
                "same_size_ambiguous_read_count": len(ambiguous),
            }
        )
        notes.append(
            "same-size clustering: "
            f"eligible={total_eligible}, candidates={len(candidates)}, "
            f"selected={len(selected)}, assigned={assigned_count}, "
            f"unassigned={len(unassigned)}, ambiguous={len(ambiguous)}, "
            f"selected coverage={coverage:.4f}"
        )
        if len(candidates) > config.same_size_max_reported_clusters:
            flags.append("same_size_too_many_supported_clusters")
        if unassigned or ambiguous:
            flags.append("same_size_minor_or_unassigned_reads_present")

        cluster_calls: list[
            tuple[tuple[str, ...], list[ReadEvidence], list[str], list[str]]
        ] = []
        for group in groups:
            path, eligible, events, ties = _interruption_consensus(
                group,
                expected_size,
                expected_motif,
                config.same_size_within_cluster_consensus_min_fraction,
            )
            if path is not None:
                cluster_calls.append((path, eligible, events, ties))

        # Prefer SVG-panel order when supported.  Otherwise retain the original
        # deterministic pooled-call order, without withholding or extra REVIEW.
        cluster_calls.sort(
            key=lambda call: (
                -_interruption_burden(call[0], expected_motif),
                -len(call[1]),
                rle_motifs(call[0]),
            )
        )
        if len(cluster_calls) >= 2:
            if coverage < config.same_size_min_top_two_coverage:
                flags.append("same_size_top_two_coverage_below_threshold")
            if cluster_calls[0][0] == cluster_calls[1][0]:
                flags.append("same_size_selected_clusters_collapse_to_same_consensus")
            panel_ids = sorted(panel_declared_sizes)[:2]
            assignment = _assign_same_size_calls_to_panels(
                cluster_calls[:2], eligible_pool, panel_ids, config
            )
            identical = cluster_calls[0][0] == cluster_calls[1][0]
            panel_order_resolved = (
                not identical and len(assignment) == 2
                and len(set(assignment.values())) == 2
            )
            ordered_calls = (
                [cluster_calls[assignment[panel]] for panel in panel_ids]
                if panel_order_resolved else cluster_calls[:2]
            )
            if not identical and not panel_order_resolved:
                notes.append(
                    "Pooled same-size calls retained in deterministic reporting "
                    "order (interruption burden, support, then motif text); "
                    "their labels do not establish top/bottom SVG-panel assignment."
                )
            for number, (path, eligible, events, ties) in enumerate(ordered_calls, start=1):
                set_pooled_metrics(
                    number,
                    f"pooled_same_size_cluster_{number}_consensus" if not panel_order_resolved else
                    f"pooled_same_size_cluster_{number}_panel_{panel_ids[number - 1]}_consensus",
                    len(eligible),
                )
                register_call(
                    number,
                    path,
                    eligible,
                    events,
                    ties,
                    config.same_size_within_cluster_consensus_min_fraction,
                )
            flags.append("same_size_two_clusters_reported")
            if panel_order_resolved:
                flags.append("same_size_panel_order_from_read_support")
        elif len(cluster_calls) == 1:
            path, eligible, events, ties = cluster_calls[0]
            if coverage >= config.same_size_min_top_two_coverage:
                set_pooled_metrics(1, "pooled_same_size_cluster_1_consensus", len(eligible))
                register_call(
                    1, path, eligible, events, ties,
                    config.same_size_within_cluster_consensus_min_fraction,
                )
                set_pooled_metrics(2, "pooled_same_size_single_cluster_copy", len(eligible))
                register_call(
                    2,
                    path,
                    eligible,
                    events,
                    ties,
                    config.same_size_within_cluster_consensus_min_fraction,
                )
                flags.append("same_size_single_cluster_reported_for_both_alleles")
            else:
                flags.append("same_size_top_two_coverage_below_threshold")
                panel_ids = sorted(panel_declared_sizes)[:2]
                assignment = _assign_same_size_calls_to_panels(
                    cluster_calls, eligible_pool, panel_ids, config
                )
                for number in (1, 2):
                    set_pooled_metrics(number, "pooled_same_size_unresolved_allele", 0)
                if len(assignment) == 1:
                    panel = next(iter(assignment))
                    number = panel_ids.index(panel) + 1
                    set_pooled_metrics(
                        number, f"pooled_same_size_cluster_1_panel_{panel}_consensus", len(eligible)
                    )
                    register_call(
                        number, path, eligible, events, ties,
                        config.same_size_within_cluster_consensus_min_fraction,
                    )
                    flags.append("same_size_panel_order_from_read_support")
                else:
                    set_pooled_metrics(1, "pooled_same_size_cluster_1_consensus", len(eligible))
                    register_call(
                        1, path, eligible, events, ties,
                        config.same_size_within_cluster_consensus_min_fraction,
                    )
                    notes.append(
                        "The single supported pooled structure is retained as "
                        "allele1, following the original reporting convention; "
                        "its label does not establish SVG-panel assignment."
                    )
        else:
            for number in (1, 2):
                set_pooled_metrics(number, "pooled_same_size_no_supported_cluster", 0)
                register_call(number, None, [], [], [], None)
            flags.append("same_size_no_supported_cluster")

    elif genotype_class == "hemizygous":
        clusters, exact_count = cluster_motif_paths(callable_evidence, config, protected)
        cluster_count = len(clusters)
        strong_count = len(_strong_clusters(clusters, len(callable_evidence), config))
        minor_count = max(0, cluster_count - 1)
        expected_size = genotype[0] if genotype else next(iter(panel_declared_sizes.values()), 0)
        path, eligible, events, ties = _interruption_consensus(
            callable_evidence,
            expected_size,
            expected_motif,
            config.heterozygous_interruption_min_fraction,
        )
        set_pooled_metrics(1, "single_allele_consensus", len(eligible))
        register_call(
            1, path, eligible, events, ties,
            config.heterozygous_interruption_min_fraction,
        )
    else:
        flags.append("genotype_class_unresolved")

    if genotype_class == "heterozygous-size" and preserve_panel_order:
        expected_sizes = [
            panel_declared_sizes[panel]
            for panel in sorted(panel_declared_sizes)[:2]
        ]
    elif genotype_class == "heterozygous-size":
        expected_sizes = sorted(genotype)
    else:
        expected_sizes = list(genotype)
    for number in (1, 2):
        metrics = allele_metrics[number]
        main[f"svg_read_rows_allele{number}"] = metrics["svg_rows"]
        main[f"full_spanning_read_depth_allele{number}"] = metrics["full"]
        main[f"callable_full_spanning_read_depth_allele{number}"] = metrics["callable"]
        main[f"consensus_eligible_read_depth_allele{number}"] = metrics["eligible"]
        main[f"evidence_scope_allele{number}"] = metrics["scope"]
        threshold = allele_thresholds[number]
        if threshold is not None:
            main[f"interruption_threshold_allele{number}"] = f"{threshold:.4f}"
        main[f"consensus_interruption_positions_allele{number}"] = (
            semicolon_join(allele_events[number]) if allele_events[number] else "none"
        )
        path = allele_paths[number]
        if path is None:
            continue
        support = allele_support[number] or 0
        minimum_call_support = max(
            config.min_spanning_reads,
            config.same_size_min_cluster_reads
            if genotype_class == "homozygous-size"
            else config.min_cluster_reads,
        )
        if support < minimum_call_support:
            flags.append(
                f"allele{number}_same_size_low_cluster_support"
                if genotype_class == "homozygous-size"
                else f"allele{number}_low_read_support"
            )
            allele_paths[number] = None
            allele_events[number] = []
            main[f"consensus_interruption_positions_allele{number}"] = "none"
            continue
        main[f"allele{number}"] = rle_motifs(path)
        main[f"allele{number}_size"] = len(path)
        main[f"read_depth_allele{number}"] = support
        main[f"supporting_read_depth_allele{number}"] = support
        main[f"allele{number}_support_fraction"] = format_fraction(
            support, allele_denominator[number]
        )
        if number <= len(expected_sizes) and len(path) != expected_sizes[number - 1]:
            flags.append(f"allele{number}_reconstructed_size_differs_from_eh")

    resolved_allele_count = sum(path is not None for path in allele_paths.values())
    if genotype_class in {"heterozygous-size", "homozygous-size"} and resolved_allele_count == 1:
        flags.append("only_one_allele_reconstructed_when_two_expected")
    if resolved_allele_count == 0:
        flags.append("no_allele_structure_reconstructed")
    length_mismatch_count = sum(
        item.motifs is not None
        and genotype
        and len(item.motifs) not in set(genotype)
        for item in callable_evidence
    )
    if length_mismatch_count:
        flags.append("callable_read_length_differs_from_expected")
        notes.append(f"callable full-spanning rows excluded from consensus for length mismatch={length_mismatch_count}")

    details = {
        "n_full_spanning_reads_raw": raw_full_count,
        "n_total_full_spanning_reads": raw_full_count - duplicates_removed,
        "n_callable_reads": len(callable_evidence) - duplicates_removed,
        "n_cross_panel_duplicates_removed": duplicates_removed,
        "svg_read_rows_allele1": main["svg_read_rows_allele1"],
        "svg_read_rows_allele2": main["svg_read_rows_allele2"],
        "full_spanning_read_depth_allele1": main["full_spanning_read_depth_allele1"],
        "full_spanning_read_depth_allele2": main["full_spanning_read_depth_allele2"],
        "callable_full_spanning_read_depth_allele1": main["callable_full_spanning_read_depth_allele1"],
        "callable_full_spanning_read_depth_allele2": main["callable_full_spanning_read_depth_allele2"],
        "consensus_eligible_read_depth_allele1": main["consensus_eligible_read_depth_allele1"],
        "consensus_eligible_read_depth_allele2": main["consensus_eligible_read_depth_allele2"],
        "supporting_read_depth_allele1": main["supporting_read_depth_allele1"],
        "supporting_read_depth_allele2": main["supporting_read_depth_allele2"],
        "interruption_threshold_allele1": main["interruption_threshold_allele1"],
        "interruption_threshold_allele2": main["interruption_threshold_allele2"],
        "consensus_interruption_positions_allele1": main["consensus_interruption_positions_allele1"],
        "consensus_interruption_positions_allele2": main["consensus_interruption_positions_allele2"],
        "evidence_scope_allele1": main["evidence_scope_allele1"],
        "evidence_scope_allele2": main["evidence_scope_allele2"],
        "n_exact_motif_composition_clusters": exact_count,
        "n_motif_composition_clusters": cluster_count,
        "n_strong_motif_composition_clusters": strong_count,
        **same_size_audit,
        "allele1_support_fraction": main["allele1_support_fraction"],
        "allele2_support_fraction": main["allele2_support_fraction"],
        "minor_cluster_count": minor_count,
        "pure_or_interrupted_allele1": classify_structure(allele_paths[1] or [], expected_motif),
        "pure_or_interrupted_allele2": classify_structure(allele_paths[2] or [], expected_motif),
        "allele1_motif_classification": _motif_classification(allele_paths[1] or [], benign, pathogenic),
        "allele2_motif_classification": _motif_classification(allele_paths[2] or [], benign, pathogenic),
    }
    if not callable_evidence:
        notes.append("No callable full-spanning reads remained after sequence validation.")
    return main, details, flags, notes


INFORMATIONAL_QC_FLAGS = {
    "genotype_inferred_from_svg_labels",
    "orientation_inferred_from_gene_strand",
    "compound_primary_component_selected",
    "compound_components_combined",
    "compound_companion_components_from_catalog",
    "exact_clustering_used",
    "same_size_two_clusters_reported",
    "same_size_single_cluster_reported_for_both_alleles",
    "same_size_minor_or_unassigned_reads_present",
    "same_size_panel_order_from_read_support",
    # Excluded rows remain auditable in qc_flags/notes, but their mere presence
    # is not a reason for REVIEW when enough eligible reads support the calls.
    "uncallable_full_spanning_reads_present",
}


def _qc_status(flags: Sequence[str], error: bool = False) -> str:
    if error:
        return "ERROR"
    substantive = [flag for flag in flags if flag not in INFORMATIONAL_QC_FLAGS]
    return "REVIEW" if substantive else "PASS"


def _all_allele_rows_have_target_indel(
    main: Mapping[str, Any],
    genotype_class: str,
    parsed: ParsedSvg,
    svg_component_index: int,
    allele_number: int,
) -> bool:
    """Return true only when every relevant row for an unresolved allele has an indel.

    A deletion can make a row fail the geometric full-spanning test, so a row is
    relevant here when it either spans the repeat and both flanks or already has
    a target insertion/deletion annotation.  Ordinary gaps and missing-flank rows
    are not used to establish an all-row indel consensus.
    """
    if str(main.get(f"allele{allele_number}", "NA")) != "NA":
        return False

    scope = str(main.get(f"evidence_scope_allele{allele_number}", "NA"))
    if genotype_class == "heterozygous-size":
        match = re.match(r"panel_(\d+)_", scope)
        allowed_panels = {int(match.group(1))} if match else set()
    elif genotype_class == "homozygous-size":
        allowed_panels = {panel.index for panel in parsed.panels}
    elif genotype_class == "hemizygous":
        allowed_panels = (
            {panel.index for panel in parsed.panels} if allele_number == 1 else set()
        )
    else:
        allowed_panels = (
            {panel.index for panel in parsed.panels} if allele_number == 1 else set()
        )

    relevant_indel_states: list[bool] = []
    for panel in parsed.panels:
        if panel.index not in allowed_panels or svg_component_index >= len(panel.blocks):
            continue
        for read in panel.reads:
            reason_values = {
                value
                for value in read.block_indel_reasons[svg_component_index].split(";")
                if value
            }
            has_target_indel = bool(reason_values & TARGET_INDEL_REASONS)
            if read.full_spanning or has_target_indel:
                relevant_indel_states.append(has_target_indel)
    return bool(relevant_indel_states) and all(relevant_indel_states)


def _component_record(
    svg_path: Path,
    batch: str,
    sample: str,
    svg_locus: str,
    catalog: CatalogLocus | None,
    annotation: Annotation | None,
    parsed: ParsedSvg,
    catalog_component_index: int | None,
    svg_component_index: int,
    genotype_override: Sequence[int] | None,
    config: RunConfig,
    genotype_class_override: str | None = None,
    preserve_panel_order: bool = True,
) -> dict[str, Any]:
    flags: list[str] = []
    notes: list[str] = []
    component = (
        catalog.components[catalog_component_index]
        if catalog
        and catalog_component_index is not None
        and catalog_component_index < len(catalog.components)
        else None
    )
    component_number = (
        catalog_component_index + 1
        if catalog_component_index is not None
        else svg_component_index + 1
    )
    svg_component_number = svg_component_index + 1
    output_locus = component.variant_id if component else svg_locus
    parent_locus = catalog.locus_id if catalog else svg_locus
    main = _empty_main(sample)

    if catalog is None:
        flags.append("catalog_locus_not_found")
    elif len(catalog.components) != len(parsed.panels[0].blocks):
        flags.append("catalog_svg_component_count_mismatch")
        flags.append("catalog_svg_optional_component_omitted")
        notes.append(
            f"catalogue component {component_number} maps to SVG block "
            f"{svg_component_number} by ordered motif/scaffold matching"
        )
    elif len(catalog.components) > 1 and config.component_mode == "primary":
        flags.append("compound_primary_component_selected")
    if parsed.warnings:
        flags.append("svg_structure_warning")
        notes.extend(parsed.warnings)

    first_block = parsed.panels[0].blocks[svg_component_index]
    inferred_motif = first_block.scaffold_sequence[: first_block.motif_length]
    if not inferred_motif:
        inferred_motif = "N" * first_block.motif_length

    if annotation is None:
        flags.append("reporting_orientation_unresolved")
        if config.unresolved_strand_policy == "error":
            raise ToolError(
                f"no strand/reporting annotation for {svg_locus}; add it to the companion file"
            )
        if config.unresolved_strand_policy == "technical":
            annotation = Annotation(
                key=svg_locus,
                strand="",
                orientation="reference",
                reported_motif="",
                gene=catalog.gene if catalog else svg_locus,
                raw={},
            )
            flags.append("unresolved_orientation_reported_in_technical_orientation")
        else:
            panel_sizes = {
                panel.index: panel.blocks[svg_component_index].declared_units
                for panel in parsed.panels
                if svg_component_index < len(panel.blocks)
            }
            genotype, genotype_source, genotype_class, genotype_flags = _infer_genotype(
                parsed.panels, svg_component_index, genotype_override
            )
            flags.extend(genotype_flags)
            qc = {
                "batch": batch,
                "sampleID": sample,
                "locus": output_locus,
                "parent_locus": parent_locus,
                "component_index": component_number,
                "svg_component_index": svg_component_number,
                "source_svg": str(svg_path),
                "EH_genotype": "|".join(map(str, genotype)) or "NA",
                "genotype_source": genotype_source,
                "genotype_class": genotype_class,
                "EH_repeat_unit": component.motif if component else inferred_motif,
                "reported_repeat_unit": "NA",
                "gene": catalog.gene if catalog else svg_locus,
                "gene_strand": "NA",
                "reporting_orientation": "unresolved",
                "catalog_status": "resolved" if catalog else "missing",
                "annotation_status": "missing",
                "n_panels": len(panel_sizes),
                **{
                    key: "NA"
                    for key in QC_COLUMNS
                    if key.startswith("n_")
                    or key.startswith("same_size_")
                    or key.endswith("_fraction")
                    or key.startswith("minor_")
                    or key.startswith("pure_or_")
                    or key.startswith("allele1_motif")
                    or key.startswith("allele2_motif")
                },
                "qc_status": "REVIEW",
                "qc_flags": semicolon_join(flags),
                "notes": "No biological call emitted because the reporting strand is unresolved.",
            }
            return {
                "main": main,
                "combined": {
                    "batch": batch,
                    "locus": output_locus,
                    "parent_locus": parent_locus,
                    "component_index": component_number,
                    "svg_component_index": svg_component_number,
                    **main,
                    "source_svg": str(svg_path),
                },
                "qc": qc,
                "reads": [],
            }

    assert annotation is not None
    if annotation.raw and not first_value(
        annotation.raw, ["reporting_orientation", "biological_orientation"]
    ):
        flags.append("orientation_inferred_from_gene_strand")
    if not annotation.orientation:
        raise ToolError(f"reporting orientation remains unresolved for {svg_locus}")

    technical_motif, reported_motif = _reported_expected_motif(
        component, inferred_motif, annotation
    )
    motif_length = len(reported_motif)
    if motif_length != first_block.motif_length:
        flags.append("catalog_svg_motif_length_mismatch")
        notes.append(
            f"reported motif length {motif_length}; SVG unit length {first_block.motif_length}"
        )

    benign = {
        transform_motif(motif, annotation.orientation)
        for motif in (catalog.benign_motifs if catalog else [])
    }
    pathogenic = {
        transform_motif(motif, annotation.orientation)
        for motif in (catalog.pathogenic_motifs if catalog else [])
    }
    incompatible_known = sorted(
        motif for motif in benign | pathogenic if len(motif) != motif_length
    )
    if incompatible_known:
        flags.append("variable_length_known_motif_not_classifiable_from_svg")
        notes.append(
            "catalogued motif lengths differ from the selected repeat unit: "
            + ",".join(incompatible_known)
        )
    benign = {motif for motif in benign if len(motif) == motif_length}
    pathogenic = {motif for motif in pathogenic if len(motif) == motif_length}
    evidence, read_audit = _make_read_evidence(
        parsed.panels, svg_component_index, annotation.orientation, motif_length
    )
    panel_sizes = {
        panel.index: panel.blocks[svg_component_index].declared_units
        for panel in parsed.panels
        if svg_component_index < len(panel.blocks)
    }
    panel_total_rows = {
        panel.index: len(panel.reads)
        for panel in parsed.panels
        if svg_component_index < len(panel.blocks)
    }
    genotype, genotype_source, genotype_class, genotype_flags = _infer_genotype(
        parsed.panels, svg_component_index, genotype_override
    )
    if genotype_class_override:
        genotype_class = genotype_class_override
        genotype_source += "+compound_all_components"
        notes.append(
            "Allele-inference branch uses the complete SVG repeat-size vector "
            "in each panel, including blocks not selected for reporting: "
            + "; ".join(
                f"panel_{panel.index}="
                + ",".join(str(block.declared_units) for block in panel.blocks)
                for panel in parsed.panels
            )
        )
    flags.extend(genotype_flags)
    if len(panel_sizes) > 2:
        flags.append("more_than_two_reviewer_panels")
    if config.max_near_motif_distance == 0:
        flags.append("exact_clustering_used")

    inferred_main, details, inference_flags, inference_notes = _infer_alleles(
        sample=sample,
        genotype_class=genotype_class,
        genotype=genotype,
        evidence=evidence,
        panel_declared_sizes=panel_sizes,
        expected_motif=reported_motif,
        benign=benign,
        pathogenic=pathogenic,
        config=config,
        panel_total_rows=panel_total_rows,
        preserve_panel_order=preserve_panel_order,
    )
    main.update(inferred_main)
    flags.extend(inference_flags)
    notes.extend(inference_notes)
    for number in (1, 2):
        if _all_allele_rows_have_target_indel(
            main=main,
            genotype_class=genotype_class,
            parsed=parsed,
            svg_component_index=svg_component_index,
            allele_number=number,
        ):
            main[f"allele{number}"] = INDEL_SENTINEL
            main[f"allele{number}_size"] = "NA"
            flags.append(INDEL_REVIEW_FLAG)
    excluded_counts = Counter()
    for panel in parsed.panels:
        excluded_counts.update(panel.excluded_read_counts)
    # Parser-level counts cover rows that are not geometrically spanning. Add
    # full-spanning but sequence-uncallable rows (notably vertical insertions)
    # so biologists can see why a visually long row was not used.
    for row in read_audit:
        if str(row.get("full_spanning", "")).lower() != "yes":
            continue
        reason = str(row.get("exclusion_reason", "")).strip()
        for value in reason.split(";"):
            if value:
                excluded_counts[value] += 1
    if excluded_counts:
        notes.append(
            "excluded SVG rows: "
            + ", ".join(f"{key}={value}" for key, value in sorted(excluded_counts.items()))
        )

    for row in read_audit:
        row.update(
            {
                "batch": batch,
                "sampleID": sample,
                "locus": output_locus,
                "parent_locus": parent_locus,
                "component_index": component_number,
                "svg_component_index": svg_component_number,
            }
        )

    qc = {
        "batch": batch,
        "sampleID": sample,
        "locus": output_locus,
        "parent_locus": parent_locus,
        "component_index": component_number,
        "svg_component_index": svg_component_number,
        "locus_structure": catalog.locus_structure if catalog else "NA",
        "component_count": len(catalog.components) if catalog else len(parsed.panels[0].blocks),
        "allele1_components": f"{output_locus}:{main.get('allele1', 'NA')}",
        "allele2_components": f"{output_locus}:{main.get('allele2', 'NA')}",
        "source_svg": str(svg_path),
        "EH_genotype": "|".join(map(str, genotype)) or "NA",
        "genotype_source": genotype_source,
        "genotype_class": genotype_class,
        "EH_repeat_unit": technical_motif,
        "reported_repeat_unit": reported_motif,
        "gene": annotation.gene or (catalog.gene if catalog else svg_locus),
        "gene_strand": annotation.strand or "NA",
        "reporting_orientation": annotation.orientation,
        "catalog_status": "resolved" if catalog else "missing_svg_derived_model",
        "annotation_status": "resolved",
        "n_panels": len(panel_sizes),
        **details,
        "qc_status": "",
        "qc_flags": "",
        "notes": " | ".join(notes),
    }
    qc["qc_status"] = _qc_status(flags)
    qc["qc_flags"] = semicolon_join(flags)
    return {
        "main": main,
        "combined": {
            "batch": batch,
            "locus": output_locus,
            "parent_locus": parent_locus,
            "component_index": component_number,
            "svg_component_index": svg_component_number,
            **main,
            "source_svg": str(svg_path),
        },
        "qc": qc,
        "reads": read_audit,
    }


def _compound_genotype_class(
    parsed: ParsedSvg,
) -> str:
    """Compare every rendered repeat block, not only the reported component.

    A selected block with equal sizes (for example 9/9) must not cause pooling
    when a neighbouring repeat differs between panels (for example 4/27).
    Component selection changes reporting, not the full-locus size comparison.
    Equality here means homozygous *size*, not necessarily identical sequences.
    """
    vectors = [
        tuple(block.declared_units for block in panel.blocks)
        for panel in parsed.panels
    ]
    if len(vectors) == 1:
        return "hemizygous"
    if len(vectors) >= 2 and vectors[0] == vectors[1]:
        return "homozygous-size"
    if len(vectors) >= 2:
        return "heterozygous-size"
    return "unresolved"


def _row_matches_composition(row: Mapping[str, Any], composition: str) -> bool:
    observed = parse_rle_motifs(str(row.get("motif_composition", "")))
    called = parse_rle_motifs(composition)
    return bool(observed and called) and len(observed) == len(called) and all(
        motif_matches(value, pattern) for value, pattern in zip(observed, called)
    )


def _combine_compound_records(
    svg_path: Path,
    batch: str,
    sample: str,
    catalog: CatalogLocus,
    annotation: Annotation | None,
    parsed: ParsedSvg,
    mappings: Sequence[ComponentMapping],
    component_records: Sequence[dict[str, Any]],
    genotype_class: str,
    config: RunConfig,
) -> dict[str, Any]:
    """Assemble one catalogue-ordered allele call from every rendered component."""
    if len(mappings) != len(component_records):
        raise ToolError("compound component record count does not match component mapping")
    main = _empty_main(sample)
    record_by_catalog = {
        mapping.catalog_index: record
        for mapping, record in zip(mappings, component_records)
        if mapping.catalog_index is not None
    }
    mapping_by_catalog = {
        mapping.catalog_index: mapping
        for mapping in mappings
        if mapping.catalog_index is not None
    }
    primary_record = record_by_catalog.get(catalog.primary_index)
    if primary_record is None:
        raise ToolError("primary catalogue component is not rendered in compound SVG")
    panel_by_id = {panel.index: panel for panel in parsed.panels}
    panel_ids = sorted(panel_by_id)
    reporting_order = list(range(len(catalog.components)))
    if annotation and annotation.orientation == "reverse_complement":
        reporting_order.reverse()

    component_breakdowns: dict[int, list[str]] = {1: [], 2: []}
    component_compositions: dict[int, dict[int, str]] = {1: {}, 2: {}}
    combined_paths: dict[int, str] = {1: "", 2: ""}
    combined_sizes: dict[int, int] = {1: 0, 2: 0}
    component_unresolved: dict[int, bool] = {1: False, 2: False}
    for catalog_index in reporting_order:
        component = catalog.components[catalog_index]
        record = record_by_catalog.get(catalog_index)
        reported_label = transform_motif(
            component.motif,
            annotation.orientation if annotation else "reference",
        )
        for number in (1, 2):
            if record is None:
                component_breakdowns[number].append(f"{component.variant_id}:{reported_label}=0")
                component_compositions[number][catalog_index] = ""
                continue
            if catalog_index == catalog.primary_index:
                composition = str(record["main"].get(f"allele{number}", "NA"))
                size = safe_int(record["main"].get(f"allele{number}_size"))
            else:
                mapping = mapping_by_catalog[catalog_index]
                if genotype_class == "heterozygous-size":
                    panel_id = panel_ids[number - 1] if len(panel_ids) >= number else None
                elif genotype_class in {"homozygous-size", "hemizygous"}:
                    panel_id = panel_ids[0] if panel_ids and number == 1 else (
                        panel_ids[0] if panel_ids and genotype_class == "homozygous-size" else None
                    )
                else:
                    panel_id = None
                if panel_id is None:
                    composition, size = "NA", None
                else:
                    size = panel_by_id[panel_id].blocks[mapping.svg_index].declared_units
                    composition = rle_motifs([reported_label] * size) if size > 0 else "NA"
            component_breakdowns[number].append(
                f"{component.variant_id}:{composition}"
            )
            component_compositions[number][catalog_index] = composition
            if composition in {"", "NA", INDEL_SENTINEL} or size is None:
                component_unresolved[number] = True
            else:
                combined_paths[number] += composition
                combined_sizes[number] += size

    structure_tokens = _locus_structure_tokens(catalog.locus_structure)
    if annotation and annotation.orientation == "reverse_complement":
        structure_tokens = list(reversed(structure_tokens))
    for number in (1, 2):
        structured_parts: list[str] = []
        structured_breakdown: list[str] = []
        for token_type, value in structure_tokens:
            if token_type == "literal":
                literal = (
                    reverse_complement(str(value))
                    if annotation and annotation.orientation == "reverse_complement"
                    else str(value)
                )
                structured_parts.append(literal)
                structured_breakdown.append(f"FIXED:{literal}")
                continue
            catalog_index = int(value)
            composition = component_compositions[number].get(catalog_index, "")
            if composition and composition != INDEL_SENTINEL:
                structured_parts.append(composition)
            component = catalog.components[catalog_index]
            structured_breakdown.append(
                f"{component.variant_id}:{composition or transform_motif(component.motif, annotation.orientation if annotation else 'reference')}=0"
                if not composition
                else f"{component.variant_id}:{composition}"
            )
        main[f"allele{number}_components"] = ";".join(structured_breakdown)
        if not component_unresolved[number] and structured_parts:
            main[f"allele{number}"] = "".join(structured_parts)
            main[f"allele{number}_size"] = combined_sizes[number]
        if (
            str(primary_record["main"].get(f"allele{number}", "NA"))
            == INDEL_SENTINEL
        ):
            main[f"allele{number}"] = INDEL_SENTINEL
            main[f"allele{number}_size"] = "NA"

    row_maps: list[dict[tuple[int, str, str], dict[str, Any]]] = []
    for record in component_records:
        row_maps.append(
            {
                (
                    safe_int(row.get("panel")) or 0,
                    str(row.get("read_row_y", "")),
                    str(row.get("read_signature", "")),
                ): row
                for row in record.get("reads", [])
            }
        )
    base_keys = set(row_maps[0]) if row_maps else set()
    component_pairs = list(zip(mappings, component_records))
    component_events: dict[int, list[str]] = {1: [], 2: []}

    for number in (1, 2):
        if genotype_class == "heterozygous-size":
            allowed_panels = {panel_ids[number - 1]} if len(panel_ids) >= number else set()
            scope = (
                f"panel_{panel_ids[number - 1]}_compound_consensus"
                if allowed_panels else "NA"
            )
        elif genotype_class == "homozygous-size":
            allowed_panels = set(panel_ids)
            scope = f"pooled_same_size_compound_allele_{number}_consensus"
        else:
            allowed_panels = set(panel_ids)
            scope = "single_allele_compound_consensus" if number == 1 else "NA"
        scoped_keys = {key for key in base_keys if key[0] in allowed_panels}
        complete_keys: set[tuple[int, str, str]] = set()
        callable_keys: set[tuple[int, str, str]] = set()
        eligible_keys: set[tuple[int, str, str]] = set()
        supporting_keys: set[tuple[int, str, str]] = set()
        resolved_components = all(
            str(record["main"].get(f"allele{number}", "NA"))
            not in {"", "NA", INDEL_SENTINEL}
            for _, record in component_pairs
        )
        for key in scoped_keys:
            rows = [row_map.get(key) for row_map in row_maps]
            if not rows or any(row is None for row in rows):
                continue
            present_rows = [row for row in rows if row is not None]
            if all(str(row.get("full_spanning", "")).lower() == "yes" for row in present_rows):
                complete_keys.add(key)
            else:
                continue
            if all(str(row.get("callable", "")).lower() == "yes" for row in present_rows):
                callable_keys.add(key)
            else:
                continue
            lengths_match = all(
                safe_int(row.get("n_repeat_units"))
                == safe_int(record["main"].get(f"allele{number}_size"))
                for row, (_, record) in zip(present_rows, component_pairs)
            )
            if lengths_match:
                eligible_keys.add(key)
            else:
                continue
            if resolved_components and all(
                _row_matches_composition(
                    row,
                    str(record["main"].get(f"allele{number}", "NA")),
                )
                for row, (_, record) in zip(present_rows, component_pairs)
            ):
                supporting_keys.add(key)

        svg_rows = len(scoped_keys)
        full_depth = len(complete_keys)
        callable_depth = len(callable_keys)
        eligible_depth = len(eligible_keys)
        support_depth: Any = len(supporting_keys) if resolved_components else "NA"
        main[f"svg_read_rows_allele{number}"] = svg_rows
        main[f"full_spanning_read_depth_allele{number}"] = full_depth
        main[f"callable_full_spanning_read_depth_allele{number}"] = callable_depth
        main[f"consensus_eligible_read_depth_allele{number}"] = eligible_depth
        main[f"read_depth_allele{number}"] = support_depth
        main[f"supporting_read_depth_allele{number}"] = support_depth
        main[f"allele{number}_support_fraction"] = format_fraction(
            support_depth if isinstance(support_depth, int) else None,
            eligible_depth,
        )
        main[f"evidence_scope_allele{number}"] = scope
        thresholds = {
            str(record["main"].get(f"interruption_threshold_allele{number}", "NA"))
            for record in component_records
            if str(record["main"].get(f"interruption_threshold_allele{number}", "NA"))
            != "NA"
        }
        main[f"interruption_threshold_allele{number}"] = (
            sorted(thresholds)[0] if len(thresholds) == 1 else "NA"
        )
        for mapping, record in component_pairs:
            component = catalog.components[mapping.catalog_index or 0]
            event = str(
                record["main"].get(
                    f"consensus_interruption_positions_allele{number}", "none"
                )
            )
            if event not in {"", "none", "NA"}:
                component_events[number].append(f"{component.variant_id}[{event}]")
        main[f"consensus_interruption_positions_allele{number}"] = (
            semicolon_join(component_events[number]) if component_events[number] else "none"
        )

        # Interruption inference and support are defined by the catalogue primary
        # component.  Companion components contribute their JSON motif and SVG
        # declared copy number, preventing reference-scaffold bases (for example
        # ATXN7's terminal TCC) from being misreported as companion interruptions.
        for field in (
            "svg_read_rows",
            "full_spanning_read_depth",
            "callable_full_spanning_read_depth",
            "consensus_eligible_read_depth",
            "read_depth",
            "supporting_read_depth",
            "allele1_support_fraction" if number == 1 else "allele2_support_fraction",
        ):
            if field.startswith("allele"):
                main[field] = primary_record["main"].get(field, "NA")
            else:
                main[f"{field}_allele{number}"] = primary_record["main"].get(
                    f"{field}_allele{number}", "NA"
                )
        main[f"interruption_threshold_allele{number}"] = primary_record["main"].get(
            f"interruption_threshold_allele{number}", "NA"
        )
        main[f"consensus_interruption_positions_allele{number}"] = primary_record[
            "main"
        ].get(f"consensus_interruption_positions_allele{number}", "none")
        primary_scope = str(
            primary_record["main"].get(f"evidence_scope_allele{number}", "NA")
        )
        main[f"evidence_scope_allele{number}"] = (
            primary_scope.replace("_consensus", "_compound_primary_consensus")
            if primary_scope != "NA"
            else "NA"
        )

    flags = sorted(
        flag
        for flag in str(primary_record["qc"].get("qc_flags", "")).split(";")
        if flag and flag != "NA" and flag != "compound_primary_component_selected"
    )
    flags.append("compound_components_combined")
    flags.append("compound_companion_components_from_catalog")
    if any(component_unresolved.values()):
        flags.append("compound_component_unresolved")
    notes = []
    if str(primary_record["qc"].get("notes", "")).strip():
        notes.append(
            f"{primary_record['qc'].get('locus', 'primary component')}: "
            f"{primary_record['qc'].get('notes', '')}"
        )
    notes.append(
        "Companion components use catalogue motifs and SVG declared copy numbers; "
        "interruption consensus and supporting depth use the catalogue primary component."
    )
    component_genotypes = []
    for mapping, record in component_pairs:
        component = catalog.components[mapping.catalog_index or 0]
        component_genotypes.append(
            f"{component.variant_id}={record['qc'].get('EH_genotype', 'NA')}"
        )

    def combined_classification(field: str) -> str:
        return str(primary_record["qc"].get(field, "unresolved"))

    qc = {column: "NA" for column in QC_COLUMNS}
    qc.update(
        {
            "batch": batch,
            "sampleID": sample,
            "locus": catalog.locus_id,
            "parent_locus": catalog.locus_id,
            "component_index": "combined",
            "svg_component_index": "+".join(str(mapping.svg_index + 1) for mapping in mappings),
            "locus_structure": catalog.locus_structure,
            "component_count": len(catalog.components),
            "allele1_components": main.get("allele1_components", "NA"),
            "allele2_components": main.get("allele2_components", "NA"),
            "source_svg": str(svg_path),
            "EH_genotype": semicolon_join(component_genotypes),
            "genotype_source": "SVG panel labels across all catalogue components",
            "genotype_class": genotype_class,
            "EH_repeat_unit": "+".join(component.motif for component in catalog.components),
            "reported_repeat_unit": "+".join(
                transform_motif(
                    component.motif,
                    annotation.orientation if annotation else "reference",
                )
                for component in catalog.components
            ),
            "gene": (annotation.gene if annotation else "") or catalog.gene,
            "gene_strand": (annotation.strand if annotation else "") or "NA",
            "reporting_orientation": (
                annotation.orientation if annotation else "unresolved"
            ),
            "catalog_status": "resolved_compound",
            "annotation_status": "resolved" if annotation else "missing",
            "n_panels": len(parsed.panels),
            "n_full_spanning_reads_raw": main["full_spanning_read_depth_allele1"],
            "n_total_full_spanning_reads": main["full_spanning_read_depth_allele1"],
            "n_callable_reads": main["callable_full_spanning_read_depth_allele1"],
            "n_cross_panel_duplicates_removed": safe_int(
                primary_record["qc"].get("n_cross_panel_duplicates_removed")
            ) or 0,
            "n_exact_motif_composition_clusters": safe_int(
                primary_record["qc"].get("n_exact_motif_composition_clusters")
            ) or 0,
            "n_motif_composition_clusters": safe_int(
                primary_record["qc"].get("n_motif_composition_clusters")
            ) or 0,
            "n_strong_motif_composition_clusters": safe_int(
                primary_record["qc"].get("n_strong_motif_composition_clusters")
            ) or 0,
            "same_size_candidate_cluster_count": primary_record["qc"].get(
                "same_size_candidate_cluster_count", "NA"
            ),
            "same_size_selected_cluster_count": primary_record["qc"].get(
                "same_size_selected_cluster_count", "NA"
            ),
            "same_size_selected_cluster_coverage": primary_record["qc"].get(
                "same_size_selected_cluster_coverage", "NA"
            ),
            "same_size_assigned_read_count": primary_record["qc"].get(
                "same_size_assigned_read_count", "NA"
            ),
            "same_size_unassigned_read_count": primary_record["qc"].get(
                "same_size_unassigned_read_count", "NA"
            ),
            "same_size_ambiguous_read_count": primary_record["qc"].get(
                "same_size_ambiguous_read_count", "NA"
            ),
            "minor_cluster_count": safe_int(
                primary_record["qc"].get("minor_cluster_count")
            ) or 0,
            "pure_or_interrupted_allele1": combined_classification("pure_or_interrupted_allele1"),
            "pure_or_interrupted_allele2": combined_classification("pure_or_interrupted_allele2"),
            "allele1_motif_classification": primary_record["qc"].get(
                "allele1_motif_classification", "NA"
            ),
            "allele2_motif_classification": primary_record["qc"].get(
                "allele2_motif_classification", "NA"
            ),
            "qc_flags": semicolon_join(flags),
            "notes": " | ".join(notes),
        }
    )
    for number in (1, 2):
        for field in (
            "svg_read_rows",
            "full_spanning_read_depth",
            "callable_full_spanning_read_depth",
            "consensus_eligible_read_depth",
            "supporting_read_depth",
            "interruption_threshold",
            "consensus_interruption_positions",
            "evidence_scope",
        ):
            qc[f"{field}_allele{number}"] = main[f"{field}_allele{number}"]
        qc[f"allele{number}_support_fraction"] = main[f"allele{number}_support_fraction"]
    qc["qc_status"] = _qc_status(flags)
    main["locus_structure"] = catalog.locus_structure
    main["component_count"] = len(catalog.components)
    combined = {
        "batch": batch,
        "locus": catalog.locus_id,
        "parent_locus": catalog.locus_id,
        "component_index": "combined",
        "svg_component_index": "+".join(str(mapping.svg_index + 1) for mapping in mappings),
        **main,
        "source_svg": str(svg_path),
    }
    return {
        "main": main,
        "combined": combined,
        "qc": qc,
        "reads": [row for record in component_records for row in record.get("reads", [])],
    }


def _error_record(
    svg_path: Path,
    batch: str,
    sample: str,
    locus: str,
    message: str,
    trace: str = "",
) -> dict[str, Any]:
    main = _empty_main(sample)
    qc = {column: "NA" for column in QC_COLUMNS}
    qc.update(
        {
            "batch": batch,
            "sampleID": sample,
            "locus": locus,
            "parent_locus": locus,
            "component_index": "NA",
            "svg_component_index": "NA",
            "source_svg": str(svg_path),
            "catalog_status": "unknown",
            "annotation_status": "unknown",
            "qc_status": "ERROR",
            "qc_flags": "svg_processing_error",
            "notes": message + (f" | {trace}" if trace else ""),
        }
    )
    return {
        "main": main,
        "combined": {
            "batch": batch,
            "locus": locus,
            "parent_locus": locus,
            "component_index": "NA",
            "svg_component_index": "NA",
            **main,
            "source_svg": str(svg_path),
        },
        "qc": qc,
        "reads": [],
    }


def _worker_initialise(context: WorkerContext) -> None:
    global _WORKER_CONTEXT, _WORKER_CATALOG_INDEX, _WORKER_ANNOTATION_INDEX
    _WORKER_CONTEXT = context
    _WORKER_CATALOG_INDEX = build_catalog_index(context.catalog_loci)
    _WORKER_ANNOTATION_INDEX = build_annotation_index(context.annotation_rows)


def _selectable_locus_keys(
    svg_locus: str, catalog: CatalogLocus | None
) -> set[str]:
    keys = {normalise_key(svg_locus)}
    if catalog:
        keys.add(normalise_key(catalog.locus_id))
        keys.add(normalise_key(catalog.gene))
        keys.update(normalise_key(component.variant_id) for component in catalog.components)
    return {key for key in keys if key}


def _process_svg_worker(path_string: str) -> list[dict[str, Any]]:
    if _WORKER_CONTEXT is None:
        raise RuntimeError("worker context was not initialised")
    context = _WORKER_CONTEXT
    svg_path = Path(path_string)
    svg_root = Path(context.svg_root)
    override = context.file_overrides.get(str(svg_path.resolve()))
    batch, sample, svg_locus = "", svg_path.parent.name, "UNKNOWN"
    try:
        batch, sample, svg_locus = parse_svg_identity(svg_path, svg_root, override)
        selected = {normalise_key(value) for value in context.selected_loci}
        catalog = resolve_catalog_locus(svg_locus, _WORKER_CATALOG_INDEX)
        selectable = _selectable_locus_keys(svg_locus, catalog)
        if selected and not (selected & selectable):
            return []
        annotation = resolve_annotation(svg_locus, catalog, _WORKER_ANNOTATION_INDEX)
        parsed = parse_reviewer_svg(svg_path)
        explicitly_selected_components: set[int] = set()
        if catalog and selected:
            parent_keys = {
                normalise_key(svg_locus),
                normalise_key(catalog.locus_id),
                normalise_key(catalog.gene),
            }
            if not (selected & parent_keys):
                explicitly_selected_components = {
                    component.index
                    for component in catalog.components
                    if normalise_key(component.variant_id) in selected
                }
        primary_repeat_only = bool(
            catalog and catalog.locus_id.upper() in PRIMARY_REPEAT_ONLY_LOCI
        )
        mapping_mode = (
            "all"
            if explicitly_selected_components
            else "primary"
            if primary_repeat_only
            else context.config.component_mode
        )
        component_config = replace(context.config, component_mode=mapping_mode)
        component_mappings = _select_component_mappings(catalog, parsed, mapping_mode)
        if explicitly_selected_components:
            component_mappings = [
                mapping
                for mapping in component_mappings
                if mapping.catalog_index in explicitly_selected_components
            ]
        if not component_mappings:
            raise ToolError("no repeat component could be selected")
        combine_components = bool(
            catalog
            and len(catalog.components) > 1
            and mapping_mode == "combined"
            and not explicitly_selected_components
        )
        compound_genotype_class = (
            _compound_genotype_class(parsed)
            if (
                any(len(panel.blocks) > 1 for panel in parsed.panels)
                or bool(catalog and len(catalog.components) > 1)
            )
            else None
        )
        output: list[dict[str, Any]] = []
        for mapping in component_mappings:
            component_locus = (
                catalog.components[mapping.catalog_index].variant_id
                if catalog and mapping.catalog_index is not None
                else svg_locus
            )
            record_keys = _selectable_locus_keys(svg_locus, catalog)
            record_keys.add(normalise_key(component_locus))
            if selected and not (selected & record_keys):
                continue
            genotype_override = (
                context.genotype_overrides.get(_genotype_key(sample, component_locus))
                or context.genotype_overrides.get(_genotype_key(sample, svg_locus))
                or (
                    context.genotype_overrides.get(_genotype_key(sample, catalog.locus_id))
                    if catalog
                    else None
                )
            )
            output.append(
                _component_record(
                    svg_path=svg_path,
                    batch=batch,
                    sample=sample,
                    svg_locus=svg_locus,
                    catalog=catalog,
                    annotation=annotation,
                    parsed=parsed,
                    catalog_component_index=mapping.catalog_index,
                    svg_component_index=mapping.svg_index,
                    genotype_override=genotype_override,
                    config=component_config,
                    genotype_class_override=compound_genotype_class,
                    preserve_panel_order=True,
                )
            )
        if combine_components and catalog:
            return [
                _combine_compound_records(
                    svg_path=svg_path,
                    batch=batch,
                    sample=sample,
                    catalog=catalog,
                    annotation=annotation,
                    parsed=parsed,
                    mappings=component_mappings,
                    component_records=output,
                    genotype_class=compound_genotype_class or "unresolved",
                    config=component_config,
                )
            ]
        return output
    except Exception as exc:  # file-level fault isolation is deliberate
        trace = traceback.format_exc(limit=3).replace("\n", " ")
        return [_error_record(svg_path, batch, sample, svg_locus, str(exc), trace)]


def discover_svg_files(root: Path) -> list[Path]:
    files = list(root.rglob("*.svg")) + list(root.rglob("*.svg.gz"))
    unique = {str(path.resolve()): path for path in files if path.is_file()}
    return sorted(unique.values(), key=lambda path: natural_key(str(path.relative_to(root))))


def _write_csv_atomic(path: Path, rows: Sequence[Mapping[str, Any]], columns: Sequence[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    with temporary.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(columns), extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({column: row.get(column, "NA") for column in columns})
    os.replace(temporary, path)


def _write_json_atomic(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    with temporary.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, indent=2, sort_keys=True)
        handle.write("\n")
    os.replace(temporary, path)


def _add_qc_flag(row: dict[str, Any], flag: str) -> None:
    flags = set(str(row.get("qc_flags", "")).split(";")) - {"", "NA"}
    flags.add(flag)
    row["qc_flags"] = semicolon_join(flags)
    if row.get("qc_status") == "PASS":
        row["qc_status"] = "REVIEW"


def _mark_duplicate_keys(records: Sequence[dict[str, Any]]) -> list[tuple[str, str]]:
    key_counts = Counter(
        (
            str(record["combined"].get("batch", "")),
            str(record["main"].get("sampleID", "")),
            str(record["combined"].get("locus", "")),
        )
        for record in records
    )
    for record in records:
        key = (
            str(record["combined"].get("batch", "")),
            str(record["main"].get("sampleID", "")),
            str(record["combined"].get("locus", "")),
        )
        if key_counts[key] > 1:
            _add_qc_flag(record["qc"], "duplicate_batch_sample_locus_key")
    sample_locus_batches: dict[tuple[str, str], set[str]] = defaultdict(set)
    for record in records:
        sample_locus_batches[
            (
                str(record["main"].get("sampleID", "")),
                str(record["combined"].get("locus", "")),
            )
        ].add(str(record["combined"].get("batch", "")))
    conflicts = sorted(
        [key for key, batches in sample_locus_batches.items() if len(batches) > 1],
        key=lambda key: (natural_key(key[1]), natural_key(key[0])),
    )
    conflict_set = set(conflicts)
    for record in records:
        key = (
            str(record["main"].get("sampleID", "")),
            str(record["combined"].get("locus", "")),
        )
        if key in conflict_set:
            _add_qc_flag(record["qc"], "duplicate_sample_locus_across_batches")
    return conflicts


def build_input_audit(
    svg_files: Sequence[Path],
    svg_root: Path,
    catalog_loci: Sequence[CatalogLocus],
    annotations: Sequence[Annotation],
    file_overrides: Mapping[str, Mapping[str, str]],
) -> list[dict[str, Any]]:
    catalog_index = build_catalog_index(catalog_loci)
    annotation_index = build_annotation_index(annotations)
    loci = Counter()
    identity_failures = Counter()
    for path in svg_files:
        try:
            override = file_overrides.get(str(path.resolve()))
            _, _, locus = parse_svg_identity(path, svg_root, override)
            loci[locus] += 1
        except Exception as exc:
            identity_failures[str(exc)] += 1
    rows: list[dict[str, Any]] = []
    for svg_locus, count in sorted(loci.items(), key=lambda item: natural_key(item[0])):
        catalog = resolve_catalog_locus(svg_locus, catalog_index)
        annotation = resolve_annotation(svg_locus, catalog, annotation_index)
        status = "ready"
        notes: list[str] = []
        if catalog is None:
            status = "unconfigured"
            notes.append("catalog locus missing")
        if annotation is None:
            status = "unconfigured"
            notes.append("strand/reporting annotation missing")
        elif not annotation.orientation:
            status = "unconfigured"
            notes.append("reporting orientation is blank or invalid")
        rows.append(
            {
                "svg_locus": svg_locus,
                "n_svg_files": count,
                "catalog_locus": catalog.locus_id if catalog else "NA",
                "catalog_gene": catalog.gene if catalog and catalog.gene else "NA",
                "annotation_key": annotation.key if annotation else "NA",
                "strand": annotation.strand if annotation else "NA",
                "reporting_orientation": annotation.orientation if annotation else "NA",
                "primary_component": (
                    catalog.components[catalog.primary_index].variant_id if catalog else "NA"
                ),
                "n_catalog_components": len(catalog.components) if catalog else "NA",
                "status": status,
                "notes": "; ".join(notes),
            }
        )
    for message, count in identity_failures.items():
        rows.append(
            {
                "svg_locus": "UNPARSED",
                "n_svg_files": count,
                "catalog_locus": "NA",
                "catalog_gene": "NA",
                "annotation_key": "NA",
                "strand": "NA",
                "reporting_orientation": "NA",
                "primary_component": "NA",
                "n_catalog_components": "NA",
                "status": "unconfigured",
                "notes": message,
            }
        )
    return rows


def _serialise_config(config: RunConfig) -> dict[str, Any]:
    return asdict(config)


def process_corpus(
    svg_files: Sequence[Path],
    context: WorkerContext,
    output_dir: Path,
    workers: int,
    write_read_level: bool,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    records: list[dict[str, Any]] = []
    read_path = output_dir / "read_level" / "read_evidence.tsv.gz"
    read_temp = read_path.with_name(read_path.name + ".tmp")
    read_handle = None
    read_writer = None
    read_count = 0
    if write_read_level:
        read_path.parent.mkdir(parents=True, exist_ok=True)
        read_handle = gzip.open(read_temp, "wt", encoding="utf-8", newline="")
        read_writer = csv.DictWriter(
            read_handle, fieldnames=READ_COLUMNS, delimiter="\t", extrasaction="ignore"
        )
        read_writer.writeheader()

    paths = [str(path) for path in svg_files]
    logging.info("Processing %d SVG files with %d worker(s)", len(paths), workers)
    executor: ProcessPoolExecutor | None = None
    if workers <= 1:
        _worker_initialise(context)
        iterator = map(_process_svg_worker, paths)
    else:
        try:
            executor = ProcessPoolExecutor(
                max_workers=workers,
                initializer=_worker_initialise,
                initargs=(context,),
            )
            iterator = executor.map(_process_svg_worker, paths, chunksize=1)
        except (OSError, PermissionError) as exc:
            logging.warning("Process pool unavailable (%s); falling back to one worker", exc)
            executor = None
            _worker_initialise(context)
            iterator = map(_process_svg_worker, paths)

    started = datetime.now(timezone.utc)
    try:
        for index, file_records in enumerate(iterator, start=1):
            for record in file_records:
                if read_writer is not None:
                    for row in record.pop("reads", []):
                        read_writer.writerow({column: row.get(column, "NA") for column in READ_COLUMNS})
                        read_count += 1
                else:
                    record.pop("reads", None)
                records.append(record)
            if index == 1 or index % 100 == 0 or index == len(paths):
                logging.info("Processed %d/%d SVG files", index, len(paths))
    finally:
        if executor is not None:
            executor.shutdown(wait=True)
        if read_handle is not None:
            read_handle.close()
            os.replace(read_temp, read_path)

    conflicts = _mark_duplicate_keys(records)
    if conflicts and not context.config.allow_duplicate_sample_locus:
        preview = ", ".join(f"{sample}/{locus}" for sample, locus in conflicts[:10])
        raise ToolError(
            "duplicate sampleID/locus combinations occur across batches and the required "
            f"per-locus CSV has no batch column ({preview}). Resolve the duplicates or rerun "
            "with --allow-duplicate-sample-locus after confirming they are intentional."
        )
    records.sort(
        key=lambda record: (
            natural_key(record["combined"].get("locus", "")),
            natural_key(record["combined"].get("batch", "")),
            natural_key(record["main"].get("sampleID", "")),
        )
    )
    elapsed = (datetime.now(timezone.utc) - started).total_seconds()
    metrics = {
        "n_svg_files": len(svg_files),
        "n_component_records": len(records),
        "n_read_evidence_rows": read_count,
        "elapsed_seconds": round(elapsed, 3),
    }
    return records, metrics


def _stable_flag_colour(flag: str) -> str:
    digest = hashlib.sha256(flag.encode("utf-8")).digest()
    hue = int.from_bytes(digest[:2], "big") / 65535.0
    saturation = 0.48 + digest[2] / 255.0 * 0.18
    red, green, blue = colorsys.hls_to_rgb(hue, 0.82, saturation)
    return f"#{round(red * 255):02X}{round(green * 255):02X}{round(blue * 255):02X}"


def _unique_flag_colours(flags: Iterable[str]) -> dict[str, str]:
    """Return deterministic colours with no duplicate colour inside one result set."""
    mapping: dict[str, str] = {}
    used: set[str] = set()
    for flag in sorted(set(flags)):
        attempt = 0
        while True:
            colour_key = flag if attempt == 0 else f"{flag}\x1f{attempt}"
            colour = _stable_flag_colour(colour_key)
            if colour not in used:
                mapping[flag] = colour
                used.add(colour)
                break
            attempt += 1
    return mapping


def _flag_explanation(flag: str, qc: Mapping[str, Any], config: RunConfig) -> str:
    if flag == INDEL_REVIEW_FLAG:
        return INDEL_REVIEW_REASON
    if (
        flag == "variable_length_known_motif_not_classifiable_from_svg"
        and str(qc.get("parent_locus") or qc.get("locus") or "").upper() == "RFC1"
    ):
        return (
            "RFC1 contains a known repeat motif that this SVG cannot distinguish. "
            "Manual review is required."
        )
    match = re.fullmatch(r"allele([12])_same_size_low_cluster_support", flag)
    if match:
        number = match.group(1)
        observed = qc.get(f"supporting_read_depth_allele{number}", "NA")
        return (
            f"Only {observed} clean complete read(s) exactly match allele {number}; "
            f"at least {config.same_size_min_cluster_reads} are needed."
        )
    match = re.fullmatch(r"allele([12])_insufficient_consensus_read_support", flag)
    if match:
        number = match.group(1)
        observed = qc.get(f"consensus_eligible_read_depth_allele{number}", "NA")
        if observed in {None, "", "NA"}:
            return (
                f"Allele {number} has fewer than {config.min_spanning_reads} clean "
                "complete reads of the expected repeat length."
            )
        return (
            f"Allele {number} has {observed} clean complete read(s) of the expected "
            f"repeat length; at least {config.min_spanning_reads} are needed."
        )
    match = re.fullmatch(r"allele([12])_low_read_support", flag)
    if match:
        number = match.group(1)
        observed = qc.get(f"supporting_read_depth_allele{number}", "NA")
        if observed in {None, "", "NA"}:
            return (
                f"Allele {number} has fewer than {config.min_cluster_reads} clean "
                "complete reads matching the reported repeat pattern."
            )
        return (
            f"Only {observed} clean complete read(s) exactly match allele {number}; "
            f"at least {config.min_cluster_reads} are needed."
        )
    if flag == "same_size_top_two_coverage_below_threshold":
        assigned = safe_int(qc.get("same_size_assigned_read_count"))
        unassigned = safe_int(qc.get("same_size_unassigned_read_count"))
        ambiguous = safe_int(qc.get("same_size_ambiguous_read_count"))
        if assigned is not None and unassigned is not None and ambiguous is not None:
            total = assigned + unassigned + ambiguous
            required = math.ceil(config.same_size_min_top_two_coverage * total)
            observed_percent = round(100 * assigned / total) if total else 0
            required_percent = round(100 * config.same_size_min_top_two_coverage)
            return (
                f"The reported allele pattern(s) account for {assigned} of {total} clean "
                f"complete reads ({observed_percent}%). At least {required} of {total} "
                f"({required_percent}%) are required."
            )
        return "Too few clean complete reads could be grouped with the reported allele pattern(s)."
    explanations = {
        "insufficient_full_spanning_read_support": (
            f"Fewer than {config.min_spanning_reads} rows cleanly span the complete "
            "repeat and both flanks."
        ),
        "uncallable_full_spanning_reads_present": (
            "At least one visually spanning row contained a deletion, vertical insertion, "
            "or ambiguous sequence and was not used for an allele call."
        ),
        "no_allele_structure_reconstructed": "No allele repeat pattern could be determined from clean complete reads.",
        "only_one_allele_reconstructed_when_two_expected": "A repeat pattern could be determined for only one of the two expected alleles.",
        "compound_component_unresolved": "At least one required component of the compound locus is unresolved.",
        "catalog_locus_not_found": "The SVG locus is not configured in the supplied variant catalogue.",
        "reporting_orientation_unresolved": "Biological reporting orientation is unresolved.",
        "callable_read_length_differs_from_expected": "Callable rows with a repeat length different from the expected allele length were excluded from consensus.",
        "variable_length_known_motif_not_classifiable_from_svg": "A catalogued known motif has a different width and cannot be classified from this fixed-width SVG block.",
        "multiple_strong_clusters_within_heterozygous_panel": "More than one strongly supported complete motif path occurs within a heterozygous allele panel.",
        "same_size_too_many_supported_clusters": (
            f"{qc.get('same_size_candidate_cluster_count', 'NA')} different repeat "
            "patterns have enough read support, but only "
            f"{config.same_size_max_reported_clusters} can be reported as diploid alleles."
        ),
        "same_size_no_supported_cluster": (
            "No repeated sequence pattern was supported by at least "
            f"{config.same_size_min_cluster_reads} clean complete reads and "
            f"{round(100 * config.same_size_min_cluster_fraction)}% of the usable reads."
        ),
        "same_size_selected_clusters_collapse_to_same_consensus": (
            "Two distinct same-size seed clusters produced the same within-cluster "
            "consensus and require inspection."
        ),
        "svg_processing_error": "The SVG could not be processed; see QC notes for the exception.",
        "compound_components_combined": "All rendered catalogue components were combined into one locus-level allele.",
        "compound_companion_components_from_catalog": "Companion components use catalogue motifs and SVG-declared copy numbers; interruption consensus uses the primary component.",
        "genotype_inferred_from_svg_labels": "Repeat sizes were read from the SVG panel labels.",
        "orientation_inferred_from_gene_strand": "Reporting orientation was inferred from the companion gene-strand annotation.",
        "exact_clustering_used": "Identical clean complete repeat patterns were counted together.",
        "same_size_two_clusters_reported": "Two independently supported repeat patterns were reported as the two same-size alleles.",
        "same_size_single_cluster_reported_for_both_alleles": "One well-supported repeat pattern accounted for enough reads to be reported for both same-size alleles.",
        "same_size_minor_or_unassigned_reads_present": "Less common read patterns were grouped with one clearly similar allele pattern or retained as unmatched evidence.",
    }
    return explanations.get(flag, flag.replace("_", " ").capitalize() + ".")


def _humanise_qc_notes(value: Any) -> str:
    """Translate audit-oriented QC notes into concise language for biologists."""
    text_value = str(value or "").strip()
    if not text_value:
        return "No additional QC note."
    labels = {
        "gap_over_repeat_locus": "with a gap across the repeat",
        "missing_left_flank": "that did not fully cover the left flank",
        "missing_right_flank": "that did not fully cover the right flank",
        "missing_left_blue_flank": "without a clean left blue-flank transition",
        "missing_right_blue_flank": "without a clean right blue-flank transition",
        "target_deletion": "with a deletion through the repeat",
        "target_insertion_bases_unavailable": (
            "with a vertical insertion at or within the repeat boundary"
        ),
        "ambiguous_iupac_base_in_repeat": "with an ambiguous repeat base",
        "sequence_not_divisible_by_motif_length": (
            "whose repeat sequence could not be divided into complete repeat units"
        ),
        "not_full_spanning": "that did not span the complete repeat and both flanks",
        "no_aligned_sequence": "without aligned sequence across the target",
    }
    output: list[str] = []
    for raw_part in (part.strip() for part in text_value.split(" | ")):
        if not raw_part:
            continue
        if raw_part == "No callable full-spanning reads remained after sequence validation.":
            output.append(
                "No clean complete row could be used to determine an allele sequence."
            )
            continue
        match = re.search(
            r"same-size clustering: eligible=(\d+), candidates=(\d+), "
            r"selected=(\d+), assigned=(\d+), unassigned=(\d+), "
            r"ambiguous=(\d+), selected coverage=([0-9.]+)",
            raw_part,
        )
        if match:
            eligible, candidates, selected, assigned, unassigned, ambiguous = map(
                int, match.groups()[:6]
            )
            prefix = raw_part[: match.start()].rstrip(": ")
            sentence = (
                f"Of {eligible} clean complete reads with the expected repeat length, "
                f"{assigned} were grouped into {selected} supported allele "
                f"pattern{'s' if selected != 1 else ''}."
            )
            if candidates > selected:
                sentence += (
                    f" {candidates - selected} additional repeat pattern"
                    f"{'s had' if candidates - selected != 1 else ' had'} support but "
                    "could not be reported as another diploid allele."
                )
            if unassigned:
                sentence += f" {unassigned} read{'s' if unassigned != 1 else ''} did not fit either reported pattern."
            if ambiguous:
                sentence += f" {ambiguous} read{'s were' if ambiguous != 1 else ' was'} equally close to two patterns."
            output.append(f"{prefix}: {sentence}" if prefix else sentence)
            continue
        excluded_match = re.search(r"excluded SVG rows:\s*(.+)$", raw_part)
        if excluded_match:
            prefix = raw_part[: excluded_match.start()].rstrip(": ")
            readable: list[str] = []
            for item in excluded_match.group(1).split(","):
                key, separator, count_text = item.strip().partition("=")
                count = safe_int(count_text) if separator else None
                if count is None:
                    continue
                label = labels.get(key, key.replace("_", " "))
                readable.append(f"{count} {label}")
            sentence = "Rows not used: " + "; ".join(readable) + "."
            output.append(f"{prefix}: {sentence}" if prefix else sentence)
            continue
        length_match = re.search(
            r"callable full-spanning rows excluded from consensus for length mismatch=(\d+)",
            raw_part,
        )
        if length_match:
            count = int(length_match.group(1))
            output.append(
                f"{count} clean complete row{'s had' if count != 1 else ' had'} a repeat "
                "length different from the expected allele length and were not used."
            )
            continue
        output.append(raw_part)
    return " ".join(output) if output else "No additional QC note."


def _write_calls_workbook(
    path: Path,
    rows: Sequence[Mapping[str, Any]],
    columns: Sequence[str],
    flag_colours: Mapping[str, str],
    config: RunConfig,
) -> None:
    """Write an Excel review surface; CSV remains the machine-readable source."""
    try:
        from openpyxl import Workbook  # type: ignore
        from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore
        from openpyxl.worksheet.table import Table, TableStyleInfo  # type: ignore
        from openpyxl.utils import get_column_letter  # type: ignore
    except ImportError:
        logging.warning("openpyxl is unavailable; all_loci.calls.xlsx was not written")
        return

    workbook = Workbook()
    calls = workbook.active
    calls.title = "Calls"
    calls.sheet_view.showGridLines = False
    calls.freeze_panes = "A2"
    calls.append(list(columns))
    for row in rows:
        calls.append([row.get(column, "NA") for column in columns])
    header_fill = PatternFill("solid", fgColor="17365D")
    for cell in calls[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    calls.row_dimensions[1].height = 34
    if len(rows) > 0:
        table = Table(displayName="STRCalls", ref=calls.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        calls.add_table(table)
    else:
        calls.auto_filter.ref = calls.dimensions
    column_index = {name: index + 1 for index, name in enumerate(columns)}
    status_fills = {
        "PASS": "C6EFCE",
        "REVIEW": "FFEB9C",
        "ERROR": "FFC7CE",
    }
    for row_number, row in enumerate(rows, start=2):
        status = str(row.get("qc_status", ""))
        if "qc_status" in column_index and status in status_fills:
            calls.cell(row_number, column_index["qc_status"]).fill = PatternFill(
                "solid", fgColor=status_fills[status]
            )
        primary_flag = str(row.get("primary_qc_flag", ""))
        if "primary_qc_flag" in column_index and primary_flag in flag_colours:
            flag_fill = PatternFill("solid", fgColor=flag_colours[primary_flag].lstrip("#"))
            calls.cell(row_number, column_index["primary_qc_flag"]).fill = flag_fill
            if "review_reason" in column_index:
                calls.cell(row_number, column_index["review_reason"]).fill = flag_fill
        for name in ("review_reason", "qc_notes", "qc_flags", "allele1_components", "allele2_components"):
            if name in column_index:
                calls.cell(row_number, column_index[name]).alignment = Alignment(
                    vertical="top", wrap_text=True
                )
    width_overrides = {
        "allele1": 34,
        "allele2": 34,
        "allele1_components": 42,
        "allele2_components": 42,
        "qc_flags": 48,
        "review_reason": 70,
        "qc_notes": 60,
        "source_svg": 55,
    }
    for index, name in enumerate(columns, start=1):
        calls.column_dimensions[get_column_letter(index)].width = width_overrides.get(
            name, min(24, max(11, len(name) + 2))
        )

    summary_columns = [
        "batch",
        "sampleID",
        "locus",
        "locus_structure",
        "allele1",
        "allele2",
        "allele1_size",
        "allele2_size",
        "consensus_eligible_read_depth_allele1",
        "consensus_eligible_read_depth_allele2",
        "read_depth_allele1",
        "read_depth_allele2",
        "qc_status",
        "primary_qc_flag",
        "review_reason",
        "qc_notes",
    ]
    summary = workbook.create_sheet("Review Summary", 0)
    summary.sheet_view.showGridLines = False
    summary.freeze_panes = "A2"
    summary.append(summary_columns)
    for row in rows:
        summary.append(
            [
                _humanise_qc_notes(row.get(column, ""))
                if column == "qc_notes"
                else row.get(column, "NA")
                for column in summary_columns
            ]
        )
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    summary.row_dimensions[1].height = 34
    if rows:
        summary_table = Table(displayName="STRReviewSummary", ref=summary.dimensions)
        summary_table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        summary.add_table(summary_table)
    else:
        summary.auto_filter.ref = summary.dimensions
    summary_index = {name: index + 1 for index, name in enumerate(summary_columns)}
    for row_number, row in enumerate(rows, start=2):
        summary.row_dimensions[row_number].height = 54
        status = str(row.get("qc_status", ""))
        if status in status_fills:
            summary.cell(row_number, summary_index["qc_status"]).fill = PatternFill(
                "solid", fgColor=status_fills[status]
            )
        primary_flag = str(row.get("primary_qc_flag", ""))
        if primary_flag in flag_colours:
            flag_fill = PatternFill("solid", fgColor=flag_colours[primary_flag].lstrip("#"))
            summary.cell(row_number, summary_index["primary_qc_flag"]).fill = flag_fill
            summary.cell(row_number, summary_index["review_reason"]).fill = flag_fill
        for name in ("review_reason", "qc_notes"):
            summary.cell(row_number, summary_index[name]).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
    summary_widths = [
        32, 18, 18, 24, 34, 34, 12, 12, 18, 18, 16, 16, 12, 48, 78, 70
    ]
    for index, width in enumerate(summary_widths, start=1):
        summary.column_dimensions[get_column_letter(index)].width = width

    details = workbook.create_sheet("Flag Details")
    details.sheet_view.showGridLines = False
    detail_columns = ["batch", "sampleID", "locus", "qc_status", "flag", "explanation"]
    details.append(detail_columns)
    detail_rows: list[list[Any]] = []
    for row in rows:
        for flag in str(row.get("qc_flags", "")).split(";"):
            if flag and flag != "NA":
                detail_rows.append(
                    [
                        row.get("batch", ""),
                        row.get("sampleID", ""),
                        row.get("locus", ""),
                        row.get("qc_status", ""),
                        flag,
                        _flag_explanation(flag, row, config),
                    ]
                )
    for values in detail_rows:
        details.append(values)
    details.freeze_panes = "A2"
    details.auto_filter.ref = details.dimensions
    for cell in details[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    for row_number, values in enumerate(detail_rows, start=2):
        flag = str(values[4])
        details.cell(row_number, 5).fill = PatternFill(
            "solid", fgColor=flag_colours[flag].lstrip("#")
        )
        details.cell(row_number, 6).fill = PatternFill(
            "solid", fgColor=flag_colours[flag].lstrip("#")
        )
        details.cell(row_number, 6).alignment = Alignment(wrap_text=True, vertical="top")
    for index, width in enumerate([32, 18, 18, 12, 48, 80], start=1):
        details.column_dimensions[get_column_letter(index)].width = width

    legend = workbook.create_sheet("Flag Legend")
    legend.sheet_view.showGridLines = False
    legend.append(["flag", "type", "meaning"])
    for flag in sorted(flag_colours):
        legend.append(
            [
                flag,
                "informational" if flag in INFORMATIONAL_QC_FLAGS else "review/error",
                _flag_explanation(flag, {}, config),
            ]
        )
    legend.freeze_panes = "A2"
    legend.auto_filter.ref = legend.dimensions
    for cell in legend[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    for row_number in range(2, legend.max_row + 1):
        flag = str(legend.cell(row_number, 1).value)
        legend.cell(row_number, 1).fill = PatternFill(
            "solid", fgColor=flag_colours[flag].lstrip("#")
        )
        legend.cell(row_number, 3).fill = PatternFill(
            "solid", fgColor=flag_colours[flag].lstrip("#")
        )
        legend.cell(row_number, 3).alignment = Alignment(wrap_text=True, vertical="top")
    for index, width in enumerate([50, 18, 90], start=1):
        legend.column_dimensions[get_column_letter(index)].width = width

    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp.xlsx")
    workbook.save(temporary)
    try:
        _validate_xlsx_package(temporary)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    os.replace(temporary, path)


def _validate_xlsx_package(path: Path) -> None:
    """Validate ZIP/XML integrity and reject overlapping table/sheet filters.

    Excel Tables contain their own AutoFilter.  A worksheet-level AutoFilter on
    the same sheet is redundant and causes desktop Excel to report damaged
    workbook content even though permissive readers can still import the file.
    """
    required_parts = {
        "[Content_Types].xml",
        "_rels/.rels",
        "xl/workbook.xml",
        "xl/_rels/workbook.xml.rels",
        "xl/styles.xml",
    }
    spreadsheet_namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    try:
        with zipfile.ZipFile(path) as archive:
            corrupt_member = archive.testzip()
            if corrupt_member is not None:
                raise ToolError(f"XLSX ZIP member failed CRC validation: {corrupt_member}")
            names = set(archive.namelist())
            missing = sorted(required_parts - names)
            if missing:
                raise ToolError("XLSX package is missing required part(s): " + ", ".join(missing))
            for name in sorted(names):
                if not (name.endswith(".xml") or name.endswith(".rels")):
                    continue
                try:
                    root = ElementTree.fromstring(archive.read(name))
                except ElementTree.ParseError as exc:
                    raise ToolError(f"invalid XML in XLSX part {name}: {exc}") from exc
                if not name.startswith("xl/worksheets/sheet") or not name.endswith(".xml"):
                    continue
                has_sheet_filter = root.find(f"{spreadsheet_namespace}autoFilter") is not None
                has_tables = root.find(f"{spreadsheet_namespace}tableParts") is not None
                if has_sheet_filter and has_tables:
                    raise ToolError(
                        f"Excel-incompatible duplicate AutoFilter definitions in {name}"
                    )
    except zipfile.BadZipFile as exc:
        raise ToolError(f"invalid XLSX ZIP package: {exc}") from exc


def write_tabular_outputs(
    records: Sequence[dict[str, Any]],
    output_dir: Path,
    config: RunConfig,
) -> dict[str, Any]:
    combined_columns = [
        "batch",
        "sampleID",
        "locus",
        "parent_locus",
        "component_index",
        "svg_component_index",
        "locus_structure",
        "component_count",
        "allele1_components",
        "allele2_components",
        "allele1",
        "allele2",
        "allele1_size",
        "allele2_size",
        "read_depth_allele1",
        "read_depth_allele2",
        "supporting_read_depth_allele1",
        "supporting_read_depth_allele2",
        "svg_read_rows_allele1",
        "svg_read_rows_allele2",
        "full_spanning_read_depth_allele1",
        "full_spanning_read_depth_allele2",
        "callable_full_spanning_read_depth_allele1",
        "callable_full_spanning_read_depth_allele2",
        "consensus_eligible_read_depth_allele1",
        "consensus_eligible_read_depth_allele2",
        "interruption_threshold_allele1",
        "interruption_threshold_allele2",
        "consensus_interruption_positions_allele1",
        "consensus_interruption_positions_allele2",
        "allele1_support_fraction",
        "allele2_support_fraction",
        "evidence_scope_allele1",
        "evidence_scope_allele2",
        "same_size_candidate_cluster_count",
        "same_size_selected_cluster_count",
        "same_size_selected_cluster_coverage",
        "same_size_assigned_read_count",
        "same_size_unassigned_read_count",
        "same_size_ambiguous_read_count",
        "qc_status",
        "primary_qc_flag",
        "qc_flags",
        "review_reason",
        "qc_notes",
        "source_svg",
    ]
    combined = [dict(record["combined"]) for record in records]
    qc_rows = [dict(record["qc"]) for record in records]
    all_flags = sorted(
        {
            flag
            for qc in qc_rows
            for flag in str(qc.get("qc_flags", "")).split(";")
            if flag and flag != "NA"
        }
    )
    flag_colours = _unique_flag_colours(all_flags)
    for row, qc in zip(combined, qc_rows):
        flags = [
            flag
            for flag in str(qc.get("qc_flags", "")).split(";")
            if flag and flag != "NA"
        ]
        substantive = [flag for flag in flags if flag not in INFORMATIONAL_QC_FLAGS]
        primary_flag = (substantive or flags or ["none"])[0]
        reasons = [_flag_explanation(flag, qc, config) for flag in substantive]
        if INDEL_REVIEW_FLAG in substantive:
            primary_flag = INDEL_REVIEW_FLAG
            reasons = [INDEL_REVIEW_REASON]
        row.update(
            {
                "locus_structure": qc.get("locus_structure", row.get("locus_structure", "NA")),
                "component_count": qc.get("component_count", row.get("component_count", "NA")),
                "allele1_components": qc.get("allele1_components", row.get("allele1_components", "NA")),
                "allele2_components": qc.get("allele2_components", row.get("allele2_components", "NA")),
                "same_size_candidate_cluster_count": qc.get(
                    "same_size_candidate_cluster_count", "NA"
                ),
                "same_size_selected_cluster_count": qc.get(
                    "same_size_selected_cluster_count", "NA"
                ),
                "same_size_selected_cluster_coverage": qc.get(
                    "same_size_selected_cluster_coverage", "NA"
                ),
                "same_size_assigned_read_count": qc.get(
                    "same_size_assigned_read_count", "NA"
                ),
                "same_size_unassigned_read_count": qc.get(
                    "same_size_unassigned_read_count", "NA"
                ),
                "same_size_ambiguous_read_count": qc.get(
                    "same_size_ambiguous_read_count", "NA"
                ),
                "qc_status": qc.get("qc_status", "NA"),
                "primary_qc_flag": primary_flag,
                "qc_flags": semicolon_join(flags),
                "review_reason": " | ".join(reasons) if reasons else "No manual-review flag.",
                "qc_notes": qc.get("notes", ""),
            }
        )
    _write_csv_atomic(output_dir / "all_loci.calls.csv", combined, combined_columns)
    _write_csv_atomic(output_dir / "all_loci.qc.csv", qc_rows, QC_COLUMNS)
    _write_calls_workbook(
        output_dir / "all_loci.calls.xlsx",
        combined,
        combined_columns,
        flag_colours,
        config,
    )

    by_locus: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        by_locus[str(record["combined"].get("locus", "UNKNOWN"))].append(record["main"])
    for locus, rows in sorted(by_locus.items(), key=lambda item: natural_key(item[0])):
        rows.sort(key=lambda row: natural_key(row.get("sampleID", "")))
        filename = f"{sanitise_filename(locus)}.calls.csv"
        _write_csv_atomic(output_dir / "calls_by_locus" / filename, rows, MAIN_COLUMNS)

    status_counts = Counter(str(row.get("qc_status", "UNKNOWN")) for row in qc_rows)
    flag_counts = Counter(
        flag
        for row in qc_rows
        for flag in str(row.get("qc_flags", "")).split(";")
        if flag and flag != "NA"
    )
    call_count = sum(
        row.get("allele1") not in {None, "", "NA", INDEL_SENTINEL}
        for row in combined
    )
    return {
        "n_loci": len(by_locus),
        "n_calls_with_allele1": call_count,
        "qc_status_counts": dict(sorted(status_counts.items())),
        "qc_flag_counts": dict(flag_counts.most_common()),
    }


FIGURE_PALETTE = {
    "pure": "#4477AA",
    "interrupted": "#EE7733",
    "unresolved": "#BBBBBB",
    "benign": "#228833",
    "pathogenic": "#CC3311",
    "novel": "#AA3377",
    "grid": "#D9D9D9",
    "text": "#222222",
}


def _wilson_interval(successes: int, total: int, z: float = 1.96) -> tuple[float, float, float]:
    if total <= 0:
        return 0.0, 0.0, 0.0
    proportion = successes / total
    denominator = 1.0 + z * z / total
    centre = (proportion + z * z / (2 * total)) / denominator
    margin = (
        z
        * math.sqrt(proportion * (1 - proportion) / total + z * z / (4 * total * total))
        / denominator
    )
    # In exact arithmetic a Wilson interval always contains the observed
    # proportion.  At cohort-sized denominators, floating-point cancellation
    # can nevertheless put a boundary a few ulps on the wrong side of the
    # point estimate (for example, upper=0.9999999999999999 when p=1.0).
    # Matplotlib then sees a tiny negative x-error and refuses to draw the
    # entire figure.  Clamp the bounds to both the probability range and the
    # point estimate so downstream error widths are guaranteed non-negative.
    lower = min(proportion, max(0.0, centre - margin))
    upper = max(proportion, min(1.0, centre + margin))
    return proportion, lower, upper


def build_figure_data(
    records: Sequence[dict[str, Any]], catalog_loci: Sequence[CatalogLocus]
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    catalog_by_parent = {normalise_key(locus.locus_id): locus for locus in catalog_loci}
    alleles: list[dict[str, Any]] = []
    locus_counts: dict[str, Counter[str]] = defaultdict(Counter)
    locus_samples: dict[str, set[str]] = defaultdict(set)
    locus_depths: dict[str, list[int]] = defaultdict(list)
    motif_counts: Counter[tuple[str, str, str]] = Counter()

    for record in records:
        main = record["main"]
        qc = record["qc"]
        locus = str(qc.get("locus", "UNKNOWN"))
        parent = str(qc.get("parent_locus", locus))
        locus_samples[locus].add(str(main.get("sampleID", "")))
        expected = str(qc.get("reported_repeat_unit", ""))
        genotype_class = str(qc.get("genotype_class", ""))
        expected_alleles = 1 if genotype_class == "hemizygous" else 2
        catalog = catalog_by_parent.get(normalise_key(parent))
        benign = {
            transform_motif(motif, str(qc.get("reporting_orientation", "reference")))
            for motif in (catalog.benign_motifs if catalog else [])
        }
        pathogenic = {
            transform_motif(motif, str(qc.get("reporting_orientation", "reference")))
            for motif in (catalog.pathogenic_motifs if catalog else [])
        }
        for number in range(1, expected_alleles + 1):
            composition = str(main.get(f"allele{number}", "NA"))
            size = safe_int(main.get(f"allele{number}_size"))
            depth = safe_int(main.get(f"read_depth_allele{number}"))
            analysis_composition = composition
            analysis_expected = expected
            if catalog and len(catalog.components) > 1:
                primary = catalog.components[catalog.primary_index]
                analysis_expected = transform_motif(
                    primary.motif,
                    str(qc.get("reporting_orientation", "reference")),
                )
                component_text = str(main.get(f"allele{number}_components", ""))
                prefix = primary.variant_id + ":"
                for part in component_text.split(";"):
                    if part.startswith(prefix):
                        analysis_composition = part[len(prefix) :]
                        break
            motifs = parse_rle_motifs(analysis_composition)
            if catalog and len(catalog.components) > 1 and motifs:
                size = len(motifs)
            category = str(
                qc.get(
                    f"pure_or_interrupted_allele{number}",
                    classify_structure(motifs, analysis_expected) if motifs else "unresolved",
                )
            )
            locus_counts[locus][category] += 1
            if depth is not None:
                locus_depths[locus].append(depth)
            alleles.append(
                {
                    "batch": qc.get("batch", ""),
                    "sampleID": main.get("sampleID", ""),
                    "locus": locus,
                    "allele": number,
                    "composition": composition,
                    # Keep the component selected for biological reporting as a
                    # separate plotting field.  For simple loci this is identical
                    # to composition; for compound loci it prevents a flanking
                    # repeat block from being mixed into the frequency figure.
                    "analysis_composition": analysis_composition,
                    "size": size if size is not None else "NA",
                    "read_depth": depth if depth is not None else "NA",
                    "structure_class": category,
                    "reported_repeat_unit": analysis_expected,
                    "normal_max": catalog.normal_max if catalog and catalog.normal_max is not None else "NA",
                    "pathogenic_min": catalog.pathogenic_min if catalog and catalog.pathogenic_min is not None else "NA",
                    "qc_status": qc.get("qc_status", ""),
                }
            )
            for motif in motifs:
                if motif_matches(motif, analysis_expected):
                    continue
                motif_class = (
                    "pathogenic" if motif in pathogenic else "benign" if motif in benign else "other"
                )
                motif_counts[(locus, motif, motif_class)] += 1

    summaries: list[dict[str, Any]] = []
    for locus in sorted(locus_counts, key=natural_key):
        counts = locus_counts[locus]
        resolved = counts["pure"] + counts["interrupted"]
        prevalence, lower, upper = _wilson_interval(counts["interrupted"], resolved)
        depths = locus_depths[locus]
        summaries.append(
            {
                "locus": locus,
                "n_samples": len(locus_samples[locus]),
                "n_expected_alleles": sum(counts.values()),
                "n_resolved_alleles": resolved,
                "n_pure_alleles": counts["pure"],
                "n_interrupted_alleles": counts["interrupted"],
                "n_unresolved_alleles": counts["unresolved"],
                "call_rate": resolved / sum(counts.values()) if sum(counts.values()) else 0.0,
                "interruption_prevalence": prevalence,
                "interruption_ci95_lower": lower,
                "interruption_ci95_upper": upper,
                "median_reported_cluster_depth": statistics.median(depths) if depths else "NA",
            }
        )
    motif_rows = [
        {"locus": locus, "motif": motif, "classification": classification, "count": count}
        for (locus, motif, classification), count in sorted(
            motif_counts.items(), key=lambda item: (-item[1], natural_key(item[0][0]), item[0][1])
        )
    ]
    return alleles, summaries, motif_rows


def build_sequence_composition_frequency_data(
    alleles: Sequence[dict[str, Any]],
    max_display_rows: int = SEQUENCE_COMPOSITION_FIGURE_MAX_ROWS,
) -> list[dict[str, Any]]:
    """Aggregate exact resolved allele motif paths for the per-locus figure.

    Every input row represents one reported allele.  Only motif paths that can
    be parsed completely are counted.  The complete aggregate is written to the
    figure-source CSV; ``displayed`` identifies the most frequent paths selected
    for the plot when a locus has more rows than can be shown legibly.
    """
    grouped: dict[str, dict[str, list[dict[str, Any]]]] = defaultdict(
        lambda: defaultdict(list)
    )
    for allele in alleles:
        composition = str(
            allele.get("analysis_composition") or allele.get("composition") or ""
        )
        if not parse_rle_motifs(composition):
            continue
        grouped[str(allele.get("locus", "UNKNOWN"))][composition].append(allele)

    output: list[dict[str, Any]] = []
    for locus in sorted(grouped, key=natural_key):
        locus_sample_count = len(
            {
                str(source_row.get("sampleID", ""))
                for composition_rows in grouped[locus].values()
                for source_row in composition_rows
            }
        )
        provisional: list[dict[str, Any]] = []
        for composition, source_rows in grouped[locus].items():
            motifs = parse_rle_motifs(composition)
            expected = str(source_rows[0].get("reported_repeat_unit", ""))
            all_counts = Counter(motifs)
            interruption_counts = Counter(
                motif for motif in motifs if not motif_matches(motif, expected)
            )
            provisional.append(
                {
                    "locus": locus,
                    "composition": composition,
                    "repeat_units": len(motifs),
                    "allele_count": len(source_rows),
                    "sample_count": len(
                        {str(row.get("sampleID", "")) for row in source_rows}
                    ),
                    "locus_sample_count": locus_sample_count,
                    "reported_repeat_unit": expected,
                    "structure_class": (
                        "interrupted" if interruption_counts else "pure"
                    ),
                    "motif_counts_json": json.dumps(
                        dict(sorted(all_counts.items())), separators=(",", ":")
                    ),
                    "interruption_motif_counts_json": json.dumps(
                        dict(sorted(interruption_counts.items())),
                        separators=(",", ":"),
                    ),
                }
            )

        # Select by observed frequency, then display the selected paths from
        # longest to shortest to mirror the reference sequence-composition plot.
        ranked = sorted(
            provisional,
            key=lambda row: (
                -int(row["allele_count"]),
                -int(row["repeat_units"]),
                natural_key(row["composition"]),
            ),
        )
        selected = {
            str(row["composition"])
            for row in (
                ranked if max_display_rows <= 0 else ranked[:max_display_rows]
            )
        }
        displayed = sorted(
            (row for row in provisional if str(row["composition"]) in selected),
            key=lambda row: (
                -int(row["repeat_units"]),
                natural_key(row["composition"]),
            ),
        )
        display_rank = {
            str(row["composition"]): index
            for index, row in enumerate(displayed, start=1)
        }
        for row in sorted(
            provisional,
            key=lambda item: (
                0 if str(item["composition"]) in selected else 1,
                display_rank.get(str(item["composition"]), 10**9),
                -int(item["allele_count"]),
                natural_key(item["composition"]),
            ),
        ):
            row["displayed"] = "yes" if str(row["composition"]) in selected else "no"
            row["display_rank"] = display_rank.get(str(row["composition"]), "NA")
            output.append(row)
    return output


def _compact_composition_label(composition: str) -> str:
    """Use the compact ``GGC(4)GGT(1)`` label style of the reference figure."""
    motifs = parse_rle_motifs(composition)
    if not motifs:
        return composition
    compact: list[str] = []
    for motif, group in itertools.groupby(motifs):
        compact.append(f"{motif}({sum(1 for _ in group)})")
    return "".join(compact)


def _sequence_motif_colours(
    rows: Sequence[dict[str, Any]], expected: str
) -> dict[str, str]:
    totals: Counter[str] = Counter()
    for row in rows:
        totals.update(json.loads(str(row["motif_counts_json"])))
    motifs = sorted(totals, key=lambda motif: (-totals[motif], natural_key(motif)))
    interruption_palette = [
        "#4477AA", "#EE7733", "#228833", "#CC6677", "#66CCEE",
        "#AA3377", "#BBBB44", "#882255", "#44AA99", "#332288",
    ]
    colours: dict[str, str] = {}
    colour_index = 0
    for motif in motifs:
        if expected and motif_matches(motif, expected):
            colours[motif] = "#B9BDC2"
            continue
        if colour_index < len(interruption_palette):
            colours[motif] = interruption_palette[colour_index]
        else:
            red, green, blue = colorsys.hsv_to_rgb(
                (colour_index * 0.61803398875) % 1.0, 0.58, 0.78
            )
            colours[motif] = "#{:02X}{:02X}{:02X}".format(
                round(red * 255), round(green * 255), round(blue * 255)
            )
        colour_index += 1
    return colours


def _draw_motif_pie(
    axis: Any,
    y: float,
    counts: Mapping[str, int],
    colours: Mapping[str, str],
    wedge_class: Any,
) -> None:
    total = sum(int(value) for value in counts.values())
    if total <= 0:
        axis.add_patch(
            wedge_class(
                (0, y), 0.34, 0, 360,
                facecolor="white", edgecolor="#B9BDC2", linewidth=0.8,
            )
        )
        axis.text(0, y, "-", ha="center", va="center", color="#777777", fontsize=7)
        return
    start = 90.0
    for motif, count in sorted(counts.items(), key=lambda item: natural_key(item[0])):
        extent = 360.0 * int(count) / total
        axis.add_patch(
            wedge_class(
                (0, y), 0.34, start, start + extent,
                facecolor=colours.get(motif, "#999999"),
                edgecolor="white", linewidth=0.45,
            )
        )
        start += extent
    axis.add_patch(
        wedge_class(
            (0, y), 0.34, 0, 360,
            facecolor="none", edgecolor="#777777", linewidth=0.45,
        )
    )


def _sequence_frequency_ticks(maximum: int) -> list[int]:
    candidates = [0, 1, 2, 5, 10, 20, 50, 100, 200, 500, 1000, 2000, 5000, 10000]
    ticks = [value for value in candidates if value <= maximum]
    if maximum not in ticks:
        ticks.append(maximum)
    return sorted(set(ticks))


def _plot_sequence_composition_frequency(
    plt: Any,
    locus: str,
    rows: Sequence[dict[str, Any]],
    base: Path,
    formats: Sequence[str],
    dpi: int,
) -> None:
    """Draw one reference-style sequence-composition frequency plot."""
    from matplotlib.patches import Patch, Wedge

    displayed = sorted(
        (row for row in rows if str(row.get("displayed")) == "yes"),
        key=lambda row: int(row.get("display_rank", 10**9)),
    )
    if not displayed:
        return
    expected = str(displayed[0].get("reported_repeat_unit", ""))
    colours = _sequence_motif_colours(displayed, expected)
    n_rows = len(displayed)
    legend_rows = max(1, math.ceil(len(colours) / 8))
    height = max(5.0, 2.25 + 0.34 * n_rows + 0.30 * legend_rows)
    figure = plt.figure(figsize=(16.0, height))
    grid = figure.add_gridspec(
        1, 5,
        width_ratios=[7.8, 0.62, 0.62, 0.78, 4.2],
        left=0.025, right=0.985, top=0.84, bottom=min(0.30, 0.11 + 0.035 * legend_rows),
        wspace=0.10,
    )
    axis_label = figure.add_subplot(grid[0, 0])
    axis_all = figure.add_subplot(grid[0, 1])
    axis_interruptions = figure.add_subplot(grid[0, 2])
    axis_size = figure.add_subplot(grid[0, 3])
    axis_frequency = figure.add_subplot(grid[0, 4])
    axes = [axis_label, axis_all, axis_interruptions, axis_size, axis_frequency]
    for axis in axes:
        axis.set_ylim(n_rows - 0.5, -0.5)

    axis_label.set_xlim(0, 1)
    axis_label.axis("off")
    axis_label.set_title("Reported allele motif path", fontsize=9, pad=8)
    for y, row in enumerate(displayed):
        label = _compact_composition_label(str(row["composition"]))
        if len(label) > 105:
            label = label[:102] + "..."
        axis_label.text(
            0.99, y, label, ha="right", va="center",
            fontsize=6.7, family="DejaVu Sans Mono",
        )

    for axis, title in (
        (axis_all, "All motifs"),
        (axis_interruptions, "Interruptions"),
    ):
        axis.set_xlim(-0.55, 0.55)
        axis.set_aspect("equal", adjustable="box")
        axis.axis("off")
        # With only one or two sequence rows an equal-aspect pie axis becomes
        # vertically compact.  Put the heading at the top of the grid cell so
        # it stays aligned with the other column headings for single samples.
        cell = grid[0, 1 if axis is axis_all else 2].get_position(figure)
        figure.text(
            (cell.x0 + cell.x1) / 2, 0.855, title,
            ha="center", va="bottom", fontsize=9,
        )
    for y, row in enumerate(displayed):
        all_counts = json.loads(str(row["motif_counts_json"]))
        interruption_counts = json.loads(str(row["interruption_motif_counts_json"]))
        _draw_motif_pie(axis_all, y, all_counts, colours, Wedge)
        _draw_motif_pie(axis_interruptions, y, interruption_counts, colours, Wedge)

    axis_size.set_xlim(0, 1)
    axis_size.axis("off")
    axis_size.set_title("Repeat units", fontsize=9, pad=8)
    for y, row in enumerate(displayed):
        axis_size.text(
            0.5, y, str(row["repeat_units"]), ha="center", va="center", fontsize=7.5
        )

    frequencies = [int(row["allele_count"]) for row in displayed]
    maximum = max(frequencies)
    bars = axis_frequency.barh(
        range(n_rows), frequencies, height=0.56,
        color=FIGURE_PALETTE["novel"], alpha=0.92,
    )
    if maximum >= 10:
        axis_frequency.set_xscale("symlog", linthresh=1.0, linscale=0.7, base=10)
        ticks = _sequence_frequency_ticks(maximum)
        axis_frequency.set_xticks(ticks, [str(value) for value in ticks])
    axis_frequency.set_xlim(0, maximum * 1.35 if maximum > 1 else 1.6)
    axis_frequency.set_yticks([])
    axis_frequency.set_xlabel("Number of reported alleles")
    axis_frequency.set_title("Allele frequency", fontsize=9, pad=8)
    axis_frequency.grid(axis="x", color=FIGURE_PALETTE["grid"], linewidth=0.55)
    axis_frequency.set_axisbelow(True)
    for bar, count in zip(bars, frequencies):
        axis_frequency.annotate(
            str(count),
            xy=(bar.get_width(), bar.get_y() + bar.get_height() / 2),
            xytext=(4, 0), textcoords="offset points",
            va="center", ha="left", fontsize=6.8,
        )

    n_alleles = sum(int(row["allele_count"]) for row in rows)
    n_samples = max((int(row["locus_sample_count"]) for row in rows), default=0)
    allele_word = "allele call" if n_alleles == 1 else "allele calls"
    sample_word = "sample" if n_samples == 1 else "samples"
    truncated = sum(str(row.get("displayed")) != "yes" for row in rows)
    title_suffix = (
        f" (top {n_rows} structures shown)" if truncated else ""
    )
    figure.suptitle(
        f"{locus}: frequency of reconstructed allele sequence compositions{title_suffix}",
        x=0.025, y=0.97, ha="left", fontsize=13, fontweight="bold",
    )
    figure.text(
        0.025, 0.91,
        f"Each row is one exact resolved motif path; {n_alleles} {allele_word} "
        f"from {n_samples} {sample_word}. The interruption pie excludes the "
        f"catalogue repeat unit ({expected or 'unresolved'}).",
        ha="left", va="center", fontsize=8, color="#444444",
    )
    handles = [
        Patch(
            facecolor=colour,
            edgecolor="none",
            label=(f"{motif} (catalogue repeat)" if expected and motif_matches(motif, expected) else motif),
        )
        for motif, colour in sorted(colours.items(), key=lambda item: natural_key(item[0]))
    ]
    if handles:
        figure.legend(
            handles=handles,
            loc="lower center", bbox_to_anchor=(0.5, 0.018),
            ncol=min(8, len(handles)), frameon=False,
            title="Repeat-unit colour key", fontsize=7.5, title_fontsize=8,
        )
    _save_figure(figure, base, formats, dpi)
    plt.close(figure)


def _configure_matplotlib() -> Any:
    import matplotlib

    matplotlib.use("Agg", force=True)
    import matplotlib.pyplot as plt

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 9,
            "axes.titlesize": 11,
            "axes.labelsize": 9,
            "xtick.labelsize": 8,
            "ytick.labelsize": 8,
            "legend.fontsize": 8,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.linewidth": 0.8,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "svg.fonttype": "none",
            "savefig.facecolor": "white",
        }
    )
    return plt


def _save_figure(figure: Any, base: Path, formats: Sequence[str], dpi: int) -> None:
    base.parent.mkdir(parents=True, exist_ok=True)
    for extension in formats:
        destination = base.parent / f"{base.name}.{extension}"
        figure.savefig(
            destination,
            dpi=dpi if extension == "png" else None,
            bbox_inches="tight",
            facecolor="white",
        )


class SimpleDrawing:
    """Tiny deterministic vector scene with SVG, PDF, and Pillow renderers."""

    def __init__(self, width: float, height: float) -> None:
        self.width = width
        self.height = height
        self.ops: list[tuple[str, dict[str, Any]]] = []

    def line(self, x1: float, y1: float, x2: float, y2: float, color: str = "#777777", width: float = 1.0) -> None:
        self.ops.append(("line", locals() | {"self": None}))

    def rect(self, x: float, y: float, width: float, height: float, fill: str, stroke: str = "none", stroke_width: float = 0.0) -> None:
        self.ops.append(("rect", locals() | {"self": None}))

    def circle(self, x: float, y: float, radius: float, fill: str, stroke: str = "none", stroke_width: float = 0.0) -> None:
        self.ops.append(("circle", locals() | {"self": None}))

    def text(self, x: float, y: float, value: Any, size: float = 9.0, color: str = "#222222", anchor: str = "start", bold: bool = False) -> None:
        self.ops.append(("text", locals() | {"self": None, "value": str(value)}))

    def _svg(self) -> str:
        lines = [
            f'<svg xmlns="http://www.w3.org/2000/svg" width="{self.width:.2f}pt" height="{self.height:.2f}pt" viewBox="0 0 {self.width:.2f} {self.height:.2f}">',
            '<rect width="100%" height="100%" fill="white"/>',
        ]
        for kind, op in self.ops:
            if kind == "line":
                lines.append(
                    f'<line x1="{op["x1"]:.2f}" y1="{op["y1"]:.2f}" x2="{op["x2"]:.2f}" y2="{op["y2"]:.2f}" stroke="{op["color"]}" stroke-width="{op["width"]:.2f}"/>'
                )
            elif kind == "rect":
                lines.append(
                    f'<rect x="{op["x"]:.2f}" y="{op["y"]:.2f}" width="{max(0, op["width"]):.2f}" height="{max(0, op["height"]):.2f}" fill="{op["fill"]}" stroke="{op["stroke"]}" stroke-width="{op["stroke_width"]:.2f}"/>'
                )
            elif kind == "circle":
                lines.append(
                    f'<circle cx="{op["x"]:.2f}" cy="{op["y"]:.2f}" r="{op["radius"]:.2f}" fill="{op["fill"]}" stroke="{op["stroke"]}" stroke-width="{op["stroke_width"]:.2f}"/>'
                )
            elif kind == "text":
                anchor = {"start": "start", "middle": "middle", "end": "end"}.get(op["anchor"], "start")
                weight = "500" if op["bold"] else "400"
                lines.append(
                    f'<text x="{op["x"]:.2f}" y="{op["y"]:.2f}" fill="{op["color"]}" font-family="DejaVu Sans,Arial,sans-serif" font-size="{op["size"]:.2f}" font-weight="{weight}" text-anchor="{anchor}">{html.escape(op["value"])}</text>'
                )
        lines.append("</svg>")
        return "\n".join(lines) + "\n"

    def save(self, base: Path, formats: Sequence[str], dpi: int) -> None:
        base.parent.mkdir(parents=True, exist_ok=True)
        if "svg" in formats:
            (base.parent / f"{base.name}.svg").write_text(self._svg(), encoding="utf-8")
        if "png" in formats:
            self._save_png(base.parent / f"{base.name}.png", dpi)
        if "pdf" in formats:
            self._save_pdf(base.parent / f"{base.name}.pdf")

    def _save_png(self, path: Path, dpi: int) -> None:
        from PIL import Image, ImageDraw, ImageFont

        scale = dpi / 72.0
        image = Image.new("RGB", (max(1, round(self.width * scale)), max(1, round(self.height * scale))), "white")
        draw = ImageDraw.Draw(image)
        font_cache: dict[tuple[int, bool], Any] = {}

        def font(size: float, bold: bool) -> Any:
            key = (max(7, round(size * scale)), bold)
            if key not in font_cache:
                names = ["DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf", "Arial.ttf"]
                for name in names:
                    try:
                        font_cache[key] = ImageFont.truetype(name, key[0])
                        break
                    except OSError:
                        continue
                else:
                    font_cache[key] = ImageFont.load_default()
            return font_cache[key]

        for kind, op in self.ops:
            if kind == "line":
                draw.line(
                    tuple(round(value * scale) for value in [op["x1"], op["y1"], op["x2"], op["y2"]]),
                    fill=op["color"], width=max(1, round(op["width"] * scale)),
                )
            elif kind == "rect":
                box = tuple(round(value * scale) for value in [op["x"], op["y"], op["x"] + op["width"], op["y"] + op["height"]])
                draw.rectangle(box, fill=op["fill"], outline=None if op["stroke"] == "none" else op["stroke"], width=max(1, round(op["stroke_width"] * scale)))
            elif kind == "circle":
                box = tuple(round(value * scale) for value in [op["x"] - op["radius"], op["y"] - op["radius"], op["x"] + op["radius"], op["y"] + op["radius"]])
                draw.ellipse(box, fill=op["fill"], outline=None if op["stroke"] == "none" else op["stroke"], width=max(1, round(op["stroke_width"] * scale)))
            elif kind == "text":
                fnt = font(op["size"], op["bold"])
                x, y = op["x"] * scale, op["y"] * scale
                bbox = draw.textbbox((0, 0), op["value"], font=fnt)
                text_width = bbox[2] - bbox[0]
                if op["anchor"] == "middle":
                    x -= text_width / 2
                elif op["anchor"] == "end":
                    x -= text_width
                draw.text((round(x), round(y - op["size"] * scale * 0.82)), op["value"], font=fnt, fill=op["color"])
        image.save(path, dpi=(dpi, dpi), optimize=True)

    def _save_pdf(self, path: Path) -> None:
        from reportlab.pdfgen import canvas
        from reportlab.lib.colors import HexColor

        pdf = canvas.Canvas(str(path), pagesize=(self.width, self.height), pageCompression=1)
        for kind, op in self.ops:
            if kind == "line":
                pdf.setStrokeColor(HexColor(op["color"]))
                pdf.setLineWidth(op["width"])
                pdf.line(op["x1"], self.height - op["y1"], op["x2"], self.height - op["y2"])
            elif kind == "rect":
                pdf.setFillColor(HexColor(op["fill"]))
                stroke = op["stroke"] != "none"
                if stroke:
                    pdf.setStrokeColor(HexColor(op["stroke"]))
                    pdf.setLineWidth(op["stroke_width"])
                pdf.rect(op["x"], self.height - op["y"] - op["height"], op["width"], op["height"], fill=1, stroke=int(stroke))
            elif kind == "circle":
                pdf.setFillColor(HexColor(op["fill"]))
                stroke = op["stroke"] != "none"
                if stroke:
                    pdf.setStrokeColor(HexColor(op["stroke"]))
                    pdf.setLineWidth(op["stroke_width"])
                pdf.circle(op["x"], self.height - op["y"], op["radius"], fill=1, stroke=int(stroke))
            elif kind == "text":
                pdf.setFillColor(HexColor(op["color"]))
                pdf.setFont("Helvetica-Bold" if op["bold"] else "Helvetica", op["size"])
                y = self.height - op["y"]
                if op["anchor"] == "middle":
                    pdf.drawCentredString(op["x"], y, op["value"])
                elif op["anchor"] == "end":
                    pdf.drawRightString(op["x"], y, op["value"])
                else:
                    pdf.drawString(op["x"], y, op["value"])
        pdf.showPage()
        pdf.save()


def _simple_axes(drawing: SimpleDrawing, left: float, top: float, width: float, height: float, x_ticks: Sequence[tuple[float, str]]) -> None:
    drawing.line(left, top + height, left + width, top + height, FIGURE_PALETTE["text"], 0.8)
    for fraction, label in x_ticks:
        x = left + fraction * width
        drawing.line(x, top, x, top + height, FIGURE_PALETTE["grid"], 0.5)
        drawing.text(x, top + height + 15, label, 8, anchor="middle")


def _generate_figures_fallback(
    alleles: Sequence[dict[str, Any]],
    summaries: Sequence[dict[str, Any]],
    motif_rows: Sequence[dict[str, Any]],
    sequence_composition_rows: Sequence[dict[str, Any]],
    output_dir: Path,
    formats: Sequence[str],
    dpi: int,
    per_locus: bool,
) -> dict[str, Any]:
    figures_dir = output_dir / "figures"
    row_height = 15.0

    ordered = sorted(summaries, key=lambda row: (row["interruption_prevalence"], natural_key(row["locus"])))
    height = 70 + row_height * len(ordered)
    drawing = SimpleDrawing(620, height)
    drawing.text(12, 22, "Cohort repeat-interruption prevalence by locus", 13, bold=True)
    left, top, width = 130.0, 40.0, 450.0
    _simple_axes(drawing, left, top, width, row_height * len(ordered), [(0, "0"), (0.25, "0.25"), (0.5, "0.50"), (0.75, "0.75"), (1, "1.00")])
    for index, row in enumerate(ordered):
        y = top + index * row_height + row_height / 2
        drawing.text(left - 7, y + 3, row["locus"], 7.5, anchor="end")
        low = left + row["interruption_ci95_lower"] * width
        high = left + row["interruption_ci95_upper"] * width
        centre = left + row["interruption_prevalence"] * width
        drawing.line(low, y, high, y, "#777777", 0.9)
        drawing.line(low, y - 2.5, low, y + 2.5, "#777777", 0.9)
        drawing.line(high, y - 2.5, high, y + 2.5, "#777777", 0.9)
        drawing.circle(centre, y, 2.7, FIGURE_PALETTE["interrupted"])
    drawing.text(left + width / 2, height - 8, "Interrupted alleles among resolved alleles (proportion, Wilson 95% CI)", 8.5, anchor="middle")
    drawing.save(figures_dir / "cohort_interruption_prevalence", formats, dpi)

    ordered_qc = sorted(summaries, key=lambda row: (row["call_rate"], natural_key(row["locus"])))
    height = 76 + row_height * len(ordered_qc)
    drawing = SimpleDrawing(780, height)
    drawing.text(12, 22, "Cohort callability and supporting-read overview", 13, bold=True)
    left, top, width_call, gap, width_depth = 130.0, 44.0, 280.0, 70.0, 230.0
    _simple_axes(drawing, left, top, width_call, row_height * len(ordered_qc), [(0, "0"), (0.5, "0.5"), (1, "1.0")])
    depths = [safe_float(row["median_reported_cluster_depth"]) or 0 for row in ordered_qc]
    depth_max = max(depths + [1.0])
    depth_left = left + width_call + gap
    _simple_axes(drawing, depth_left, top, width_depth, row_height * len(ordered_qc), [(0, "0"), (0.5, f"{depth_max/2:g}"), (1, f"{depth_max:g}")])
    drawing.text(left, 37, "A  Callability", 9, bold=True)
    drawing.text(depth_left, 37, "B  Median cluster depth", 9, bold=True)
    for index, (row, depth) in enumerate(zip(ordered_qc, depths)):
        y = top + index * row_height + 3
        drawing.text(left - 7, y + 7, row["locus"], 7.5, anchor="end")
        drawing.rect(left, y, row["call_rate"] * width_call, 8.5, FIGURE_PALETTE["pure"])
        drawing.circle(depth_left + depth / depth_max * width_depth, y + 4.25, 2.6, FIGURE_PALETTE["interrupted"])
    drawing.save(figures_dir / "cohort_qc_overview", formats, dpi)

    if motif_rows:
        top_rows = list(motif_rows[:30])
        height = 70 + row_height * len(top_rows)
        drawing = SimpleDrawing(620, height)
        drawing.text(12, 22, "Most frequent non-canonical repeat motifs", 13, bold=True)
        left, top, width = 180.0, 40.0, 390.0
        maximum = max(int(row["count"]) for row in top_rows)
        _simple_axes(drawing, left, top, width, row_height * len(top_rows), [(0, "0"), (0.5, f"{maximum/2:g}"), (1, str(maximum))])
        for index, row in enumerate(reversed(top_rows)):
            y = top + index * row_height + 3
            label = f"{row['locus']}  {row['motif']}"
            colour = FIGURE_PALETTE["pathogenic" if row["classification"] == "pathogenic" else "benign" if row["classification"] == "benign" else "novel"]
            drawing.text(left - 7, y + 7, label, 7.5, anchor="end")
            bar_width = int(row["count"]) / maximum * width
            drawing.rect(left, y, bar_width, 8.5, colour)
            drawing.text(left + bar_width + 4, y + 7, row["count"], 7.5)
        drawing.save(figures_dir / "cohort_interruption_motif_spectrum", formats, dpi)

    n_locus_summary_figures = 0
    n_locus_sequence_figures = 0
    if per_locus:
        by_locus: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in alleles:
            by_locus[str(row["locus"])].append(row)
        for locus in sorted(by_locus, key=natural_key):
            rows = [row for row in by_locus[locus] if safe_int(row["size"]) is not None]
            if not rows:
                continue
            drawing = SimpleDrawing(840, 340)
            drawing.text(12, 22, f"{locus}: repeat length and interruption architecture", 13, bold=True)
            size_counts: dict[int, Counter[str]] = defaultdict(Counter)
            for row in rows:
                size_counts[int(row["size"])][str(row["structure_class"])] += 1
            size_values = sorted(size_counts)
            x0, y0, w, h = 55.0, 60.0, 300.0, 220.0
            max_count = max(sum(size_counts[size].values()) for size in size_values)
            drawing.text(x0, 45, "A  Allele-size distribution", 9, bold=True)
            drawing.line(x0, y0, x0, y0 + h, FIGURE_PALETTE["text"], 0.8)
            drawing.line(x0, y0 + h, x0 + w, y0 + h, FIGURE_PALETTE["text"], 0.8)
            for fraction, label in [(0.0, "0"), (0.5, f"{max_count / 2:g}"), (1.0, str(max_count))]:
                y_tick = y0 + h - fraction * h
                if fraction:
                    drawing.line(x0, y_tick, x0 + w, y_tick, FIGURE_PALETTE["grid"], 0.5)
                drawing.text(x0 - 6, y_tick + 3, label, 7, anchor="end")
            bar_space = w / max(1, len(size_values))
            bar_width = min(12.0, bar_space * 0.72)
            for index, size in enumerate(size_values):
                x = x0 + (index + 0.5) * bar_space - bar_width / 2
                pure = size_counts[size]["pure"]
                interrupted = size_counts[size]["interrupted"]
                pure_height = pure / max_count * h
                interrupted_height = interrupted / max_count * h
                drawing.rect(x, y0 + h - pure_height, bar_width, pure_height, FIGURE_PALETTE["pure"])
                drawing.rect(x, y0 + h - pure_height - interrupted_height, bar_width, interrupted_height, FIGURE_PALETTE["interrupted"])
                if len(size_values) <= 20 or index % max(1, len(size_values) // 10) == 0:
                    drawing.text(x + bar_width / 2, y0 + h + 13, size, 7, anchor="middle")
            drawing.text(x0 + w / 2, 320, "Reconstructed repeat size (units)", 8, anchor="middle")

            structure_counts = Counter(str(row["composition"]) for row in rows)
            top = structure_counts.most_common(12)
            x1, y1, w1 = 505.0, 60.0, 285.0
            drawing.text(405, 45, "B  Most supported motif compositions", 9, bold=True)
            drawing.rect(670, 37, 9, 7, FIGURE_PALETTE["pure"])
            drawing.text(683, 44, "Pure", 7)
            drawing.rect(725, 37, 9, 7, FIGURE_PALETTE["interrupted"])
            drawing.text(738, 44, "Interrupted", 7)
            maximum = max(count for _, count in top)
            for index, (composition, count) in enumerate(reversed(top)):
                y = y1 + index * 17
                label = composition if len(composition) <= 38 else composition[:35] + "..."
                categories = Counter(
                    str(row["structure_class"])
                    for row in rows
                    if str(row["composition"]) == composition
                )
                category = categories.most_common(1)[0][0] if categories else "unresolved"
                drawing.text(x1 - 7, y + 8, label, 6.8, anchor="end")
                bar_width_value = count / maximum * w1
                drawing.rect(x1, y, bar_width_value, 10, FIGURE_PALETTE[category])
                drawing.text(x1 + bar_width_value + 4, y + 8, count, 7)
            drawing.text(x1 + w1 / 2, 320, "Allele count", 8, anchor="middle")
            drawing.save(figures_dir / "locus" / f"{sanitise_filename(locus)}.summary", formats, dpi)
            n_locus_summary_figures += 1

        sequence_by_locus: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in sequence_composition_rows:
            if str(row.get("displayed")) == "yes":
                sequence_by_locus[str(row["locus"])].append(row)
        for locus in sorted(sequence_by_locus, key=natural_key):
            frequency_rows = sorted(
                sequence_by_locus[locus],
                key=lambda row: int(row.get("display_rank", 10**9)),
            )
            if not frequency_rows:
                continue
            expected = str(frequency_rows[0].get("reported_repeat_unit", ""))
            colours = _sequence_motif_colours(frequency_rows, expected)
            row_height = 18.0
            height = 105.0 + row_height * len(frequency_rows)
            drawing = SimpleDrawing(1120, height)
            drawing.text(
                14, 23,
                f"{locus}: frequency of reconstructed allele sequence compositions",
                13, bold=True,
            )
            drawing.text(390, 50, "All motifs", 8, anchor="middle", bold=True)
            drawing.text(520, 50, "Interruptions", 8, anchor="middle", bold=True)
            drawing.text(655, 50, "Repeat units", 8, anchor="middle", bold=True)
            drawing.text(890, 50, "Allele frequency", 8, anchor="middle", bold=True)
            maximum = max(int(row["allele_count"]) for row in frequency_rows)
            for index, row in enumerate(frequency_rows):
                y = 64 + index * row_height
                label = _compact_composition_label(str(row["composition"]))
                if len(label) > 66:
                    label = label[:63] + "..."
                drawing.text(335, y + 9, label, 6.7, anchor="end")
                all_counts = json.loads(str(row["motif_counts_json"]))
                interruption_counts = json.loads(
                    str(row["interruption_motif_counts_json"])
                )
                for left, counts in ((350.0, all_counts), (480.0, interruption_counts)):
                    total = sum(int(value) for value in counts.values())
                    if total:
                        cursor = left
                        for motif, count in sorted(counts.items(), key=lambda item: natural_key(item[0])):
                            width = 80.0 * int(count) / total
                            drawing.rect(cursor, y + 2, width, 10, colours.get(motif, "#999999"))
                            cursor += width
                    else:
                        drawing.rect(left, y + 2, 80, 10, "#FFFFFF", "#B9BDC2", 0.6)
                drawing.text(655, y + 9, row["repeat_units"], 7.2, anchor="middle")
                bar_width = 335.0 * int(row["allele_count"]) / maximum
                drawing.rect(700, y + 2, bar_width, 10, FIGURE_PALETTE["novel"])
                drawing.text(704 + bar_width, y + 9, row["allele_count"], 7)
            drawing.text(
                700, height - 16,
                "Fallback renderer: coloured strips show motif proportions; bars use a linear scale.",
                7, color="#555555",
            )
            drawing.save(
                figures_dir / "locus" /
                f"{sanitise_filename(locus)}.sequence_composition_frequency",
                formats, dpi,
            )
            n_locus_sequence_figures += 1
    return {
        "backend": "stdlib-svg+pillow+reportlab",
        "n_cohort_figure_bases": 3 if motif_rows else 2,
        "n_locus_summary_figure_bases": n_locus_summary_figures,
        "n_locus_sequence_composition_figure_bases": n_locus_sequence_figures,
        "n_locus_figure_bases": n_locus_summary_figures + n_locus_sequence_figures,
        "formats": list(formats),
    }


def generate_figures(
    records: Sequence[dict[str, Any]],
    catalog_loci: Sequence[CatalogLocus],
    output_dir: Path,
    formats: Sequence[str],
    dpi: int,
    per_locus: bool,
) -> dict[str, Any]:
    alleles, summaries, motif_rows = build_figure_data(records, catalog_loci)
    sequence_composition_rows = build_sequence_composition_frequency_data(alleles)
    figure_data = output_dir / "figure_data"
    _write_csv_atomic(
        figure_data / "allele_level.csv",
        alleles,
        [
            "batch", "sampleID", "locus", "allele", "composition",
            "analysis_composition", "size",
            "read_depth", "structure_class", "reported_repeat_unit", "normal_max",
            "pathogenic_min", "qc_status",
        ],
    )
    _write_csv_atomic(
        figure_data / "locus_summary.csv",
        summaries,
        [
            "locus", "n_samples", "n_expected_alleles", "n_resolved_alleles",
            "n_pure_alleles", "n_interrupted_alleles", "n_unresolved_alleles",
            "call_rate", "interruption_prevalence", "interruption_ci95_lower",
            "interruption_ci95_upper", "median_reported_cluster_depth",
        ],
    )
    _write_csv_atomic(
        figure_data / "interruption_motif_spectrum.csv",
        motif_rows,
        ["locus", "motif", "classification", "count"],
    )
    _write_csv_atomic(
        figure_data / "sequence_composition_frequency.csv",
        sequence_composition_rows,
        [
            "locus", "composition", "repeat_units", "allele_count",
            "sample_count", "locus_sample_count", "reported_repeat_unit", "structure_class",
            "motif_counts_json", "interruption_motif_counts_json",
            "displayed", "display_rank",
        ],
    )
    if not summaries:
        return {"n_figures": 0, "warning": "no figure data"}

    try:
        plt = _configure_matplotlib()
    except ImportError:
        return _generate_figures_fallback(
            alleles, summaries, motif_rows, sequence_composition_rows,
            output_dir, formats, dpi, per_locus
        )

    figures_dir = output_dir / "figures"
    ordered = sorted(summaries, key=lambda row: (row["interruption_prevalence"], natural_key(row["locus"])))
    height = max(5.0, min(26.0, 0.30 * len(ordered) + 1.8))

    # Forest plot: interruption prevalence and Wilson 95% confidence intervals.
    figure, axis = plt.subplots(figsize=(8.3, height), constrained_layout=True)
    y_positions = list(range(len(ordered)))
    proportions = [row["interruption_prevalence"] for row in ordered]
    # Keep this final clamp at the plotting boundary as a defensive guard for
    # figure-data imported from older runs or other numeric implementations.
    lower_errors = [
        max(0.0, p - row["interruption_ci95_lower"])
        for p, row in zip(proportions, ordered)
    ]
    upper_errors = [
        max(0.0, row["interruption_ci95_upper"] - p)
        for p, row in zip(proportions, ordered)
    ]
    axis.errorbar(
        proportions,
        y_positions,
        xerr=[lower_errors, upper_errors],
        fmt="o",
        markersize=4,
        color=FIGURE_PALETTE["interrupted"],
        ecolor="#777777",
        elinewidth=1,
        capsize=2,
    )
    axis.set_yticks(y_positions, [row["locus"] for row in ordered])
    axis.set_xlim(-0.02, 1.02)
    axis.set_xlabel("Interrupted alleles among resolved alleles (proportion, 95% CI)")
    axis.set_title("Cohort repeat-interruption prevalence by locus", loc="left", weight="bold")
    axis.grid(axis="x", color=FIGURE_PALETTE["grid"], linewidth=0.6)
    _save_figure(figure, figures_dir / "cohort_interruption_prevalence", formats, dpi)
    plt.close(figure)

    # Callability and support overview on aligned locus axes.
    ordered_qc = sorted(summaries, key=lambda row: (row["call_rate"], natural_key(row["locus"])))
    figure, (axis_call, axis_depth) = plt.subplots(
        ncols=2,
        figsize=(11.7, height),
        sharey=True,
        gridspec_kw={"width_ratios": [1.25, 1]},
        constrained_layout=True,
    )
    y_positions = list(range(len(ordered_qc)))
    call_rates = [row["call_rate"] for row in ordered_qc]
    axis_call.barh(y_positions, call_rates, color=FIGURE_PALETTE["pure"], height=0.68)
    axis_call.set_yticks(y_positions, [row["locus"] for row in ordered_qc])
    axis_call.set_xlim(0, 1)
    axis_call.set_xlabel("Resolved expected alleles (proportion)")
    axis_call.set_title("A  Callability", loc="left", weight="bold")
    axis_call.grid(axis="x", color=FIGURE_PALETTE["grid"], linewidth=0.6)
    depth_values = [
        safe_float(row["median_reported_cluster_depth"])
        for row in ordered_qc
    ]
    x_depth = [value if value is not None else 0 for value in depth_values]
    axis_depth.scatter(x_depth, y_positions, color=FIGURE_PALETTE["interrupted"], s=18)
    axis_depth.set_xlabel("Median supporting full-spanning reads")
    axis_depth.set_title("B  Reported-cluster support", loc="left", weight="bold")
    axis_depth.grid(axis="x", color=FIGURE_PALETTE["grid"], linewidth=0.6)
    _save_figure(figure, figures_dir / "cohort_qc_overview", formats, dpi)
    plt.close(figure)

    # Most recurrent non-canonical motifs.
    if motif_rows:
        top_motifs = motif_rows[:30]
        figure_height = max(4.0, 0.28 * len(top_motifs) + 1.5)
        figure, axis = plt.subplots(figsize=(8.3, figure_height), constrained_layout=True)
        labels = [f"{row['locus']}  {row['motif']}" for row in reversed(top_motifs)]
        counts = [row["count"] for row in reversed(top_motifs)]
        colours = [
            FIGURE_PALETTE[
                "pathogenic" if row["classification"] == "pathogenic"
                else "benign" if row["classification"] == "benign"
                else "novel"
            ]
            for row in reversed(top_motifs)
        ]
        bars = axis.barh(range(len(labels)), counts, color=colours, height=0.7)
        axis.set_yticks(range(len(labels)), labels)
        axis.set_xlabel("Occurrences across reported allele motif paths")
        axis.set_title("Most frequent non-canonical repeat motifs", loc="left", weight="bold")
        axis.grid(axis="x", color=FIGURE_PALETTE["grid"], linewidth=0.6)
        for bar, count in zip(bars, counts):
            axis.text(bar.get_width(), bar.get_y() + bar.get_height() / 2, f"  {count}", va="center")
        _save_figure(figure, figures_dir / "cohort_interruption_motif_spectrum", formats, dpi)
        plt.close(figure)

    n_locus_summary_figures = 0
    n_locus_sequence_figures = 0
    if per_locus:
        allele_by_locus: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in alleles:
            allele_by_locus[str(row["locus"])].append(row)
        for locus in sorted(allele_by_locus, key=natural_key):
            resolved = [row for row in allele_by_locus[locus] if safe_int(row["size"]) is not None]
            if not resolved:
                continue
            figure, (axis_size, axis_structure) = plt.subplots(
                ncols=2,
                figsize=(11.7, 4.4),
                gridspec_kw={"width_ratios": [1, 1.35]},
                constrained_layout=True,
            )
            pure_sizes = [int(row["size"]) for row in resolved if row["structure_class"] == "pure"]
            interrupted_sizes = [
                int(row["size"]) for row in resolved if row["structure_class"] == "interrupted"
            ]
            all_sizes = pure_sizes + interrupted_sizes
            minimum, maximum = min(all_sizes), max(all_sizes)
            bins = [value - 0.5 for value in range(minimum, maximum + 2)]
            axis_size.hist(
                [pure_sizes, interrupted_sizes],
                bins=bins,
                stacked=True,
                color=[FIGURE_PALETTE["pure"], FIGURE_PALETTE["interrupted"]],
                label=["Pure", "Interrupted"],
                edgecolor="white",
                linewidth=0.4,
            )
            normal_max = safe_int(resolved[0].get("normal_max"))
            pathogenic_min = safe_int(resolved[0].get("pathogenic_min"))
            if normal_max is not None:
                axis_size.axvline(normal_max, color=FIGURE_PALETTE["benign"], linestyle="--", linewidth=1.1, label="Normal max")
            if pathogenic_min is not None:
                axis_size.axvline(pathogenic_min, color=FIGURE_PALETTE["pathogenic"], linestyle=":", linewidth=1.3, label="Pathogenic min")
            axis_size.set_xlabel("Reconstructed repeat size (units)")
            axis_size.set_ylabel("Allele count")
            axis_size.set_title("A  Allele-size distribution", loc="left", weight="bold")
            axis_size.legend(frameon=False)
            axis_size.grid(axis="y", color=FIGURE_PALETTE["grid"], linewidth=0.6)

            structures = Counter(str(row["composition"]) for row in resolved)
            top = structures.most_common(12)
            labels = [composition for composition, _ in reversed(top)]
            counts = [count for _, count in reversed(top)]
            colours = [
                FIGURE_PALETTE[
                    Counter(
                        str(row["structure_class"])
                        for row in resolved
                        if str(row["composition"]) == composition
                    ).most_common(1)[0][0]
                ]
                for composition in labels
            ]
            bars = axis_structure.barh(range(len(labels)), counts, color=colours, height=0.72)
            display_labels = [label if len(label) <= 55 else label[:52] + "..." for label in labels]
            axis_structure.set_yticks(range(len(labels)), display_labels)
            axis_structure.set_xlabel("Allele count")
            axis_structure.set_title("B  Most supported motif compositions", loc="left", weight="bold")
            axis_structure.grid(axis="x", color=FIGURE_PALETTE["grid"], linewidth=0.6)
            for bar, count in zip(bars, counts):
                axis_structure.text(bar.get_width(), bar.get_y() + bar.get_height() / 2, f"  {count}", va="center")
            figure.suptitle(
                f"{locus}: repeat length and interruption architecture",
                x=0.01,
                ha="left",
                fontsize=13,
                fontweight="bold",
            )
            _save_figure(
                figure,
                figures_dir / "locus" / f"{sanitise_filename(locus)}.summary",
                formats,
                dpi,
            )
            plt.close(figure)
            n_locus_summary_figures += 1

        sequence_by_locus: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in sequence_composition_rows:
            sequence_by_locus[str(row["locus"])].append(row)
        for locus in sorted(sequence_by_locus, key=natural_key):
            if not any(str(row.get("displayed")) == "yes" for row in sequence_by_locus[locus]):
                continue
            _plot_sequence_composition_frequency(
                plt,
                locus,
                sequence_by_locus[locus],
                figures_dir / "locus" /
                f"{sanitise_filename(locus)}.sequence_composition_frequency",
                formats,
                dpi,
            )
            n_locus_sequence_figures += 1

    return {
        "backend": "matplotlib",
        "n_cohort_figure_bases": 3 if motif_rows else 2,
        "n_locus_summary_figure_bases": n_locus_summary_figures,
        "n_locus_sequence_composition_figure_bases": n_locus_sequence_figures,
        "n_locus_figure_bases": n_locus_summary_figures + n_locus_sequence_figures,
        "formats": list(formats),
    }


def _default_workers() -> int:
    for variable in ("PBS_NCPUS", "SLURM_CPUS_PER_TASK"):
        value = safe_int(os.environ.get(variable))
        if value and value > 0:
            return value
    return min(4, os.cpu_count() or 1)


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Reconstruct read-level STR nucleotide sequences from ReViewer SVGs "
            "and infer allele-level repeat interruption structures."
        ),
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument("--svg-root", required=True, type=Path, help="Root containing batch/sample/SVG files")
    parser.add_argument("--catalog", required=True, type=Path, help="ExpansionHunter variant catalogue JSON")
    parser.add_argument("--annotation", required=True, type=Path, help="Companion XLSX/CSV with locus or Gene and strand")
    parser.add_argument("--annotation-sheet", default=None, help="Worksheet name; first sheet is used by default")
    parser.add_argument("--output", required=True, type=Path, help="New or explicitly reusable output directory")
    parser.add_argument("--genotypes", type=Path, default=None, help="Optional CSV/TSV/XLSX genotype override table")
    parser.add_argument("--file-map", type=Path, default=None, help="Optional path/sampleID/locus metadata table")
    parser.add_argument(
        "--locus",
        action="append",
        default=[],
        help="Limit to a locus/component; repeat option or use comma-separated values",
    )
    parser.add_argument(
        "--component-mode",
        choices=["combined", "primary", "all"],
        default="combined",
        help=(
            "combined emits one catalogue-ordered multi-component allele per locus; "
            "primary retains only MainReferenceRegion; all emits separate component rows"
        ),
    )
    parser.add_argument("--workers", type=int, default=_default_workers())
    parser.add_argument(
        "--min-spanning-reads", type=int, default=3,
        help="Minimum total full-spanning rows required to avoid an insufficient-depth REVIEW flag",
    )
    parser.add_argument(
        "--min-cluster-reads", type=int, default=3,
        help="Minimum reads supporting a motif path for strong-cluster/allele support",
    )
    parser.add_argument(
        "--max-near-motif-distance",
        type=int,
        default=0,
        help="0 keeps exact complete motif-path clusters (recommended)",
    )
    parser.add_argument("--near-cluster-parent-ratio", type=float, default=3.0)
    parser.add_argument(
        "--strong-cluster-fraction",
        type=float,
        default=0.0,
        help="Optional fraction criterion in addition to --min-cluster-reads; 0 disables",
    )
    parser.add_argument(
        "--heterozygous-interruption-min-fraction", type=float, default=0.50,
        help=(
            "Within each heterozygous panel, a specific non-reference nucleotide is called "
            "at a repeat column only when its fraction is strictly greater than this value"
        ),
    )
    parser.add_argument(
        "--same-size-min-cluster-reads", type=int, default=2,
        help=(
            "For a same-size genotype, an exact motif-path cluster must have at least "
            "this many reads to seed a reported sequence allele"
        ),
    )
    parser.add_argument(
        "--same-size-min-cluster-fraction", type=float, default=0.20,
        help=(
            "For a same-size genotype, an exact motif-path cluster must also contain at "
            "least this fraction of all expected-length reads to seed an allele"
        ),
    )
    parser.add_argument(
        "--same-size-min-top-two-coverage", type=float, default=0.90,
        help=(
            "Minimum fraction of expected-length same-size reads uniquely assigned to the "
            "selected one or two allele clusters"
        ),
    )
    parser.add_argument(
        "--same-size-within-cluster-consensus-min-fraction", type=float, default=0.50,
        help=(
            "Within each selected same-size cluster, a non-reference nucleotide is called "
            "only when its fraction is strictly greater than this value"
        ),
    )
    parser.add_argument(
        "--same-size-max-reported-clusters", type=int, choices=[1, 2], default=2,
        help="Maximum supported same-size sequence clusters reported as diploid alleles",
    )
    parser.add_argument(
        "--same-size-max-assignment-distance", type=int, default=2,
        help=(
            "Maximum motif-unit Hamming distance for assigning a minor same-size read path "
            "to its unique nearest selected allele seed"
        ),
    )
    parser.add_argument(
        "--homozygous-interruption-min-fraction",
        type=float,
        default=None,
        help=argparse.SUPPRESS,
    )
    parser.add_argument(
        "--homozygous-deduplication",
        choices=["none", "cross_panel_exact"],
        default="none",
        help="SVGs have no read IDs; exact-row deduplication is therefore opt-in",
    )
    parser.add_argument(
        "--unresolved-strand-policy",
        choices=["skip", "technical", "error"],
        default="skip",
        help="Never guess missing biological reporting orientations by default",
    )
    parser.add_argument(
        "--write-read-level", action=argparse.BooleanOptionalAction, default=True
    )
    parser.add_argument("--figures", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument(
        "--per-locus-figures", action=argparse.BooleanOptionalAction, default=True
    )
    parser.add_argument("--figure-formats", default="png,pdf,svg")
    parser.add_argument("--dpi", type=int, default=300)
    parser.add_argument("--validate-only", action="store_true")
    parser.add_argument("--strict", action="store_true", help="Exit non-zero on errors or unconfigured loci")
    parser.add_argument("--fail-on-review", action="store_true", help="Exit non-zero if any call requires review")
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Remove prior tool-owned outputs before reusing a non-empty output directory",
    )
    parser.add_argument(
        "--allow-duplicate-sample-locus",
        action="store_true",
        help="Explicitly permit duplicate sampleID/locus rows across batches",
    )
    parser.add_argument("--log-level", choices=["DEBUG", "INFO", "WARNING", "ERROR"], default="INFO")
    parser.add_argument("--version", action="version", version=f"%(prog)s {VERSION}")
    return parser


def _validate_arguments(args: argparse.Namespace) -> None:
    for path, label in [
        (args.svg_root, "SVG root"),
        (args.catalog, "catalogue"),
        (args.annotation, "annotation"),
    ]:
        if not path.exists():
            raise ToolError(f"{label} does not exist: {path}")
    if not args.svg_root.is_dir():
        raise ToolError(f"SVG root is not a directory: {args.svg_root}")
    if args.workers < 1:
        raise ToolError("--workers must be at least 1")
    if (
        args.min_spanning_reads < 1
        or args.min_cluster_reads < 1
        or args.same_size_min_cluster_reads < 1
    ):
        raise ToolError("read support thresholds must be at least 1")
    if args.max_near_motif_distance < 0:
        raise ToolError("--max-near-motif-distance cannot be negative")
    if args.near_cluster_parent_ratio <= 1:
        raise ToolError("--near-cluster-parent-ratio must be greater than 1")
    if not 0 <= args.strong_cluster_fraction <= 1:
        raise ToolError("--strong-cluster-fraction must be between 0 and 1")
    if not 0 <= args.heterozygous_interruption_min_fraction < 1:
        raise ToolError("--heterozygous-interruption-min-fraction must be at least 0 and less than 1")
    if not 0 <= args.same_size_min_cluster_fraction <= 1:
        raise ToolError("--same-size-min-cluster-fraction must be between 0 and 1")
    if not 0 <= args.same_size_min_top_two_coverage <= 1:
        raise ToolError("--same-size-min-top-two-coverage must be between 0 and 1")
    if not 0 <= args.same_size_within_cluster_consensus_min_fraction < 1:
        raise ToolError(
            "--same-size-within-cluster-consensus-min-fraction must be at least 0 and less than 1"
        )
    if args.same_size_max_assignment_distance < 0:
        raise ToolError("--same-size-max-assignment-distance cannot be negative")
    if (
        args.homozygous_interruption_min_fraction is not None
        and not 0 <= args.homozygous_interruption_min_fraction < 1
    ):
        raise ToolError("--homozygous-interruption-min-fraction must be at least 0 and less than 1")
    formats = [item.strip().lower() for item in args.figure_formats.split(",") if item.strip()]
    invalid_formats = sorted(set(formats) - {"png", "pdf", "svg"})
    if invalid_formats:
        raise ToolError(f"unsupported figure formats: {', '.join(invalid_formats)}")
    if args.output.exists() and any(args.output.iterdir()) and not args.overwrite:
        raise ToolError(
            f"output directory is not empty: {args.output}; choose a new directory or add --overwrite"
        )


def _configure_logging(output: Path, level: str) -> None:
    output.mkdir(parents=True, exist_ok=True)
    handlers: list[logging.Handler] = [logging.StreamHandler(sys.stderr)]
    handlers.append(logging.FileHandler(output / "run.log", encoding="utf-8"))
    logging.basicConfig(
        level=getattr(logging, level),
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=handlers,
        force=True,
    )


def _prepare_output_directory(output: Path, overwrite: bool) -> None:
    """Create a clean run surface while preserving unrelated user files."""
    output.mkdir(parents=True, exist_ok=True)
    if not overwrite:
        return
    for name in (
        "all_loci.calls.csv",
        "all_loci.calls.xlsx",
        "all_loci.qc.csv",
        "input_validation.csv",
        "run_manifest.json",
        "run.log",
    ):
        target = output / name
        if target.is_file() or target.is_symlink():
            target.unlink()
    for name in ("calls_by_locus", "read_level", "figure_data", "figures"):
        target = output / name
        if target.is_dir():
            shutil.rmtree(target)


def run(args: argparse.Namespace) -> int:
    _validate_arguments(args)
    started = datetime.now(timezone.utc)
    catalog_loci = load_catalog(args.catalog)
    annotations = load_annotations(args.annotation, args.annotation_sheet)
    svg_files = discover_svg_files(args.svg_root)
    if not svg_files:
        raise ToolError(f"no .svg or .svg.gz files found recursively under {args.svg_root}")
    genotype_overrides = load_genotype_overrides(args.genotypes)
    file_overrides = load_file_overrides(args.file_map, args.svg_root)
    selected_loci = [
        value.strip()
        for group in args.locus
        for value in group.split(",")
        if value.strip()
    ]
    if selected_loci:
        selected_names = {
            normalise_key(value): value for value in selected_loci if normalise_key(value)
        }
        selected_keys = set(selected_names)
        matched_keys: set[str] = set()
        catalog_index = build_catalog_index(catalog_loci)
        scoped_files: list[Path] = []
        for svg_path in svg_files:
            try:
                override = file_overrides.get(str(svg_path.resolve()))
                _, _, svg_locus = parse_svg_identity(svg_path, args.svg_root, override)
            except Exception:
                continue
            catalog = resolve_catalog_locus(svg_locus, catalog_index)
            matched = selected_keys & _selectable_locus_keys(svg_locus, catalog)
            if matched:
                scoped_files.append(svg_path)
                matched_keys.update(matched)
        unmatched = selected_keys - matched_keys
        if unmatched:
            raise ToolError(
                "--locus value(s) matched no SVG/catalog locus: "
                + ", ".join(selected_names[key] for key in sorted(unmatched))
            )
        svg_files = scoped_files
        if not svg_files:
            raise ToolError(
                "no SVG files match --locus " + ", ".join(selected_loci)
            )
    same_size_consensus_threshold = (
        args.homozygous_interruption_min_fraction
        if args.homozygous_interruption_min_fraction is not None
        else args.same_size_within_cluster_consensus_min_fraction
    )
    config = RunConfig(
        min_spanning_reads=args.min_spanning_reads,
        min_cluster_reads=args.min_cluster_reads,
        max_near_motif_distance=args.max_near_motif_distance,
        near_cluster_parent_ratio=args.near_cluster_parent_ratio,
        strong_cluster_fraction=args.strong_cluster_fraction,
        heterozygous_interruption_min_fraction=(
            args.heterozygous_interruption_min_fraction
        ),
        same_size_min_cluster_reads=args.same_size_min_cluster_reads,
        same_size_min_cluster_fraction=args.same_size_min_cluster_fraction,
        same_size_min_top_two_coverage=args.same_size_min_top_two_coverage,
        same_size_within_cluster_consensus_min_fraction=same_size_consensus_threshold,
        same_size_max_reported_clusters=args.same_size_max_reported_clusters,
        same_size_max_assignment_distance=args.same_size_max_assignment_distance,
        homozygous_deduplication=args.homozygous_deduplication,
        component_mode=args.component_mode,
        unresolved_strand_policy=args.unresolved_strand_policy,
        allow_duplicate_sample_locus=args.allow_duplicate_sample_locus,
    )
    audit_rows = build_input_audit(
        svg_files, args.svg_root, catalog_loci, annotations, file_overrides
    )
    audit_columns = [
        "svg_locus", "n_svg_files", "catalog_locus", "catalog_gene",
        "annotation_key", "strand", "reporting_orientation", "primary_component",
        "n_catalog_components", "status", "notes",
    ]
    unconfigured = [row for row in audit_rows if row["status"] != "ready"]
    _prepare_output_directory(args.output, args.overwrite)
    _configure_logging(args.output, args.log_level)
    logging.info("MotifSTaR version %s", VERSION)
    if args.homozygous_interruption_min_fraction is not None:
        logging.warning(
            "--homozygous-interruption-min-fraction is deprecated; its value is being "
            "used as --same-size-within-cluster-consensus-min-fraction"
        )
    _write_csv_atomic(args.output / "input_validation.csv", audit_rows, audit_columns)
    logging.info(
        "Input audit: %d SVGs, %d observed loci, %d unconfigured loci",
        len(svg_files),
        len(audit_rows),
        len(unconfigured),
    )

    manifest: dict[str, Any] = {
        "software": "MotifSTaR",
        "version": VERSION,
        "started_utc": started.isoformat(),
        "python": sys.version.split()[0],
        "inputs": {
            "svg_root": str(args.svg_root.resolve()),
            "catalog": str(args.catalog.resolve()),
            "catalog_sha256": sha256_file(args.catalog),
            "annotation": str(args.annotation.resolve()),
            "annotation_sha256": sha256_file(args.annotation),
            "genotypes": str(args.genotypes.resolve()) if args.genotypes else None,
            "file_map": str(args.file_map.resolve()) if args.file_map else None,
        },
        "configuration": _serialise_config(config),
        "compound_pooling_policy": (
            "Compare the complete ordered size vectors of all rendered repeat "
            "blocks, including unreported companions; pool only when equal."
        ),
        "allele_order_policy": (
            "Panel-specific calls follow SVG order: top is allele1, bottom is "
            "allele2. Identical pooled structures may be reported for both. "
            "Distinct pooled same-size structures follow panel order when "
            "uniquely supported by panel-specific exact matches of at "
            "least max(min_spanning_reads, same_size_min_cluster_reads) and "
            "strictly more than 50% of that panel's eligible reads. Otherwise "
            "retain the original deterministic pooled-call order without an "
            "additional REVIEW or no-call solely for ordering uncertainty. "
            "Pooled labels do not guarantee panel assignment or parental phase."
        ),
        "orientation_note": (
            "When reporting_orientation is absent, companion strand '+' is explicitly "
            "mapped to reference and '-' to reverse_complement and flagged as inferred."
        ),
        "depth_definition_note": (
            "read_depth_alleleN and supporting_read_depth_alleleN count reads matching the "
            "reported position-wise consensus structure; full_spanning_read_depth_alleleN counts continuous "
            "left-flank/repeat/right-flank rows; callable_full_spanning_read_depth_alleleN "
            "also requires an unambiguous, motif-divisible repeat sequence; consensus_eligible_read_depth_alleleN "
            "additionally requires the expected allele length and, for same-size genotypes, unique assignment "
            "to that allele's selected sequence cluster. Different-size panels are evaluated separately with "
            "a strict > threshold. Same-size panels are pooled, exact-path allele seeds must meet both support "
            "thresholds, minor paths are assigned only to a unique nearby seed, and each selected cluster gets "
            "its own strict > consensus. No allele is forced to an uninterrupted reference path."
        ),
        "selected_loci": selected_loci,
        "workers": args.workers,
        "input_audit": {
            "n_svg_files": len(svg_files),
            "n_observed_loci": len(audit_rows),
            "n_unconfigured_loci": len(unconfigured),
            "unconfigured_loci": [row["svg_locus"] for row in unconfigured],
        },
    }
    if args.validate_only:
        manifest["completed_utc"] = datetime.now(timezone.utc).isoformat()
        manifest["status"] = "validation_only"
        _write_json_atomic(args.output / "run_manifest.json", manifest)
        return 2 if args.strict and unconfigured else 0

    context = WorkerContext(
        svg_root=str(args.svg_root.resolve()),
        catalog_loci=catalog_loci,
        annotation_rows=annotations,
        genotype_overrides=genotype_overrides,
        file_overrides=file_overrides,
        selected_loci=selected_loci,
        config=config,
    )
    records, processing_metrics = process_corpus(
        svg_files=svg_files,
        context=context,
        output_dir=args.output,
        workers=args.workers,
        write_read_level=args.write_read_level,
    )
    output_metrics = write_tabular_outputs(records, args.output, config)
    figure_metrics: dict[str, Any] = {"enabled": False}
    figure_warning = ""
    if args.figures:
        try:
            formats = [item.strip().lower() for item in args.figure_formats.split(",") if item.strip()]
            figure_metrics = {
                "enabled": True,
                **generate_figures(
                    records,
                    catalog_loci,
                    args.output,
                    formats,
                    args.dpi,
                    args.per_locus_figures,
                ),
            }
        except Exception as exc:
            figure_warning = str(exc)
            logging.exception("Figure generation failed")
            figure_metrics = {"enabled": True, "error": figure_warning}
            if args.strict:
                raise

    manifest.update(
        {
            "processing": processing_metrics,
            "outputs": output_metrics,
            "figures": figure_metrics,
            "completed_utc": datetime.now(timezone.utc).isoformat(),
            "status": "completed_with_figure_warning" if figure_warning else "completed",
        }
    )
    _write_json_atomic(args.output / "run_manifest.json", manifest)
    error_count = sum(record["qc"].get("qc_status") == "ERROR" for record in records)
    review_count = sum(record["qc"].get("qc_status") == "REVIEW" for record in records)
    logging.info(
        "Completed: %d component calls, %d errors, %d review calls",
        len(records), error_count, review_count,
    )
    if args.strict and (error_count or unconfigured):
        return 2
    if args.fail_on_review and review_count:
        return 3
    return 0


def main(argv: Sequence[str] | None = None) -> int:
    parser = build_argument_parser()
    args = parser.parse_args(argv)
    try:
        return run(args)
    except ToolError as exc:
        parser.error(str(exc))
    except KeyboardInterrupt:
        logging.error("Interrupted")
        return 130
    return 2


if __name__ == "__main__":
    raise SystemExit(main())
